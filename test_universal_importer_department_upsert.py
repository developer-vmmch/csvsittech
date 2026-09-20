"""
Comprehensive verification tests for Universal Importer:
Department Unique Constraint and Master Entity Safe Upsert.

Covers all 6 required test scenarios:
Test 1: Excel contains an existing department (e.g. EMR / EMERGENCY MEDICINE matching EM / EMERGENCY MEDICINE).
        Expected: UPDATE, no unique constraint error.
Test 2: Excel contains a completely new department.
        Expected: CREATE.
Test 3: Excel contains the same department twice.
        Expected: No duplicate, duplicate detected in preview, safe update in import.
Test 4: Excel contains: "   GENERAL   MEDICINE   ", DB contains "GENERAL MEDICINE".
        Expected: Match existing department as UPDATE.
Test 5: Import the same Excel file twice.
        Expected: First import creates/updates; second import does not duplicate. Idempotent.
Test 6: Force an import error halfway through.
        Expected: Transaction rollback, no partial database changes.
"""

import os
import sys
import io
import django
from openpyxl import load_workbook

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'vmmc_erp.settings')
django.setup()

from apps.patients.models import Department
from apps.lab.models import Diagnosis, Investigation, Parameter
from apps.lab.universal_importer import (
    validate_import,
    import_universal_master,
    generate_blank_template,
    UniversalImportError,
    DepartmentResolver,
    normalize_text
)


def create_workbook_with_departments(dept_rows):
    """
    Helper to create an in-memory Excel workbook based on the standard template,
    populating the 'Departments' sheet with given rows: [(name, code, status), ...].
    """
    raw_tpl = generate_blank_template()
    wb = load_workbook(raw_tpl)
    ws = wb['Departments']
    for row in dept_rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_1_existing_department_update():
    print("\n--- Running Test 1: Existing department update (no UNIQUE constraint error) ---")
    # In DB, EMERGENCY MEDICINE exists with code 'EM'
    em_dept = Department.objects.filter(name__iexact='EMERGENCY MEDICINE').first()
    assert em_dept is not None, "EMERGENCY MEDICINE must exist in DB"
    initial_dept_count = Department.objects.count()
    orig_code = em_dept.code
    orig_active = em_dept.is_active

    # Excel provides 'EMR' and 'EMERGENCY MEDICINE'
    buf = create_workbook_with_departments([
        ('EMERGENCY MEDICINE', 'EMR', 'Active')
    ])

    # 1. Preview
    val = validate_import(buf)
    assert val['valid'] is True, f"Validation failed: {val['errors']}"
    d_stat = val['sheet_stats']['Departments']
    print(f"Preview Stats: {d_stat}")
    assert d_stat['existing'] == 1, f"Expected 1 existing (update), got {d_stat['existing']}"
    assert d_stat['new'] == 0, f"Expected 0 new, got {d_stat['new']}"

    # 2. Import
    buf.seek(0)
    counts = import_universal_master(buf)
    print(f"Import Counts: {counts}")
    assert counts['departments_updated'] == 1, f"Expected 1 updated, got {counts['departments_updated']}"
    assert counts['departments_created'] == 0, f"Expected 0 created, got {counts['departments_created']}"
    assert Department.objects.count() == initial_dept_count, "Total department count should not increase"

    # Reload department from DB
    updated_dept = Department.objects.filter(name__iexact='EMERGENCY MEDICINE').first()
    assert updated_dept is not None
    assert updated_dept.id == em_dept.id, "Should update the exact same record"
    print("✓ Test 1 PASSED: Existing department updated cleanly without constraint error.")


def test_2_new_department_create():
    print("\n--- Running Test 2: Completely new department creation ---")
    test_code = 'TEST_CARD'
    test_name = 'TEST CARDIOLOGY SPECIALTY'
    Department.objects.filter(models_q_code_or_name(test_code, test_name)).delete()

    initial_dept_count = Department.objects.count()

    buf = create_workbook_with_departments([
        (test_name, test_code, 'Active')
    ])

    # 1. Preview
    val = validate_import(buf)
    assert val['valid'] is True, f"Validation failed: {val['errors']}"
    d_stat = val['sheet_stats']['Departments']
    print(f"Preview Stats: {d_stat}")
    assert d_stat['new'] == 1, f"Expected 1 new, got {d_stat['new']}"
    assert d_stat['existing'] == 0, f"Expected 0 existing, got {d_stat['existing']}"

    # 2. Import
    buf.seek(0)
    counts = import_universal_master(buf)
    print(f"Import Counts: {counts}")
    assert counts['departments_created'] == 1, f"Expected 1 created, got {counts['departments_created']}"
    assert Department.objects.count() == initial_dept_count + 1

    created_dept = Department.objects.filter(code=test_code).first()
    assert created_dept is not None
    assert created_dept.name == test_name

    # Clean up test department
    created_dept.delete()
    print("✓ Test 2 PASSED: Completely new department created successfully.")


def models_q_code_or_name(code, name):
    from django.db.models import Q
    return Q(code__iexact=code) | Q(name__iexact=name)


def test_3_duplicate_inside_excel():
    print("\n--- Running Test 3: Duplicate inside the Excel file ---")
    initial_dept_count = Department.objects.count()

    # Provide GENERAL MEDICINE twice in the same workbook
    buf = create_workbook_with_departments([
        ('GENERAL MEDICINE', 'GM', 'Active'),
        ('GENERAL MEDICINE', 'GM', 'Active'),
    ])

    # 1. Preview
    val = validate_import(buf)
    assert val['valid'] is True
    d_stat = val['sheet_stats']['Departments']
    print(f"Preview Stats: {d_stat}")
    assert d_stat['duplicates'] >= 1, f"Expected >= 1 duplicates, got {d_stat['duplicates']}"

    # 2. Import
    buf.seek(0)
    counts = import_universal_master(buf)
    print(f"Import Counts: {counts}")
    assert counts['departments_created'] == 0
    assert counts['departments_updated'] == 2  # Both rows processed/merged into update
    assert Department.objects.count() == initial_dept_count, "No duplicate department should be created"
    print("✓ Test 3 PASSED: In-file duplicate handled cleanly without duplicate creation.")


def test_4_whitespace_normalization():
    print("\n--- Running Test 4: Whitespace normalization ---")
    initial_dept_count = Department.objects.count()

    # Excel contains unusual leading/trailing/multiple internal whitespace
    weird_name = "   GENERAL    MEDICINE   "
    buf = create_workbook_with_departments([
        (weird_name, 'GM', 'Active')
    ])

    # 1. Preview
    val = validate_import(buf)
    assert val['valid'] is True
    d_stat = val['sheet_stats']['Departments']
    print(f"Preview Stats: {d_stat}")
    assert d_stat['existing'] == 1, f"Expected 1 existing update, got {d_stat['existing']}"
    assert d_stat['new'] == 0

    # 2. Import
    buf.seek(0)
    counts = import_universal_master(buf)
    print(f"Import Counts: {counts}")
    assert counts['departments_created'] == 0
    assert counts['departments_updated'] == 1
    assert Department.objects.count() == initial_dept_count

    gm = Department.objects.filter(code='GM').first()
    assert gm is not None
    assert gm.name == "GENERAL MEDICINE", f"Should normalize name, got '{gm.name}'"
    print("✓ Test 4 PASSED: Whitespace normalized and existing department matched.")


def test_5_idempotent_import():
    print("\n--- Running Test 5: Idempotent re-import ---")
    test_code = 'TEST_IDEMP'
    test_name = 'TEST IDEMPOTENT DEPT'
    Department.objects.filter(models_q_code_or_name(test_code, test_name)).delete()

    buf_data = [
        ('GENERAL MEDICINE', 'GM', 'Active'),
        (test_name, test_code, 'Active')
    ]

    # First import: creates 1, updates 1
    buf1 = create_workbook_with_departments(buf_data)
    counts1 = import_universal_master(buf1)
    print(f"First Import: {counts1}")
    assert counts1['departments_created'] == 1
    assert counts1['departments_updated'] == 1

    count_after_first = Department.objects.count()

    # Second import with identical file: creates 0, updates 2
    buf2 = create_workbook_with_departments(buf_data)
    counts2 = import_universal_master(buf2)
    print(f"Second Import: {counts2}")
    assert counts2['departments_created'] == 0, f"Expected 0 created, got {counts2['departments_created']}"
    assert counts2['departments_updated'] == 2, f"Expected 2 updated, got {counts2['departments_updated']}"
    assert Department.objects.count() == count_after_first, "Count must remain identical on second import"

    # Cleanup
    Department.objects.filter(code=test_code).delete()
    print("✓ Test 5 PASSED: Re-importing identical workbook is 100% idempotent.")


def test_6_transaction_rollback_on_failure():
    print("\n--- Running Test 6: Transaction rollback on failure ---")
    test_code = 'TEST_ROLLBACK_DEPT'
    test_name = 'TEST ROLLBACK DEPT'
    Department.objects.filter(models_q_code_or_name(test_code, test_name)).delete()

    initial_dept_count = Department.objects.count()

    # Build workbook with 1 valid new department, but we will monkeypatch to simulate
    # an unexpected failure halfway through the atomic block (e.g. during Parameters)
    buf = create_workbook_with_departments([
        (test_name, test_code, 'Active')
    ])

    import apps.lab.universal_importer as ui
    orig_param_upsert = ui.ParameterResolver.upsert

    def faulty_upsert(*args, **kwargs):
        raise RuntimeError("Simulated mid-import catastrophic failure")

    ui.ParameterResolver.upsert = faulty_upsert

    try:
        # Also add a parameter so the faulty method is triggered
        raw_wb = load_workbook(buf)
        raw_wb['Parameters'].append(['Faulty Param', 'F_PARAM', 'NUMERIC', 'mg', 'Active'])
        buf_fail = io.BytesIO()
        raw_wb.save(buf_fail)
        buf_fail.seek(0)

        rolled_back = False
        try:
            import_universal_master(buf_fail)
        except UniversalImportError as uie:
            rolled_back = True
            print(f"Caught expected UniversalImportError: {uie.to_dict()}")

        assert rolled_back, "Expected UniversalImportError to be raised"

        # Verify DB: the new department in row 1 MUST NOT exist because transaction rolled back!
        assert not Department.objects.filter(code=test_code).exists(), "Partially imported department must not exist after rollback!"
        assert Department.objects.count() == initial_dept_count, "Department count must not change on rollback!"
        print("✓ Test 6 PASSED: Transaction rollback verified, 0 partial changes committed.")
    finally:
        ui.ParameterResolver.upsert = orig_param_upsert
        Department.objects.filter(models_q_code_or_name(test_code, test_name)).delete()


if __name__ == '__main__':
    print("==================================================")
    print("STARTING UNIVERSAL IMPORTER DEPARTMENT TESTS")
    print("==================================================")
    test_1_existing_department_update()
    test_2_new_department_create()
    test_3_duplicate_inside_excel()
    test_4_whitespace_normalization()
    test_5_idempotent_import()
    test_6_transaction_rollback_on_failure()
    print("==================================================")
    print("ALL 6 TESTS COMPLETED SUCCESSFULLY!")
    print("==================================================")

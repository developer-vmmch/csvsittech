import re
import os

with open('apps/lab/import_views.py', 'r') as f:
    content = f.read()

new_agegroup_import = """def api_agegroup_preview(request):
    if request.method != 'POST': return JsonResponse({'success': False, 'message': 'Invalid method'})
    if 'file' not in request.FILES: return JsonResponse({'success': False, 'message': 'No file uploaded'})
        
    file = request.FILES['file']
    filename = file.name
    
    try:
        import pandas as pd
        if filename.endswith('.csv'): 
            import csv
            import io
            decoded_file = file.read().decode('utf-8')
            reader = csv.reader(io.StringIO(decoded_file))
            raw_headers = next(reader, [])
            rows_data = list(reader)
        elif filename.endswith('.xlsx'): 
            from openpyxl import load_workbook
            try:
                workbook = load_workbook(file, read_only=True, data_only=True)
            except Exception as e:
                import traceback
                traceback.print_exc()
                print("AGE GROUP IMPORT ERROR:", repr(e))
                return JsonResponse({'success': False, 'error_type': 'VALIDATION_ERROR', 'errors': [{'row': 0, 'field': 'File', 'value': '', 'message': 'Could not read Excel file'}]})
            if 'Template' not in workbook.sheetnames:
                return JsonResponse({'success': False, 'error_type': 'VALIDATION_ERROR', 'errors': [{'row': 0, 'field': 'Worksheet', 'value': '', 'message': "Import worksheet 'Template' not found."}]})
            worksheet = workbook['Template']
            raw_headers = []
            rows_data = []
            for i, row in enumerate(worksheet.iter_rows(values_only=True)):
                if i == 0:
                    raw_headers = list(row)
                else:
                    if any(cell is not None and str(cell).strip() != '' for cell in row):
                        rows_data.append(list(row))
        else: 
            return JsonResponse({'success': False, 'message': 'Unsupported format'})
            
        required = ['age_group_code', 'age_group_name', 'minimum_age', 'maximum_age', 'age_unit', 'gender', 'pregnancy_applicable', 'display_order', 'active']
        headers = [str(h).strip().lower() if h else "" for h in raw_headers]
        
        for req in required:
            if req not in headers:
                return JsonResponse({'success': False, 'error_type': 'VALIDATION_ERROR', 'errors': [{'row': 1, 'field': 'Header', 'value': '', 'message': f'Missing required column: {req}'}]})
                
        col_idx = {req: headers.index(req) for req in required}
        
        preview_data = []
        validation_errors = []
        file_codes = set()
        
        from .models import AgeGroup
        existing = set(AgeGroup.objects.values_list('code', flat=True))
        
        valid_rows = 0
        invalid_rows = 0
        duplicate_rows = 0
        
        def parse_optional_number(value):
            if value is None: return None
            val_str = str(value).strip()
            if val_str == "": return None
            try:
                f = float(val_str)
                if f.is_integer(): return int(f)
                return f
            except Exception:
                return "INVALID"
                
        def safe_str(val):
            if val is None: return ""
            return str(val).strip()
        
        for i, row in enumerate(rows_data):
            row_num = i + 2
            while len(row) < len(headers): row.append(None)
            
            code = safe_str(row[col_idx['age_group_code']])
            name = safe_str(row[col_idx['age_group_name']])
            min_age = parse_optional_number(row[col_idx['minimum_age']])
            max_age = parse_optional_number(row[col_idx['maximum_age']])
            disp_order = parse_optional_number(row[col_idx['display_order']])
            
            age_unit = safe_str(row[col_idx['age_unit']]).capitalize()
            gender = safe_str(row[col_idx['gender']]).capitalize()
            preg_str = safe_str(row[col_idx['pregnancy_applicable']]).lower()
            active_str = safe_str(row[col_idx['active']]).lower()
            
            row_errors = []
            
            if not code: row_errors.append({'row': row_num, 'field': 'age_group_code', 'value': '', 'message': 'Age group code required'})
            if not name: row_errors.append({'row': row_num, 'field': 'age_group_name', 'value': '', 'message': 'Age group name required'})
            
            if min_age == 'INVALID': row_errors.append({'row': row_num, 'field': 'minimum_age', 'value': row[col_idx['minimum_age']], 'message': 'INVALID_NUMERIC_VALUE'})
            if max_age == 'INVALID': row_errors.append({'row': row_num, 'field': 'maximum_age', 'value': row[col_idx['maximum_age']], 'message': 'INVALID_NUMERIC_VALUE'})
            if disp_order == 'INVALID': row_errors.append({'row': row_num, 'field': 'display_order', 'value': row[col_idx['display_order']], 'message': 'INVALID_NUMERIC_VALUE'})
            
            if min_age != 'INVALID' and max_age != 'INVALID' and min_age is not None and max_age is not None and min_age > max_age:
                row_errors.append({'row': row_num, 'field': 'age_range', 'value': f'{min_age}-{max_age}', 'message': 'INVALID_AGE_RANGE: minimum_age > maximum_age'})
                
            if age_unit not in ['Days', 'Months', 'Years']:
                row_errors.append({'row': row_num, 'field': 'age_unit', 'value': row[col_idx['age_unit']], 'message': 'INVALID_AGE_UNIT: Allowed values: Days, Months, Years.'})
            if gender not in ['All', 'Male', 'Female']:
                row_errors.append({'row': row_num, 'field': 'gender', 'value': row[col_idx['gender']], 'message': 'INVALID_GENDER: Allowed values: All, Male, Female.'})
            
            preg = preg_str in ['yes', 'y', '1', 'true']
            if preg_str not in ['yes', 'no', 'y', 'n', '1', '0', 'true', 'false', '']:
                row_errors.append({'row': row_num, 'field': 'pregnancy_applicable', 'value': row[col_idx['pregnancy_applicable']], 'message': 'INVALID_PREGNANCY_VALUE'})
                
            active = active_str in ['yes', 'y', '1', 'true', '']
            if active_str not in ['yes', 'no', 'y', 'n', '1', '0', 'true', 'false', '']:
                row_errors.append({'row': row_num, 'field': 'active', 'value': row[col_idx['active']], 'message': 'INVALID_ACTIVE_VALUE'})
            
            is_duplicate = False
            if code in existing:
                is_duplicate = True
            elif code in file_codes:
                row_errors.append({'row': row_num, 'field': 'age_group_code', 'value': code, 'message': 'DUPLICATE_AGE_GROUP_CODE_IN_FILE'})
                
            if code: file_codes.add(code)
            
            if row_errors:
                status = 'Error'
                invalid_rows += 1
                validation_errors.extend(row_errors)
                error_msg = row_errors[0]['message']
            elif is_duplicate:
                status = 'Duplicate'
                duplicate_rows += 1
                error_msg = 'EXISTING_AGE_GROUP'
            else:
                status = 'Valid'
                valid_rows += 1
                error_msg = ''
                
            preview_data.append({
                'row_num': row_num,
                'code': code,
                'name': name,
                'min_age': min_age if min_age != 'INVALID' and min_age is not None else '',
                'max_age': max_age if max_age != 'INVALID' and max_age is not None else '',
                'unit': age_unit,
                'gender': gender,
                'preg': preg,
                'disp_order': disp_order if disp_order != 'INVALID' and disp_order is not None else '',
                'active': active,
                'status': status,
                'error': error_msg
            })
            
        if validation_errors:
            return JsonResponse({
                'success': False,
                'error_type': 'VALIDATION_ERROR',
                'errors': validation_errors,
                'data': preview_data,
                'filename': filename
            })
            
        return JsonResponse({
            'success': True,
            'total_rows': len(rows_data),
            'valid_rows': valid_rows,
            'invalid_rows': invalid_rows,
            'duplicates': duplicate_rows,
            'errors': [],
            'data': preview_data,
            'filename': filename
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("AGE GROUP IMPORT ERROR:", repr(e))
        return JsonResponse({'success': False, 'message': str(e)})

def api_agegroup_import(request):
    import json
    from django.http import JsonResponse
    if request.method != 'POST': return JsonResponse({'success': False, 'message': 'Invalid method'})
    try:
        data = json.loads(request.body)
        filename = data.get('filename', 'Unknown')
        rows = data.get('rows', [])
        update_duplicates = data.get('update_duplicates', False)
        
        imported = updated = skipped = failed = 0
        from django.db import transaction
        from .models import AgeGroup, AgeGroupImportHistory
        
        with transaction.atomic():
            for row in rows:
                if row['status'] == 'Error':
                    failed += 1; continue
                try:
                    ag = AgeGroup.objects.filter(code=row['code']).first()
                    min_age = row['min_age'] if row['min_age'] != '' else 0
                    max_age = row['max_age'] if row['max_age'] != '' else None
                    disp = row['disp_order'] if row['disp_order'] != '' else 1
                    
                    if ag:
                        if update_duplicates or row['status'] == 'Duplicate':
                            if update_duplicates:
                                ag.label = row['name']
                                ag.min_age_value = min_age
                                ag.min_age_unit = row['unit']
                                ag.max_age_value = max_age
                                ag.max_age_unit = row['unit']
                                ag.gender = row['gender']
                                ag.pregnancy_applicable = row['preg']
                                ag.sort_order = disp
                                ag.is_active = row['active']
                                ag.save()
                                updated += 1
                            else:
                                skipped += 1
                        else:
                            skipped += 1
                    else:
                        AgeGroup.objects.create(
                            code=row['code'],
                            label=row['name'],
                            min_age_value=min_age,
                            min_age_unit=row['unit'],
                            max_age_value=max_age,
                            max_age_unit=row['unit'],
                            gender=row['gender'],
                            pregnancy_applicable=row['preg'],
                            sort_order=disp,
                            is_active=row['active']
                        )
                        imported += 1
                except Exception as e:
                    import traceback
                    traceback.print_exc()
                    print("AGE GROUP IMPORT ROW ERROR:", repr(e))
                    failed += 1
                    
            history = AgeGroupImportHistory.objects.create(
                file_name=filename, uploaded_by=request.user if request.user.is_authenticated else None,
                total_records=len(rows), imported=imported, updated=updated, skipped=skipped, failed=failed
            )
        return JsonResponse({'success': True, 'summary': {
            'total': len(rows), 'imported': imported, 'updated': updated, 'skipped': skipped, 'failed': failed, 'history_id': history.id
        }})
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("AGE GROUP IMPORT ERROR:", repr(e))
        return JsonResponse({'success': False, 'message': str(e)})"""

start_idx = content.find('def api_agegroup_preview(request):')
end_idx = content.find('# --- Parameter Import ---')

if start_idx == -1 or end_idx == -1:
    print("Could not find boundaries for age group import.")
else:
    final_content = content[:start_idx] + new_agegroup_import + "\n\n" + content[end_idx:]
    with open('apps/lab/import_views.py', 'w') as f:
        f.write(final_content)
    print("Age group import functions replaced successfully.")

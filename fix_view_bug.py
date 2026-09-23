import os
import sys

views_path = r'c:\Users\Admin\Desktop\erp1\apps\lab\views.py'
with open(views_path, 'r', encoding='utf-8') as f:
    views_content = f.read()

target_class = """class ReferenceRangeGridView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'lab_master.reference_range.view'
    template_name = 'lab/master/reference_range_grid.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['reference_ranges'] = ParameterReferenceRange.objects.select_related('investigation_parameter', 'investigation_parameter__investigation', 'age_group').order_by('-id')
        context['parameters'] = InvestigationParameter.objects.select_related('investigation').filter(is_active=True).order_by('name')
        context['age_groups'] = AgeGroup.objects.filter(is_active=True).order_by('sort_order', 'label')
        return context

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.reference_range.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:reference_range_grid')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io
        from .models import ParameterReferenceRange

        queryset = ParameterReferenceRange.objects.all().select_related('investigation_parameter__investigation', 'age_group')
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['parameter_code', 'age_group_code', 'gender', 'pregnancy', 'range_type', 'min_value', 'max_value', 'reference_text', 'unit', 'method', 'remarks', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.investigation_parameter.code if obj.investigation_parameter else '',
                obj.age_group.code if obj.age_group else '',
                obj.gender or 'All',
                'Yes' if obj.pregnancy else 'No',
                obj.range_type or 'Numeric',
                obj.min_value if obj.min_value is not None else '',
                obj.max_value if obj.max_value is not None else '',
                obj.reference_text or '',
                obj.unit or '',
                obj.method or '',
                obj.remarks or '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="reference_range_export.xlsx"'
        return response

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['investigations'] = Investigation.objects.filter(is_active=True)
        context['age_groups'] = AgeGroup.objects.filter(is_active=True).order_by('sort_order')
        context['diagnoses'] = Diagnosis.objects.filter(is_active=True)
        return context"""

new_class = """class ReferenceRangeGridView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.reference_range.view'
    template_name = 'lab/master/reference_range_grid.html'
    model = ParameterReferenceRange
    context_object_name = 'reference_ranges'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset().select_related('investigation_parameter', 'investigation_parameter__investigation', 'age_group').order_by('-id')
        search = self.request.GET.get('search', '').strip()
        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(investigation_parameter__name__icontains=search) | 
                Q(investigation_parameter__code__icontains=search) |
                Q(reference_text__icontains=search) |
                Q(age_group__label__icontains=search)
            )
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['parameters'] = InvestigationParameter.objects.select_related('investigation').filter(is_active=True).order_by('name')
        context['investigations'] = Investigation.objects.filter(is_active=True)
        context['age_groups'] = AgeGroup.objects.filter(is_active=True).order_by('sort_order', 'label')
        context['diagnoses'] = Diagnosis.objects.filter(is_active=True)
        return context

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.reference_range.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:reference_range_grid')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io
        from .models import ParameterReferenceRange

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['parameter_code', 'age_group_code', 'gender', 'pregnancy', 'range_type', 'min_value', 'max_value', 'reference_text', 'unit', 'method', 'remarks', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.investigation_parameter.code if obj.investigation_parameter else '',
                obj.age_group.code if obj.age_group else '',
                obj.gender or 'All',
                'Yes' if obj.pregnancy else 'No',
                obj.range_type or 'Numeric',
                obj.min_value if obj.min_value is not None else '',
                obj.max_value if obj.max_value is not None else '',
                obj.reference_text or '',
                obj.unit or '',
                obj.method or '',
                obj.remarks or '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="reference_range_export.xlsx"'
        return response"""

if target_class in views_content:
    views_content = views_content.replace(target_class, new_class)
    with open(views_path, 'w', encoding='utf-8') as f:
        f.write(views_content)
    print("Updated views.py successfully")
else:
    print("Could not find the target class in views.py")

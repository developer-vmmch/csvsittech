
from django.utils import timezone
from datetime import timedelta

def get_default_date_range(request, from_param='from_date', to_param='to_param'):
    from_raw = request.GET.get(from_param)
    to_raw = request.GET.get(to_param)
    
    if from_raw is None:
        from_date = (timezone.localdate() - timedelta(days=6)).strftime('%Y-%m-%d')
    else:
        from_date = from_raw.strip()
        
    if to_raw is None:
        to_date = timezone.localdate().strftime('%Y-%m-%d')
    else:
        to_date = to_raw.strip()
        
    return from_date, to_date

from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse_lazy, reverse
from django.views.generic import ListView, CreateView, UpdateView, DeleteView, TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.auth.decorators import login_required
from apps.core.mixins import MenuAccessRequiredMixin, GranularPermissionRequiredMixin
from django.contrib import messages
from django.http import JsonResponse
import json
from decimal import Decimal
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from apps.patients.models import Patient, PatientVisit
from django.utils import timezone

from .models import (
    Diagnosis, LabDiagnosis, Investigation, Parameter, AgeGroup,
    InvestigationParameter,
    ParameterReferenceRange, PatientInvestigationOrder, PatientInvestigationResult,
    LabDepartment, SampleType,
    ChiefComplaint, HospitalService, StagingDiagnosis, StagingInvestigation,
    ServiceRequest, ServiceRequestDiagnosis, ServiceRequestInvestigation,
    PatientVisitDiagnosis
)
from .forms import (
    DiagnosisForm, LabDiagnosisForm, InvestigationForm, ParameterForm, AgeGroupForm,
    InvestigationParameterForm, ServiceRequestForm
)
from apps.patients.models import Patient

# --- Master Entities CRUD ---

class BaseLabMasterView(LoginRequiredMixin, MenuAccessRequiredMixin):
    menu_key = 'administration' # Fallback for now

class DiagnosisListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.diagnosis.view'
    model = Diagnosis
    template_name = 'lab/master/diagnosis_list.html'
    context_object_name = 'diagnoses'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(name__icontains=search) | models.Q(code__icontains=search) | models.Q(chapter__icontains=search))
        return qs
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(
                models.Q(name__icontains=search) | 
                models.Q(code__icontains=search) | 
                models.Q(chapter__icontains=search)
            )
        return qs


    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.diagnosis.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:diagnosis_list')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['icd_code', 'diagnosis_name', 'category', 'synonyms', 'class_kind', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.code,
                obj.name,
                obj.chapter or '',
                obj.synonyms or '',
                obj.class_kind or '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="diagnosis_export.xlsx"'
        return response

    def render_to_response(self, context, **response_kwargs):
        if self.request.headers.get('Accept') == 'application/json' or self.request.GET.get('format') == 'json':
            qs = self.get_queryset()
            # If no pagination applied to json, return all or paginate according to request
            results = [{
                'id': d.id,
                'name': d.name,
                'code': d.code,
                'status': 'Active' if d.is_active else 'Inactive'
            } for d in qs]
            return JsonResponse({'diagnoses': results})
        return super().render_to_response(context, **response_kwargs)

@csrf_exempt
@login_required
def api_diagnosis_save(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            diag_id = data.get('id')
            name = data.get('name', '').strip()
            code = data.get('code', '').strip()
            chapter = data.get('chapter', '').strip()
            is_active = data.get('is_active', True)
            
            if not name:
                return JsonResponse({'status': 'error', 'message': 'Diagnosis Name is required.'})
                
            if code == '': code = None
                
            # Duplicate check
            duplicate_code = Diagnosis.objects.filter(code__iexact=code) if code else Diagnosis.objects.none()
            duplicate_name = Diagnosis.objects.filter(name__iexact=name)
            if diag_id:
                duplicate_code = duplicate_code.exclude(id=diag_id)
                duplicate_name = duplicate_name.exclude(id=diag_id)
                
            if duplicate_name.exists():
                return JsonResponse({'status': 'error', 'message': f'A Diagnosis with Name "{name}" already exists.'})
                
            if duplicate_code.exists():
                return JsonResponse({'status': 'error', 'message': f'A Diagnosis with ICD Code {code} already exists.'})
                
            if diag_id:
                diag = Diagnosis.objects.get(id=diag_id)
                diag.name = name
                diag.code = code
                diag.chapter = chapter
                diag.is_active = is_active
                diag.save()
                msg = "Primary Diagnosis updated successfully."
            else:
                diag = Diagnosis.objects.create(
                    name=name,
                    code=code,
                    chapter=chapter,
                    is_active=is_active
                )
                msg = "Primary Diagnosis saved successfully."
                
            return JsonResponse({
                'status': 'success',
                'message': msg,
                'data': {
                    'id': diag.id,
                    'name': diag.name,
                    'code': diag.code,
                    'chapter': diag.chapter or '-',
                    'is_active': diag.is_active
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method'})

@csrf_exempt
@login_required
def api_diagnosis_delete(request, pk):
    if request.method == 'POST':
        try:
            diag = Diagnosis.objects.get(id=pk)
            diag.delete()
            return JsonResponse({'status': 'success', 'message': 'Primary Diagnosis deleted successfully.'})
        except Diagnosis.DoesNotExist:
            return JsonResponse({'status': 'error', 'message': 'Primary Diagnosis not found.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method'})

class DiagnosisCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_master.diagnosis.create'
    model = Diagnosis
    form_class = DiagnosisForm
    template_name = 'lab/master/diagnosis_form.html'
    success_url = reverse_lazy('lab:diagnosis_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Diagnosis created successfully!")
        return super().form_valid(form)

class DiagnosisUpdateView(LoginRequiredMixin, GranularPermissionRequiredMixin, UpdateView):
    permission_required = 'lab_master.diagnosis.update'
    model = Diagnosis
    form_class = DiagnosisForm
    template_name = 'lab/master/diagnosis_form.html'
    success_url = reverse_lazy('lab:diagnosis_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Diagnosis updated successfully!")
        return super().form_valid(form)

class LabDiagnosisListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.diagnosis.view'
    model = LabDiagnosis
    template_name = 'lab/master/lab_diagnosis_list.html'
    context_object_name = 'diagnoses'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(name__icontains=search) | models.Q(code__icontains=search))
        return qs

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.diagnosis.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:lab_diagnosis_list')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['icd_code', 'diagnosis_name', 'category', 'synonyms', 'class_kind', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.code,
                obj.name,
                obj.chapter or '',
                obj.synonyms or '',
                obj.class_kind or '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="diagnosis_export.xlsx"'
        return response

    def render_to_response(self, context, **response_kwargs):
        if self.request.headers.get('Accept') == 'application/json' or self.request.GET.get('format') == 'json':
            qs = self.get_queryset()
            results = [{
                'id': d.id,
                'name': d.name,
                'code': d.code,
                'status': 'Active' if d.is_active else 'Inactive'
            } for d in qs]
            return JsonResponse({'diagnoses': results})
        return super().render_to_response(context, **response_kwargs)

class LabDiagnosisCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_master.diagnosis.create'
    model = LabDiagnosis
    form_class = LabDiagnosisForm
    template_name = 'lab/master/lab_diagnosis_form.html'
    success_url = reverse_lazy('lab:lab_diagnosis_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Diagnosis created successfully!")
        return super().form_valid(form)

class LabDiagnosisUpdateView(LoginRequiredMixin, GranularPermissionRequiredMixin, UpdateView):
    permission_required = 'lab_master.diagnosis.update'
    model = LabDiagnosis
    form_class = LabDiagnosisForm
    template_name = 'lab/master/lab_diagnosis_form.html'
    success_url = reverse_lazy('lab:lab_diagnosis_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Diagnosis updated successfully!")
        return super().form_valid(form)

class InvestigationListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.investigation.view'
    model = Investigation
    template_name = 'lab/master/investigation_list.html'
    context_object_name = 'investigations'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(name__icontains=search) | models.Q(code__icontains=search))
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from .models import LabDepartment, SampleType
        context['departments'] = LabDepartment.objects.all()
        context['sample_types'] = SampleType.objects.all()
        return context

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.investigation.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:investigation_list')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['investigation_code', 'investigation_name', 'category', 'specimen_or_sample', 'parameters', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            params = obj.parameters.filter(is_active=True).values_list('name', flat=True)
            params_str = "; ".join([p for p in params if p])
            ws.append([
                obj.code,
                obj.name,
                obj.department.name if obj.department else '',
                obj.sample_type.name if obj.sample_type else '',
                params_str,
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="investigation_export.xlsx"'
        return response

    def get_queryset(self):
        return super().get_queryset().prefetch_related('parameters', 'department', 'sample_type')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = LabDepartment.objects.filter(is_active=True).order_by('name')
        context['sample_types'] = SampleType.objects.filter(is_active=True).order_by('name')
        return context

@csrf_exempt
@login_required
def api_investigation_save(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            inv_id = data.get('id')
            name = data.get('name', '').strip()
            code = data.get('code', '').strip()
            department_id = data.get('department_id')
            sample_type_id = data.get('sample_type_id')
            is_active = data.get('is_active', True)
            
            if not name or not code:
                return JsonResponse({'status': 'error', 'message': 'Name and Code are required.'})
                
            duplicate = Investigation.objects.filter(code__iexact=code)
            if inv_id:
                duplicate = duplicate.exclude(id=inv_id)
            if duplicate.exists():
                return JsonResponse({'status': 'error', 'message': f'Investigation with code {code} already exists.'})
                
            department = LabDepartment.objects.filter(id=department_id).first() if department_id else None
            sample_type = SampleType.objects.filter(id=sample_type_id).first() if sample_type_id else None

            if inv_id:
                inv = Investigation.objects.get(id=inv_id)
                inv.name = name
                inv.code = code
                inv.department = department
                inv.sample_type = sample_type
                inv.is_active = is_active
                inv.save()
                msg = "Investigation updated successfully."
            else:
                inv = Investigation.objects.create(
                    name=name,
                    code=code,
                    department=department,
                    sample_type=sample_type,
                    is_active=is_active
                )
                msg = "Investigation saved successfully."
                
            return JsonResponse({
                'status': 'success',
                'message': msg,
                'data': {
                    'id': inv.id,
                    'name': inv.name,
                    'code': inv.code,
                    'department_id': inv.department_id,
                    'department_name': inv.department.name if inv.department else '-',
                    'sample_type_id': inv.sample_type_id,
                    'sample_type_name': inv.sample_type.name if inv.sample_type else '-',
                    'param_count': inv.parameters.count(),
                    'is_active': inv.is_active
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method'})

from django.views.generic import DetailView
class InvestigationDetailView(LoginRequiredMixin, GranularPermissionRequiredMixin, DetailView):
    permission_required = 'lab_master.investigation.view'
    model = Investigation
    template_name = 'lab/master/investigation_detail.html'
    context_object_name = 'investigation'

class InvestigationCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_master.investigation.create'
    model = Investigation
    form_class = InvestigationForm
    template_name = 'lab/master/investigation_form.html'
    success_url = reverse_lazy('lab:investigation_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Investigation created successfully!")
        return super().form_valid(form)

class InvestigationUpdateView(LoginRequiredMixin, GranularPermissionRequiredMixin, UpdateView):
    permission_required = 'lab_master.investigation.update'
    model = Investigation
    form_class = InvestigationForm
    template_name = 'lab/master/investigation_form.html'
    success_url = reverse_lazy('lab:investigation_list')

    def form_valid(self, form):
        messages.success(self.request, "Investigation updated successfully!")
        return super().form_valid(form)

class ParameterListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.parameter.view'
    model = Parameter
    template_name = 'lab/master/parameter_list.html'
    context_object_name = 'parameters'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(name__icontains=search) | models.Q(code__icontains=search))
        return qs

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.parameter.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:parameter_list')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['parameter_code', 'investigation_code', 'parameter_name', 'short_name', 'result_type', 'unit', 'decimal_precision', 'display_order', 'active'])
        
        for obj in queryset.select_related('investigation').iterator(chunk_size=1000):
            ws.append([
                obj.code or '',
                obj.investigation.code if obj.investigation else '',
                obj.name or '',
                obj.short_name or '',
                obj.result_type or '',
                obj.unit or '',
                obj.decimal_precision if obj.decimal_precision is not None else '',
                obj.display_order if obj.display_order is not None else '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="parameter_export.xlsx"'
        return response


    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['investigations'] = Investigation.objects.filter(is_active=True).order_by('name')
        return context

@csrf_exempt
@login_required
def api_parameter_save(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            param_id = data.get('id')
            name = data.get('name', '').strip()
            code = data.get('code', '').strip()
            unit = data.get('unit', '').strip()
            data_type = data.get('data_type', 'Numeric').strip()
            is_active = data.get('is_active', True)
            
            if not name or not code:
                return JsonResponse({'status': 'error', 'message': 'Name and Code are required.'})
                
            duplicate = Parameter.objects.filter(code__iexact=code)
            if param_id:
                duplicate = duplicate.exclude(id=param_id)
            if duplicate.exists():
                return JsonResponse({'status': 'error', 'message': f'Parameter with code {code} already exists.'})

            if param_id:
                param = Parameter.objects.get(id=param_id)
                param.name = name
                param.code = code
                param.default_unit = unit
                param.data_type = data_type
                param.is_active = is_active
                param.save()
                msg = "Parameter updated successfully."
            else:
                param = Parameter.objects.create(
                    name=name,
                    code=code,
                    default_unit=unit,
                    data_type=data_type,
                    is_active=is_active
                )
                msg = "Parameter saved successfully."
                
            return JsonResponse({
                'status': 'success',
                'message': msg,
                'data': {
                    'id': param.id,
                    'name': param.name,
                    'code': param.code,
                    'unit': param.default_unit,
                    'data_type': param.data_type,
                    'is_active': param.is_active
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method'})

class ParameterCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_master.parameter.create'
    model = InvestigationParameter
    form_class = ParameterForm
    template_name = 'lab/master/parameter_form.html'
    success_url = reverse_lazy('lab:parameter_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Parameter created successfully!")
        return super().form_valid(form)

class ParameterUpdateView(LoginRequiredMixin, GranularPermissionRequiredMixin, UpdateView):
    permission_required = 'lab_master.parameter.update'
    model = InvestigationParameter
    form_class = ParameterForm
    template_name = 'lab/master/parameter_form.html'
    success_url = reverse_lazy('lab:parameter_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['reference_ranges'] = self.object.reference_ranges.all()
        return context

    def form_valid(self, form):
        messages.success(self.request, "Parameter updated successfully!")
        return super().form_valid(form)

class AgeGroupListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.age_group.view'
    model = AgeGroup
    template_name = 'lab/master/agegroup_list.html'
    context_object_name = 'age_groups'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(name__icontains=search))
        return qs

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.age_group.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:agegroup_list')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['age_group_code', 'age_group_name', 'minimum_age', 'maximum_age', 'age_unit', 'gender', 'pregnancy_applicable', 'display_order', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.code or '',
                obj.label or '',
                obj.min_age_value if obj.min_age_value is not None else '',
                obj.max_age_value if obj.max_age_value is not None else '',
                obj.min_age_unit or 'Years',
                obj.gender or 'All',
                'Yes' if obj.pregnancy_applicable else 'No',
                obj.sort_order if obj.sort_order is not None else '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="agegroup_export.xlsx"'
        return response

class AgeGroupCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_master.age_group.create'
    model = AgeGroup
    form_class = AgeGroupForm
    template_name = 'lab/master/agegroup_form.html'
    success_url = reverse_lazy('lab:agegroup_list')
    
    def form_valid(self, form):
        messages.success(self.request, "Age Group created successfully!")
        return super().form_valid(form)

class AgeGroupUpdateView(LoginRequiredMixin, GranularPermissionRequiredMixin, UpdateView):
    permission_required = 'lab_master.age_group.update'
    model = AgeGroup
    form_class = AgeGroupForm
    template_name = 'lab/master/agegroup_form.html'
    success_url = reverse_lazy('lab:agegroup_list')

    def form_valid(self, form):
        messages.success(self.request, "Age Group updated successfully!")
        return super().form_valid(form)


# --- Reference Range Grid ---
class ReferenceRangeGridView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.reference_range.view'
    template_name = 'lab/master/reference_range_grid.html'
    model = ParameterReferenceRange
    context_object_name = 'reference_ranges'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset().select_related('investigation_parameter', 'investigation_parameter__investigation', 'age_group').order_by('-id')
        search = self.request.GET.get('search', '').strip()
        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(investigation_parameter__name__icontains=search) | 
                Q(investigation_parameter__code__icontains=search) |
                Q(reference_text__icontains=search) |
                Q(age_group__label__icontains=search)
            )
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['parameters'] = InvestigationParameter.objects.select_related('investigation').filter(is_active=True).order_by('name')
        context['investigations'] = Investigation.objects.filter(is_active=True)
        context['age_groups'] = AgeGroup.objects.filter(is_active=True).order_by('sort_order', 'label')
        context['diagnoses'] = Diagnosis.objects.filter(is_active=True)
        return context

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.reference_range.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:reference_range_grid')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io
        from .models import ParameterReferenceRange

        queryset = self.get_queryset()
        
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['parameter_code', 'age_group_code', 'gender', 'pregnancy', 'range_type', 'min_value', 'max_value', 'reference_text', 'unit', 'method', 'remarks', 'active'])
        
        for obj in queryset.iterator(chunk_size=1000):
            ws.append([
                obj.investigation_parameter.code if obj.investigation_parameter else '',
                obj.age_group.code if obj.age_group else '',
                obj.gender or 'All',
                'Yes' if obj.pregnancy else 'No',
                obj.range_type or 'Numeric',
                obj.min_value if obj.min_value is not None else '',
                obj.max_value if obj.max_value is not None else '',
                obj.reference_text or '',
                obj.unit or '',
                obj.method or '',
                obj.remarks or '',
                'Yes' if obj.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="reference_range_export.xlsx"'
        return response

class InvestigationParameterMappingView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_master.mapping.view'
    template_name = 'lab/master/investigation_parameter_mapping.html'
    model = InvestigationParameter
    context_object_name = 'mappings'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.GET.get('search', '').strip()
        if search:
            qs = qs.filter(models.Q(investigation__name__icontains=search) | models.Q(parameter__name__icontains=search))
        return qs

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            if not request.user.has_perm_code('lab_master.mapping.export'):
                from django.contrib import messages
                from django.shortcuts import redirect
                messages.error(request, "Access Denied: You do not have permission to export.")
                return redirect('lab:investigation_parameter_mapping')
            return self.export_excel(request)
        return super().get(request, *args, **kwargs)

    def export_excel(self, request):
        from openpyxl import Workbook
        from django.http import HttpResponse
        import io

        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Template')
        ws.append(['investigation', 'parameter', 'active'])
        
        for mapping in self.get_queryset().select_related('investigation', 'parameter').iterator(chunk_size=1000):
            inv_name = mapping.investigation.name if mapping.investigation else ''
            param_name = mapping.parameter.name if mapping.parameter else ''
            ws.append([
                inv_name,
                param_name,
                'Yes' if mapping.is_active else 'No'
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="investigation_parameter_mapping_export.xlsx"'
        return response
    
    def get_queryset(self):
        return super().get_queryset().select_related('investigation', 'parameter')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['investigations'] = Investigation.objects.filter(is_active=True).order_by('name')
        context['parameters'] = Parameter.objects.filter(is_active=True).order_by('name')
        return context

from apps.core.mixins import granular_permission_required

@granular_permission_required('lab_master.mapping.view')
def get_all_parameters(request):
    params = InvestigationParameter.objects.filter(is_active=True).values('code', 'name').distinct().order_by('name')
    param_list = [{'id': p['code'], 'code': p['code'], 'name': p['name']} for p in params if p['code']]
    return JsonResponse({'parameters': param_list})

@granular_permission_required('lab_master.mapping.update')
def save_investigation_parameters(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            inv_id = data.get('investigation_id')
            param_codes = data.get('parameter_ids', [])
            
            if not inv_id:
                return JsonResponse({'status': 'error', 'message': 'Investigation ID is required'})
                
            existing_mappings = InvestigationParameter.objects.filter(investigation_id=inv_id)
            existing_map_dict = {m.code: m for m in existing_mappings if m.code}
            
            target_codes = set(str(p) for p in param_codes)
            
            for code, mapping in existing_map_dict.items():
                if code not in target_codes and mapping.is_active:
                    mapping.is_active = False
                    mapping.save()
                    
            display_order = existing_mappings.count()
            for code in target_codes:
                if code in existing_map_dict:
                    mapping = existing_map_dict[code]
                    if not mapping.is_active:
                        mapping.is_active = True
                        mapping.save()
                else:
                    source = InvestigationParameter.objects.filter(code=code).first()
                    if source:
                        display_order += 1
                        InvestigationParameter.objects.create(
                            investigation_id=inv_id,
                            name=source.name,
                            code=source.code,
                            unit=source.unit,
                            result_type=source.result_type,
                            decimal_precision=source.decimal_precision,
                            display_order=display_order,
                            is_active=True
                        )
            
            return JsonResponse({'status': 'success'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'})

@csrf_exempt
@granular_permission_required('lab_master.parameter.create')
def api_add_parameter(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            name = data.get('name', '').strip()
            code = data.get('code', '').strip()
            unit = data.get('unit', '').strip()
            is_active = data.get('is_active', True)
            
            if not name:
                return JsonResponse({'status': 'error', 'message': 'Parameter Name is required.'})
                
            # Check for duplicate
            if Parameter.objects.filter(name__iexact=name).exists() or (code and Parameter.objects.filter(code__iexact=code).exists()):
                return JsonResponse({'status': 'error', 'message': 'A parameter with this name or code already exists.'})
                
            param = Parameter.objects.create(
                name=name,
                code=code,
                default_unit=unit,
                is_active=is_active
            )
            
            return JsonResponse({
                'status': 'success',
                'parameter': {
                    'id': param.id,
                    'name': param.name,
                    'code': param.code,
                    'unit': param.default_unit
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'})

@granular_permission_required('lab_master.reference_range.view')
def get_investigation_parameters(request):
    investigation_id = request.GET.get('investigation_id')
    if not investigation_id:
        return JsonResponse({'parameters': []})
    
    params = InvestigationParameter.objects.filter(investigation_id=investigation_id, is_active=True).select_related('parameter')
    data = [{'id': p.id, 'parameter_id': p.code, 'name': p.name or (p.parameter.name if p.parameter else ''), 'unit': p.unit or (p.parameter.default_unit if p.parameter else '')} for p in params]
    return JsonResponse({'parameters': data})

@granular_permission_required('lab_master.reference_range.view')
def get_reference_ranges(request):
    inv_param_id = request.GET.get('inv_param_id')
    if not inv_param_id:
        return JsonResponse({'ranges': []})
    
    ranges = ParameterReferenceRange.objects.filter(investigation_parameter_id=inv_param_id).select_related('age_group')
    
    data = []
    for r in ranges:
        data.append({
            'id': r.id,
            'age_group_id': r.age_group.id,
            'age_group_name': r.age_group.label,
            'gender': r.gender,
            'pregnancy': r.pregnancy,
            'min_value': str(r.min_value) if r.min_value is not None else '',
            'max_value': str(r.max_value) if r.max_value is not None else '',
            'unit': r.unit or '',
            'reference_text': r.reference_text or '',
            'is_active': r.is_active
        })
    return JsonResponse({'ranges': data})

@granular_permission_required('lab_master.reference_range.view')
def get_reference_ranges_bulk(request):
    investigation_id = request.GET.get('investigation_id')
    age_group_id = request.GET.get('age_group_id')
    gender = request.GET.get('gender')
    
    if not all([investigation_id, age_group_id, gender]):
        return JsonResponse({'ranges': []})
        
    ranges = ParameterReferenceRange.objects.filter(
        investigation_parameter__investigation_id=investigation_id,
        age_group_id=age_group_id,
        gender=gender
    )
    
    data = []
    for r in ranges:
        data.append({
            'id': r.id,
            'investigation_parameter_id': r.investigation_parameter_id,
            'age_group_id': r.age_group_id,
            'gender': r.gender,
            'pregnancy': r.pregnancy,
            'range_type': r.range_type,
            'min_value': str(r.min_value) if r.min_value is not None else '',
            'max_value': str(r.max_value) if r.max_value is not None else '',
            'unit': r.unit or '',
            'reference_text': r.reference_text or '',
            'method': r.method or '',
            'remarks': r.remarks or '',
            'is_active': r.is_active
        })
    return JsonResponse({'ranges': data})

@granular_permission_required('lab_master.reference_range.view')
def get_reference_ranges_all(request):
    inv_id = request.GET.get('inv_id')
    ag_id = request.GET.get('ag_id')
    gender = request.GET.get('gender')
    
    ranges = ParameterReferenceRange.objects.select_related(
        'investigation_parameter__investigation',
        'investigation_parameter',
        'age_group'
    ).all()
    
    if inv_id:
        ranges = ranges.filter(investigation_parameter__investigation_id=inv_id)
    if ag_id:
        ranges = ranges.filter(age_group_id=ag_id)
    if gender:
        ranges = ranges.filter(gender=gender)
        
    data = []
    for r in ranges:
        val = ''
        if r.range_type == 'Numeric':
            val = f"{r.min_value} - {r.max_value}"
        elif r.range_type == 'Text':
            val = r.reference_text
        else:
            val = '-'
            
        data.append({
            'id': r.id,
            'investigation': r.investigation_parameter.investigation.name,
            'investigation_id': r.investigation_parameter.investigation_id,
            'parameter': r.investigation_parameter.name,
            'age_group': r.age_group.label,
            'age_group_id': r.age_group_id,
            'gender': r.gender,
            'range': val,
            'unit': r.unit or '-',
            'method': r.method or '-',
            'status': r.is_active
        })
    return JsonResponse({'ranges': data})

@granular_permission_required('lab_master.reference_range.update')
def save_reference_range(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            range_id = data.get('id')
            inv_param_id = data.get('investigation_parameter_id')
            age_group_id = data.get('age_group_id')
            gender = data.get('gender', 'All')
            pregnancy = data.get('pregnancy', 'No')
            
            if not inv_param_id or not age_group_id:
                return JsonResponse({'status': 'error', 'message': 'Missing required fields'})
            
            # Validate Min <= Max
            min_val_raw = data.get('min_value')
            max_val_raw = data.get('max_value')
            min_val = Decimal(min_val_raw) if min_val_raw and str(min_val_raw).strip() != '' else None
            max_val = Decimal(max_val_raw) if max_val_raw and str(max_val_raw).strip() != '' else None
            
            if min_val is not None and max_val is not None:
                if min_val > max_val:
                    return JsonResponse({'status': 'error', 'message': 'Minimum value cannot be greater than maximum value.'})
            
            created = False
            if range_id:
                try:
                    range_obj = ParameterReferenceRange.objects.get(id=range_id)
                    range_obj.gender = gender
                    range_obj.pregnancy = pregnancy
                except ParameterReferenceRange.DoesNotExist:
                    return JsonResponse({'status': 'error', 'message': 'Record not found for update'})
            else:
                range_obj, created = ParameterReferenceRange.objects.get_or_create(
                    investigation_parameter_id=inv_param_id,
                    age_group_id=age_group_id,
                    gender=gender,
                    pregnancy=pregnancy
                )
            
            range_obj.range_type = data.get('range_type', 'Numeric')
            range_obj.min_value = min_val
            range_obj.max_value = max_val
            range_obj.reference_text = data.get('reference_text', '')
            range_obj.unit = data.get('unit', '')
            range_obj.method = data.get('method', '')
            range_obj.remarks = data.get('remarks', '')
            range_obj.is_active = data.get('is_active', True)
            
            range_obj.save()
            return JsonResponse({'status': 'success', 'id': range_obj.id, 'created': created})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'})

@granular_permission_required('lab_master.reference_range.create')
def save_multiple_reference_ranges(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            param_ids = data.get('parameter_ids', [])
            age_group_ids = data.get('age_group_ids', [])
            gender = data.get('gender', 'All')
            
            # Additional fields
            range_type = data.get('range_type', 'Numeric')
            min_val_raw = data.get('min_value')
            max_val_raw = data.get('max_value')
            reference_text = data.get('reference_text', '')
            unit = data.get('unit', '')
            method = data.get('method', '')
            remarks = data.get('remarks', '')
            
            update_existing = data.get('update_existing', False)
            
            if not param_ids or not age_group_ids:
                return JsonResponse({'status': 'error', 'message': 'Missing parameters or age groups'})
                
            min_val = Decimal(min_val_raw) if min_val_raw and str(min_val_raw).strip() != '' else None
            max_val = Decimal(max_val_raw) if max_val_raw and str(max_val_raw).strip() != '' else None
            
            if range_type == 'Numeric' and min_val is not None and max_val is not None:
                if min_val > max_val:
                    return JsonResponse({'status': 'error', 'message': 'Minimum value cannot be greater than maximum value.'})
                    
            created_count = 0
            updated_count = 0
            
            # Check for existing
            if not update_existing:
                for p_id in param_ids:
                    for ag_id in age_group_ids:
                        if ParameterReferenceRange.objects.filter(
                            investigation_parameter_id=p_id,
                            age_group_id=ag_id,
                            gender=gender
                        ).exists():
                            return JsonResponse({'status': 'exists', 'message': 'Reference range already exists. Update existing?'})

            for p_id in param_ids:
                for ag_id in age_group_ids:
                    range_obj, created = ParameterReferenceRange.objects.get_or_create(
                        investigation_parameter_id=p_id,
                        age_group_id=ag_id,
                        gender=gender
                    )
                    
                    if created:
                        created_count += 1
                    else:
                        updated_count += 1
                        
                    range_obj.range_type = range_type
                    range_obj.min_value = min_val
                    range_obj.max_value = max_val
                    range_obj.reference_text = reference_text
                    range_obj.unit = unit
                    range_obj.method = method
                    range_obj.remarks = remarks
                    range_obj.is_active = True
                    range_obj.save()
                    
            return JsonResponse({
                'status': 'success',
                'message': f"{created_count + updated_count} reference ranges saved successfully."
            })
            
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'})

@granular_permission_required('lab_master.reference_range.delete')
def delete_reference_range(request, pk):
    if request.method == 'POST':
        try:
            r = ParameterReferenceRange.objects.get(id=pk)
            r.delete()
            return JsonResponse({'status': 'success'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error'})

def api_add_department(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        if name:
            dept, created = LabDepartment.objects.get_or_create(name=name.strip())
            return JsonResponse({'status': 'success', 'id': dept.id, 'name': dept.name})
    return JsonResponse({'status': 'error'})

def api_add_sample_type(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        if name:
            st, created = SampleType.objects.get_or_create(name=name.strip())
            return JsonResponse({'status': 'success', 'id': st.id, 'name': st.name})
    return JsonResponse({'status': 'error'})

class LegacyMappingView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'lab_master.legacy_mapping.view'
    template_name = 'lab/master/legacy_mapping.html'

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'excel':
            from apps.lab.universal_importer import export_universal_master
            from django.http import HttpResponse
            output = export_universal_master()
            response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="universal_master_export.xlsx"'
            return response
        
        if request.GET.get('template') == 'true':
            from apps.lab.universal_importer import generate_blank_template
            from django.http import HttpResponse
            output = generate_blank_template()
            response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="universal_master_template.xlsx"'
            return response

        return super().get(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        from apps.lab.universal_importer import validate_import, commit_import, UniversalImportError
        from django.conf import settings
        from django.core.files.storage import FileSystemStorage
        import logging
        import os
        import time
        import uuid
        
        logger = logging.getLogger('apps.lab.views')
        action = request.POST.get('action')
        
        # Use a secure temporary folder
        temp_dir = os.path.join(settings.MEDIA_ROOT, 'temp_imports')
        fs = FileSystemStorage(location=temp_dir)
        
        # Helper to clean up temp files older than 2 hours
        def cleanup_old_files():
            try:
                if not os.path.exists(temp_dir):
                    return
                current_time = time.time()
                for filename in os.listdir(temp_dir):
                    file_path = os.path.join(temp_dir, filename)
                    if os.path.isfile(file_path):
                        if current_time - os.path.getmtime(file_path) > 7200:
                            os.remove(file_path)
            except Exception as e:
                logger.error(f"Failed to cleanup temp imports: {e}")

        if action == 'validate':
            file_obj = request.FILES.get('file')
            if not file_obj:
                return JsonResponse({'status': 'error', 'message': 'No file uploaded'})
            try:
                cleanup_old_files()
                
                # Perform read-only validation
                res = validate_import(file_obj)
                
                # Store the uploaded file securely with a UUID
                validation_id = str(uuid.uuid4())
                safe_filename = f"{validation_id}.xlsx"
                
                # Rewind file pointer because validate_import already read it
                file_obj.seek(0)
                fs.save(safe_filename, file_obj)
                
                res['validation_id'] = validation_id
                
                return JsonResponse({'status': 'success', 'validation': res})
            except Exception as e:
                logger.exception("Error during universal import validation")
                return JsonResponse({
                    'status': 'error', 
                    'message': f'Validation service encountered an unexpected error: {str(e)}. Please check the server logs.'
                })
                
        elif action == 'import':
            validation_id = request.POST.get('validation_id')
            if not validation_id:
                return JsonResponse({'status': 'error', 'message': 'No validation token provided. Please validate again.'})
                
            safe_filename = f"{validation_id}.xlsx"
            if not fs.exists(safe_filename):
                return JsonResponse({'status': 'error', 'message': 'Import session expired or the uploaded workbook is no longer available. Please upload the workbook again.'})
                
            file_path = fs.path(safe_filename)
            try:
                with open(file_path, 'rb') as f:
                    counts = commit_import(f)
                
                # Cleanup after successful import
                try:
                    fs.delete(safe_filename)
                except Exception as cleanup_err:
                    logger.warning(f"Failed to delete temp file {safe_filename}: {cleanup_err}")
                    
                return JsonResponse({'status': 'success', 'counts': counts})
            except UniversalImportError as uie:
                logger.error(f"Universal import error: {uie.to_dict()}")
                try:
                    if fs.exists(safe_filename):
                        fs.delete(safe_filename)
                except Exception:
                    pass
                return JsonResponse({
                    'status': 'error',
                    'message': uie.message,
                    'error_details': uie.to_dict()
                })
            except Exception as e:
                logger.exception("Error during universal import commit")
                # Attempt to cleanup on failure
                try:
                    if fs.exists(safe_filename):
                        fs.delete(safe_filename)
                except Exception:
                    pass
                return JsonResponse({
                    'status': 'error', 
                    'message': f"Import failed — no partial changes were committed. {str(e)}",
                    'error_details': {
                        'entity': 'Master Entity',
                        'entity_name': '',
                        'row_number': None,
                        'reason': str(e),
                        'action': 'Please verify workbook data and re-try.',
                        'message': str(e)
                    }
                })

        return JsonResponse({'status': 'error', 'message': 'Invalid action'})


@granular_permission_required('lab_master.legacy_mapping.update')
def process_legacy_mapping(request):
    """
    API endpoint to handle legacy mapping operations from both tabs.
    Expected JSON payload: { "tab": "diagnosis"|"investigation", "action": "...", "ids": [1,2,3], ...other fields }
    """
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Invalid method'})
        
    try:
        data = json.loads(request.body)
        tab = data.get('tab')
        action = data.get('action')
        item_ids = data.get('ids', [])
        
        if not item_ids:
            return JsonResponse({'status': 'error', 'message': 'No items selected'})

        if tab == 'diagnosis':
            staging_qs = StagingDiagnosis.objects.filter(id__in=item_ids)
            if action == 'chief_complaint':
                for st in staging_qs:
                    # Save as Chief Complaint
                    cc, created = ChiefComplaint.objects.get_or_create(
                        legacy_code=st.legacy_id,
                        defaults={'name': st.legacy_text}
                    )
                    st.migrated = True
                    st.migrated_to_type = 'ChiefComplaint'
                    st.migrated_to_id = cc.id
                    st.save()
                    
            elif action == 'true_diagnosis':
                # Single item mapping
                st = staging_qs.first()
                icd_code = data.get('icd_code')
                icd_title = data.get('icd_title')
                if st and icd_code:
                    diag, created = Diagnosis.objects.get_or_create(
                        code=icd_code,
                        defaults={
                            'name': icd_title or st.legacy_text,
                            'icd11_title': icd_title,
                            'legacy_code': st.legacy_id
                        }
                    )
                    st.migrated = True
                    st.migrated_to_type = 'Diagnosis'
                    st.migrated_to_id = diag.id
                    st.save()
                    
            elif action == 'skip':
                staging_qs.update(migrated=True, migrated_to_type='Skipped')

        elif tab == 'investigation':
            staging_qs = StagingInvestigation.objects.filter(id__in=item_ids)
            
            if action == 'hospital_service':
                category = data.get('category', HospitalService.CategoryChoices.OTHER)
                for st in staging_qs:
                    hs, created = HospitalService.objects.get_or_create(
                        code=st.legacy_code,
                        defaults={'name': st.legacy_text, 'category': category}
                    )
                    st.migrated = True
                    st.migrated_to_type = 'HospitalService'
                    st.migrated_to_id = hs.id
                    st.save()
                    
            elif action == 'true_investigation':
                # Single item mapping
                st = staging_qs.first()
                name = data.get('name')
                short_name = data.get('short_name')
                dept_id = data.get('department_id')
                sample_type_id = data.get('sample_type_id')
                is_panel = data.get('is_panel', False)
                
                if st and name:
                    inv, created = Investigation.objects.get_or_create(
                        code=f"INV-{st.legacy_code}", # ensure unique code format or just use legacy code
                        defaults={
                            'name': name,
                            'short_name': short_name,
                            'department_id': dept_id if dept_id else None,
                            'sample_type_id': sample_type_id if sample_type_id else None,
                            'is_panel': is_panel,
                            'legacy_code': st.legacy_code
                        }
                    )
                    st.migrated = True
                    st.migrated_to_type = 'Investigation'
                    st.migrated_to_id = inv.id
                    st.save()
            
            elif action == 'skip':
                staging_qs.update(migrated=True, migrated_to_type='Skipped')

        return JsonResponse({'status': 'success'})
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'status': 'error', 'message': str(e)})

# --- Transaction Screens ---
class OrderEntryView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'lab_orders.create_lab_order.view'
    template_name = 'lab/orders/order_entry.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        patient_id = self.request.GET.get('patient_id')
        if patient_id:
            context['patient'] = Patient.objects.filter(id=patient_id).first()
        context['diagnoses'] = Diagnosis.objects.filter(is_active=True)
        context['diagnosis_investigation_map_json'] = json.dumps({})
        context['all_investigations'] = Investigation.objects.filter(is_active=True)
        return context

    def post(self, request, *args, **kwargs):
        patient_id = request.POST.get('patient_id')
        diagnosis_id = request.POST.get('diagnosis_id')
        investigation_ids = request.POST.getlist('investigations')

        if not patient_id or not investigation_ids:
            messages.error(request, "Patient and at least one Investigation are required.")
            return redirect('lab:order_entry')

        patient = get_object_or_404(Patient, id=patient_id)
        diagnosis = Diagnosis.objects.filter(id=diagnosis_id).first() if diagnosis_id else None
        
        for inv_id in investigation_ids:
            inv = get_object_or_404(Investigation, id=inv_id)
            PatientInvestigationOrder.objects.create(
                patient=patient,
                diagnosis=diagnosis,
                investigation=inv,
                ordered_by=request.user,
                status=PatientInvestigationOrder.StatusChoices.PENDING
            )
            
        messages.success(request, f"Successfully created {len(investigation_ids)} order(s) for patient {patient.name}.")
        return redirect('lab:order_entry')

class ResultEntryListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_orders.lab_order_list.view'
    model = PatientInvestigationOrder
    template_name = 'lab/orders/result_entry_list.html'
    context_object_name = 'orders'

    def get_queryset(self):
        return PatientInvestigationOrder.objects.filter(status=PatientInvestigationOrder.StatusChoices.PENDING).select_related('patient', 'investigation')

class ResultEntryDetailView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'lab_orders.lab_order_list.view'
    template_name = 'lab/orders/result_entry.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        order = get_object_or_404(PatientInvestigationOrder, id=kwargs.get('pk'))
        context['order'] = order
        
        patient = order.patient
        patient_gender = patient.gender or 'All'
        patient_age_days = (patient.age_years or 0) * 365 + (patient.age_months or 0) * 30 + (patient.age_days or 0)
        
        parameters = InvestigationParameter.objects.filter(investigation=order.investigation, is_active=True).select_related('parameter')
        from .models import ParameterReferenceRange
        
        # Fetch reference range for display dynamically
        for p in parameters:
            ranges = list(ParameterReferenceRange.objects.filter(investigation_parameter=p, is_active=True).select_related('age_group'))
            matched_ref = None
            
            # Try to find a precise match
            for ref in ranges:
                if ref.gender != 'All' and ref.gender != patient_gender:
                    continue
                
                if ref.age_group:
                    ag = ref.age_group
                    min_days, max_days = 0, 999999
                    
                    if ag.min_age_unit == 'Years': min_days = (ag.min_age_value or 0) * 365
                    elif ag.min_age_unit == 'Months': min_days = (ag.min_age_value or 0) * 30
                    elif ag.min_age_unit == 'Weeks': min_days = (ag.min_age_value or 0) * 7
                    elif ag.min_age_unit == 'Days': min_days = (ag.min_age_value or 0)
                    
                    if ag.max_age_unit == 'Years': max_days = (ag.max_age_value or 0) * 365
                    elif ag.max_age_unit == 'Months': max_days = (ag.max_age_value or 0) * 30
                    elif ag.max_age_unit == 'Weeks': max_days = (ag.max_age_value or 0) * 7
                    elif ag.max_age_unit == 'Days': max_days = (ag.max_age_value or 0)
                    
                    if not (min_days <= patient_age_days <= max_days):
                        continue
                
                # Special case for pregnancy (if the patient has it or if we don't know, we might just match the first we see)
                matched_ref = ref
                break
                
            if not matched_ref:
                # Fallback to general range (age_group=None) if exact match fails
                for ref in ranges:
                    if not ref.age_group and (ref.gender == 'All' or ref.gender == patient_gender):
                        matched_ref = ref
                        break
                
            if matched_ref:
                if matched_ref.reference_text:
                    p.display_ref_range = matched_ref.reference_text
                else:
                    p.display_ref_range = f"{matched_ref.min_value} - {matched_ref.max_value}"
                    if matched_ref.unit:
                        p.display_ref_range += f" {matched_ref.unit}"
            else:
                p.display_ref_range = ''
                
        context['parameters'] = parameters
        return context

    def post(self, request, *args, **kwargs):
        order = get_object_or_404(PatientInvestigationOrder, id=kwargs.get('pk'))
        params = InvestigationParameter.objects.filter(investigation=order.investigation, is_active=True)
        
        # Simplified resolve logic: just save for now
        from .services import resolveReferenceRange
        
        for p in params:
            val = request.POST.get(f'param_{p.id}')
            if val:
                # We would normally calculate patient age here, but just mock it
                age_group_id = None # need real logic for patient age
                
                # Mock resolving logic
                # result_range = resolveReferenceRange(age_group_id, order.diagnosis_id, p.id, order.patient.gender)
                
                PatientInvestigationResult.objects.update_or_create(
                    order=order,
                    investigation_parameter=p,
                    defaults={
                        'result_value': val,
                        'flag': PatientInvestigationResult.FlagChoices.NORMAL, # would use resolved logic
                        'entered_by': request.user
                    }
                )
                
        order.status = PatientInvestigationOrder.StatusChoices.COMPLETED
        order.save()
        messages.success(request, "Results saved successfully.")
        return redirect('lab:result_entry_list')


# --- Counts APIs ---
from django.http import JsonResponse
from .models import InvestigationParameter

def api_diagnosis_count(request):
    total = Diagnosis.objects.count()
    active = Diagnosis.objects.filter(is_active=True).count()
    return JsonResponse({
        'total': total,
        'active': active,
        'inactive': total - active
    })

# --- Service Request (Doctor Window) ---
from django.db import transaction, models
from django.db.models import Q
from apps.users.models import User

class ServiceRequestListView(LoginRequiredMixin, GranularPermissionRequiredMixin, ListView):
    permission_required = 'lab_orders.service_request.view'
    model = ServiceRequest
    template_name = 'lab/orders/service_request_list.html'
    context_object_name = 'service_requests'

    def get_queryset(self):
        return super().get_queryset().select_related('patient', 'consultant', 'department')

class ServiceRequestCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, CreateView):
    permission_required = 'lab_orders.service_request.create'
    model = ServiceRequest
    form_class = ServiceRequestForm
    template_name = 'lab/orders/service_request_form.html'
    success_url = reverse_lazy('lab:service_request_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.patients.models import Department
        context['departments'] = Department.objects.filter(is_active=True)
        # Assuming we can get consultants from User model (is_staff or something), filtering all for now
        context['consultants'] = User.objects.all()
        return context

    def form_valid(self, form):
        # We handle this mainly via AJAX, but if form submits normally:
        return super().form_valid(form)

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        # We will handle AJAX POST here
        if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.content_type == 'application/json':
            try:
                data = json.loads(request.body)
                patient_id = data.get('patient_id')
                consultant_id = data.get('consultant_id')
                consultant_name = data.get('consultant_name')
                department_id = data.get('department_id')
                visit_type = data.get('visit_type', 'OP')
                visit_no = data.get('visit_no')
                request_date = data.get('request_date')
                diagnoses_data = data.get('diagnoses', [])
                investigations_data = data.get('investigations', [])
                
                if not diagnoses_data:
                    return JsonResponse({'status': 'error', 'message': 'At least one diagnosis is required to save investigations.'})
                
                from apps.patients.models import Department
                patient = Patient.objects.get(id=patient_id)
                
                # Department Handling
                department = None
                if department_id:
                    department = Department.objects.filter(id=department_id).first()
                if not department and patient.department:
                    department = Department.objects.filter(name__iexact=patient.department).first()
                    
                consultant = User.objects.filter(id=consultant_id).first() if consultant_id else None

                sr = ServiceRequest.objects.create(
                    patient=patient,
                    consultant=consultant,
                    consultant_name=consultant_name,
                    department=department,
                    visit_type=visit_type,
                    visit_no=visit_no,
                    request_date=request_date or timezone.now().date(),
                    status=ServiceRequest.StatusChoices.SAVED,
                    created_by=request.user
                )

                for idx, diag in enumerate(diagnoses_data):
                    item_type = diag.get('type')
                    item_id = diag.get('id')
                    if item_type == 'Diagnosis':
                        ServiceRequestDiagnosis.objects.create(
                            service_request=sr,
                            diagnosis_id=item_id,
                            sort_order=idx
                        )
                    elif item_type == 'ChiefComplaint':
                        ServiceRequestDiagnosis.objects.create(
                            service_request=sr,
                            chief_complaint_id=item_id,
                            sort_order=idx
                        )

                for inv in investigations_data:
                    ServiceRequestInvestigation.objects.create(
                        service_request=sr,
                        investigation_id=inv.get('id'),
                        qty=inv.get('qty', 1),
                        source=inv.get('source', 'MANUAL'),
                        is_removed=inv.get('is_removed', False)
                    )

                return JsonResponse({'status': 'success', 'service_request_id': sr.id})
            except Exception as e:
                return JsonResponse({'status': 'error', 'message': str(e)})
        else:
            return super().post(request, *args, **kwargs)

@login_required
def api_get_patients_for_request(request):
    return api_doctor_window_patients(request)

@login_required
def api_search_diagnosis(request):
    q = request.GET.get('q', '').strip()
    diag_type = request.GET.get('type', 'ALL')
    
    if not q:
        diags = Diagnosis.objects.filter(is_active=True)
        chiefs = ChiefComplaint.objects.filter(is_active=True)
    else:
        diags = Diagnosis.objects.filter(
            Q(name__icontains=q) | Q(code__icontains=q) | Q(icd11_title__icontains=q) | Q(synonyms__icontains=q),
            is_active=True
        )
        chiefs = ChiefComplaint.objects.filter(name__icontains=q, is_active=True)
    
    results = []
    if diag_type in ['ALL', 'Diagnosis']:
        for d in diags:
            results.append({
                'id': d.id,
                'text': d.name,
                'code': d.code,
                'category': d.chapter or 'General',
                'type': 'Diagnosis'
            })
            
    if diag_type in ['ALL', 'ChiefComplaint']:
        for c in chiefs:
            results.append({
                'id': c.id,
                'text': c.name,
                'type': 'ChiefComplaint'
            })
            
    return JsonResponse({'results': results})

@login_required
@csrf_exempt
def api_suggest_investigations(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            diagnosis_ids = data.get('diagnosis_ids', [])
            patient_id = data.get('patient_id')
            
            if not diagnosis_ids:
                return JsonResponse({'status': 'success', 'investigations': []})
                
            from apps.lab.models import DiagnosisInvestigationMap
            mappings = DiagnosisInvestigationMap.objects.filter(diagnosis_id__in=diagnosis_ids, is_active=True).select_related('investigation', 'age_group')
            
            patient_age_days = 0
            patient_gender = 'All'
            if patient_id:
                from apps.patients.models import Patient
                try:
                    p = Patient.objects.get(id=patient_id)
                    patient_age_days = (getattr(p, 'age_years', 0) or 0) * 365 + (getattr(p, 'age_months', 0) or 0) * 30 + (getattr(p, 'age_days', 0) or 0)
                    patient_gender = getattr(p, 'gender', 'All')
                except Patient.DoesNotExist:
                    pass
                    
            results = {}
            for m in mappings:
                inv = m.investigation
                if m.age_group:
                    ag = m.age_group
                    # Check Gender
                    if ag.gender != 'All' and ag.gender != patient_gender:
                        continue
                    
                    # Check Age if patient exists
                    if patient_id:
                        multiplier = {'Days': 1, 'Months': 30, 'Years': 365, 'Weeks': 7}
                        min_days = (ag.min_age_value or 0) * multiplier.get(ag.min_age_unit, 365)
                        max_days = (ag.max_age_value or 150) * multiplier.get(ag.max_age_unit, 365)
                        if not (min_days <= patient_age_days <= max_days):
                            continue
                            
                results[inv.id] = {
                    'id': inv.id,
                    'name': inv.name,
                    'code': getattr(inv, 'code', ''),
                    'sample': getattr(inv.sample_type, 'name', 'Whole Blood') if hasattr(inv, 'sample_type') and inv.sample_type else 'Whole Blood',
                    'available_count': 0  # To be populated below
                }
                
            # Fetch global result pool availability
            from apps.lab.services.result_allocation_service import ResultAllocationService
            inv_ids = list(results.keys())
            pool_stats = ResultAllocationService.get_pool_availability(inv_ids)
            for inv_id in inv_ids:
                if inv_id in pool_stats:
                    results[inv_id]['available_count'] = pool_stats[inv_id]['available']
                    
            return JsonResponse({'status': 'success', 'investigations': list(results.values())})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
            
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'})


def api_investigation_count(request):
    total = Investigation.objects.count()
    active = Investigation.objects.filter(is_active=True).count()
    return JsonResponse({
        'total': total,
        'active': active,
        'inactive': total - active
    })

def api_parameter_count(request):
    total = InvestigationParameter.objects.count()
    active = InvestigationParameter.objects.filter(is_active=True).count()
    return JsonResponse({
        'total': total,
        'active': active,
        'inactive': total - active
    })

def api_search_investigation(request):
    query = request.GET.get('q', '').strip()
    diagnosis_ids_str = request.GET.get('diagnosis_ids', '')
    patient_id = request.GET.get('patient_id')
        
    investigations = Investigation.objects.filter(is_active=True)
    if query:
        investigations = investigations.filter(
            Q(name__icontains=query) | Q(code__icontains=query) | Q(short_name__icontains=query)
        )
        

    results = []
    inv_ids = [inv.id for inv in investigations]
    
    # Fetch global result pool availability
    from apps.lab.services.result_allocation_service import ResultAllocationService
    pool_stats = ResultAllocationService.get_pool_availability(inv_ids)
    
    for inv in investigations:
        avail = pool_stats.get(inv.id, {}).get('available', 0)
        results.append({
            'id': inv.id,
            'text': inv.name,
            'code': inv.code,
            'sample': inv.sample_type.name if inv.sample_type else 'Whole Blood',
            'available_count': avail
        })
        
    return JsonResponse({'results': results})

# --- Doctor Window ---

class DoctorWindowView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    template_name = 'lab/orders/doctor_window.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.patients.models import Department
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        return context

@login_required
def api_doctor_window_patients(request):
    department_id = request.GET.get('department_id')
    patient_type = request.GET.get('patient_type')
    visit_type = request.GET.get('visit_type')  # OP / IP / '' (All)
    date_str = request.GET.get('date')
    
    # We will combine Patient (Visit 1) and PatientVisit (Visit > 1)
    # 1. Query Patient (Visit 1)
    patients_qs = Patient.objects.all().order_by('-id')
    if date_str:
        patients_qs = patients_qs.filter(registration_date=date_str)
    if department_id:
        patients_qs = patients_qs.filter(department_obj_id=department_id)
    if patient_type and patient_type != 'A':
        patients_qs = patients_qs.filter(patient_type=patient_type)
    if visit_type:
        patients_qs = patients_qs.filter(visit_through=visit_type)
        
    # 2. Query PatientVisit (Visit > 1)
    visits_qs = PatientVisit.objects.select_related('patient', 'department_obj').all().order_by('-id')
    if date_str:
        visits_qs = visits_qs.filter(visit_date__date=date_str)
    if department_id:
        visits_qs = visits_qs.filter(department_obj_id=department_id)
    if patient_type and patient_type != 'A':
        visits_qs = visits_qs.filter(patient__patient_type=patient_type)
    if visit_type:
        visits_qs = visits_qs.filter(visit_type=visit_type)
        
    data = []
    
    # Process Patient (Visit 1)
    for p in patients_qs[:100]:
        # Check if Visit 1 record exists in PatientVisit for diagnosis
        visit1 = PatientVisit.objects.filter(patient=p, visit_no=1).first()
        diagnosis_texts = []
        if visit1:
            for d in visit1.diagnoses.all():
                if d.diagnosis: diagnosis_texts.append(d.diagnosis.name)
                elif d.chief_complaint: diagnosis_texts.append(d.chief_complaint.name)
                
        req_count = ServiceRequest.objects.filter(patient=p, request_date=p.registration_date).count()
        data.append({
            'visit_id': f"patient_{p.id}",
            'patient_id': p.patient_id,
            'patient_pk': p.id,
            'patient_name': p.name,
            'name': p.name,
            'gender': p.gender,
            'age': p.age_years,
            'department': p.department_obj.name if p.department_obj else (p.department or '-'),
            'dept_code': p.department_obj.code if p.department_obj else (p.department[:4].upper() if p.department else '-'),
            'visit_type': p.visit_through,
            'patient_type': p.patient_type,
            'visit_no': 1,
            'consultant': p.unit_doctor or '-',
            'diagnosis': ", ".join(diagnosis_texts) if diagnosis_texts else "Select",
            'request_count': req_count
        })
        
    # Process PatientVisit (Visit > 1)
    for v in visits_qs[:100]:
        if v.visit_no == 1:
            continue # Already handled in Patient queries typically
            
        diagnosis_texts = []
        for d in v.diagnoses.all():
            if d.diagnosis: diagnosis_texts.append(d.diagnosis.name)
            elif d.chief_complaint: diagnosis_texts.append(d.chief_complaint.name)
            
        req_count = ServiceRequest.objects.filter(patient=v.patient, request_date=v.visit_date.date()).count()
        data.append({
            'visit_id': v.id,
            'patient_id': v.patient.patient_id,
            'patient_pk': v.patient.id,
            'patient_name': v.patient.name,
            'name': v.patient.name,
            'gender': v.patient.gender,
            'age': v.patient.age_years,
            'department': v.department_obj.name if v.department_obj else (v.department or '-'),
            'dept_code': v.department_obj.code if v.department_obj else (v.department[:4].upper() if v.department else '-'),
            'visit_type': v.visit_type,
            'patient_type': v.patient.patient_type,
            'visit_no': v.visit_no,
            'consultant': v.unit_doctor or '-',
            'diagnosis': ", ".join(diagnosis_texts) if diagnosis_texts else "Select",
            'request_count': req_count
        })

    return JsonResponse({'patients': data[:150]})

@login_required
def api_doctor_window_get_diagnosis(request):
    visit_id = request.GET.get('visit_id')
    diagnoses = []
    
    if str(visit_id).startswith('patient_'):
        patient_id = visit_id.split('_')[1]
        visit = PatientVisit.objects.filter(patient_id=patient_id, visit_no=1).first()
    else:
        visit = get_object_or_404(PatientVisit, id=visit_id)
        
    if visit:
        for d in visit.diagnoses.all():
            if d.diagnosis:
                diagnoses.append({'id': d.diagnosis.id, 'text': d.diagnosis.name, 'type': 'Diagnosis'})
            elif d.chief_complaint:
                diagnoses.append({'id': d.chief_complaint.id, 'text': d.chief_complaint.name, 'type': 'Complaint'})
            
    return JsonResponse({'diagnoses': diagnoses})

@login_required
@require_POST
def api_doctor_window_save_diagnosis(request):
    data = json.loads(request.body)
    visit_id = data.get('visit_id')
    diagnoses = data.get('diagnoses', [])
    
    if str(visit_id).startswith('patient_'):
        patient_id = visit_id.split('_')[1]
        p = get_object_or_404(Patient, id=patient_id)
        
        from django.utils import timezone
        from datetime import datetime
        
        # Create or get Visit 1 for this Patient to attach diagnosis
        v_date = timezone.make_aware(datetime.combine(p.registration_date, datetime.min.time())) if p.registration_date else timezone.now()
        
        visit, _ = PatientVisit.objects.get_or_create(
            patient=p,
            visit_no=1,
            defaults={
                'visit_date': v_date,
                'department_obj': p.department_obj,
                'department': p.department,
                'visit_type': p.visit_through,
                'category': p.category
            }
        )
    else:
        visit = get_object_or_404(PatientVisit, id=visit_id)
    
    with transaction.atomic():
        visit.diagnoses.all().delete()
        for d in diagnoses:
            if d['type'] == 'Diagnosis':
                PatientVisitDiagnosis.objects.create(visit=visit, diagnosis_id=d['id'])
            elif d['type'] == 'Complaint':
                PatientVisitDiagnosis.objects.create(visit=visit, chief_complaint_id=d['id'])
                
    return JsonResponse({'status': 'success'})

# --- Work Orders ---
class WorkOrdersView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'lab_orders.work_orders.view'
    template_name = 'lab/orders/work_orders.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.patients.models import Department
        from apps.lab.models import SampleType, ServiceRequest
        from django.contrib.auth import get_user_model
        User = get_user_model()

        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        context['sample_types'] = SampleType.objects.filter(is_active=True).order_by('name')

        doctor_names = set(User.objects.filter(is_active=True).values_list('username', flat=True))
        for name in ServiceRequest.objects.exclude(consultant_name__isnull=True).exclude(consultant_name='').values_list('consultant_name', flat=True).distinct():
            if name:
                doctor_names.add(name)
        context['doctors'] = sorted(list(doctor_names))
        return context

@login_required
def api_work_orders_list(request):
    from apps.lab.models import ServiceRequest, ServiceRequestInvestigation
    from django.db.models import Exists, OuterRef, Case, When, Value, CharField, Q
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger

    # Pagination parameters
    try:
        per_page = int(request.GET.get('per_page', 100))
        if per_page not in [100, 500, 1000]:
            per_page = 100
    except (ValueError, TypeError):
        per_page = 100

    try:
        page_num = int(request.GET.get('page', 1))
        if page_num < 1:
            page_num = 1
    except (ValueError, TypeError):
        page_num = 1

    # Filter parameters
    from_date = request.GET.get('from_date', '').strip()
    to_date = request.GET.get('to_date', '').strip()
    department_id = request.GET.get('department_id') or request.GET.get('department')
    patient_type = request.GET.get('patient_type', '').strip()
    status = request.GET.get('status', '').strip()
    visit_type = request.GET.get('visit_type', '').strip()
    sample = request.GET.get('sample', '').strip()
    sample_id = request.GET.get('sample_id', '').strip()
    search_name = request.GET.get('search_name', '').strip()
    doctor = request.GET.get('doctor', '').strip()
    sort_by = request.GET.get('sort_by', 'patient_id').strip()
    sort_dir = request.GET.get('sort_dir', 'asc').strip().lower()
    patient_id = request.GET.get('patient_id', '').strip()
    search = request.GET.get('search', '').strip()

    # Base query with SQL status annotation following aggregate business rule
    sub_completed = Exists(ServiceRequestInvestigation.objects.filter(service_request=OuterRef('pk'), is_removed=False, status='COMPLETED'))
    sub_received = Exists(ServiceRequestInvestigation.objects.filter(service_request=OuterRef('pk'), is_removed=False, status='RECEIVED'))
    sub_rejected = Exists(ServiceRequestInvestigation.objects.filter(service_request=OuterRef('pk'), is_removed=False, status='REJECTED'))
    sub_pending = Exists(ServiceRequestInvestigation.objects.filter(service_request=OuterRef('pk'), is_removed=False, status='PENDING'))

    qs = ServiceRequest.objects.filter(investigations__is_removed=False).distinct().select_related(
        'patient', 'department', 'consultant'
    ).annotate(
        has_completed=sub_completed,
        has_received=sub_received,
        has_rejected=sub_rejected,
        has_pending=sub_pending,
        computed_status=Case(
            When(Q(has_completed=True) & Q(has_pending=False) & Q(has_received=False) & Q(has_rejected=False), then=Value('COMPLETED')),
            When(Q(has_received=True) | Q(has_completed=True), then=Value('RECEIVED')),
            When(has_rejected=True, then=Value('REJECTED')),
            default=Value('PENDING'),
            output_field=CharField()
        )
    )

    # Apply database-level filters
    if from_date:
        qs = qs.filter(request_date__gte=from_date)
    if to_date:
        qs = qs.filter(request_date__lte=to_date)
    if department_id and department_id != 'ALL':
        qs = qs.filter(department_id=department_id)
    if patient_type and patient_type != 'ALL':
        qs = qs.filter(patient__patient_type=patient_type)
    if status and status != 'ALL':
        qs = qs.filter(computed_status=status)
    if visit_type and visit_type != 'ALL':
        if visit_type == 'IP':
            qs = qs.filter(visit_type__in=['IP', 'INPATIENT'])
        else:
            qs = qs.filter(visit_type=visit_type)
    if sample and sample != 'ALL':
        if sample.isdigit():
            qs = qs.filter(investigations__investigation__sample_type_id=sample)
        else:
            qs = qs.filter(investigations__investigation__sample_type__name__iexact=sample)
    if sample_id:
        qs = qs.filter(sample_id__icontains=sample_id)
    if search_name:
        qs = qs.filter(patient__name__icontains=search_name)
    if doctor and doctor != 'ALL':
        if doctor.isdigit():
            qs = qs.filter(Q(consultant_id=doctor) | Q(consultant_name__icontains=doctor) | Q(patient__unit_doctor__icontains=doctor))
        else:
            qs = qs.filter(Q(consultant__username__icontains=doctor) | Q(consultant_name__icontains=doctor) | Q(patient__unit_doctor__icontains=doctor))
    if patient_id:
        qs = qs.filter(patient__patient_id__icontains=patient_id)
    if search:
        qs = qs.filter(
            Q(patient__patient_id__icontains=search) |
            Q(patient__name__icontains=search) |
            Q(sample_id__icontains=search) |
            Q(receipt_no__icontains=search)
        )

    # Sorting
    prefix = '-' if sort_dir == 'desc' else ''
    if sort_by == 'patient_id':
        qs = qs.order_by(f"{prefix}patient__patient_id", '-id')
    elif sort_by == 'patient_name':
        qs = qs.order_by(f"{prefix}patient__name", '-id')
    elif sort_by == 'voucher_no':
        qs = qs.order_by(f"{prefix}receipt_no", f"{prefix}id")
    else:
        qs = qs.order_by('-request_date', '-id')

    # Django database pagination
    paginator = Paginator(qs, per_page)
    try:
        page_obj = paginator.page(page_num)
    except (EmptyPage, PageNotAnInteger):
        page_obj = paginator.page(1)
        page_num = 1

    start_index = page_obj.start_index() if paginator.count > 0 else 0
    data = []
    for idx, sr in enumerate(page_obj.object_list):
        voucher_no = sr.receipt_no
        if not voucher_no:
            date_part = sr.request_date.strftime('%y%m%d') if sr.request_date else sr.created_at.strftime('%y%m%d')
            voucher_no = f"VCH{date_part}{sr.id:03d}"

        bill_time = sr.created_at.strftime('%I:%M:%S %p') if sr.created_at else '-'
        sample_dt = sr.request_date.strftime('%d/%b/%Y') if sr.request_date else '-'
        ipno = sr.patient.ipno if getattr(sr.patient, 'ipno', None) else '--'

        amount_val = getattr(sr, 'amount', None)
        if amount_val is not None:
            amount_str = f"{float(amount_val):.2f}"
        else:
            amount_str = '-'

        v_no = sr.visit_no if sr.visit_no else 1
        age = sr.patient.age_years if sr.patient.age_years is not None else '--'
        v_type = 'IP' if sr.visit_type in ['IP', 'INPATIENT'] else 'OP'

        data.append({
            's_no': start_index + idx,
            'id': sr.id,
            'patient_id': sr.patient.patient_id,
            'patient_name': sr.patient.name,
            'v_no': v_no,
            'age': age,
            'v_type': v_type,
            'sample_id': sr.sample_id or '-',
            'voucher_no': voucher_no,
            'amount': amount_str,
            'bill_time': bill_time,
            'sample_dt': sample_dt,
            'ipno': ipno,
            'status': getattr(sr, 'computed_status', 'PENDING')
        })

    return JsonResponse({
        'status': 'success',
        'work_orders': data,
        'pagination': {
            'page': page_obj.number,
            'num_pages': paginator.num_pages,
            'total_records': paginator.count,
            'per_page': per_page,
            'has_previous': page_obj.has_previous(),
            'has_next': page_obj.has_next()
        }
    })

class WorkOrderDetailView(LoginRequiredMixin, GranularPermissionRequiredMixin, DetailView):
    permission_required = 'lab_orders.work_orders.view'
    template_name = 'lab/orders/work_order_detail.html'
    context_object_name = 'order'
    
    def get_queryset(self):
        from apps.lab.models import ServiceRequest
        return ServiceRequest.objects.select_related('patient', 'department', 'consultant')
        
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.lab.models import ServiceRequestInvestigation
        from apps.lab.services.work_order_service import get_overall_sample_status
        context['investigations'] = self.object.investigations.filter(is_removed=False).select_related(
            'investigation', 'investigation__department', 'investigation__sample_type', 'received_by'
        )
        context['overall_status'] = get_overall_sample_status(self.object)
        return context

@login_required
@require_POST
def api_work_order_receive_multiple(request):
    try:
        data = json.loads(request.body)
        test_ids = data.get('test_ids', [])
        
        from apps.lab.models import ServiceRequestInvestigation
        tests = ServiceRequestInvestigation.objects.filter(id__in=test_ids, status=ServiceRequestInvestigation.StatusChoices.PENDING)
        
        for t in tests:
            t.status = ServiceRequestInvestigation.StatusChoices.RECEIVED
            t.received_date = timezone.now().date()
            t.received_time = timezone.now().time()
            t.received_by = request.user
            t.save()
            
        return JsonResponse({'status': 'success'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)})

@login_required
@require_POST
def api_work_order_receive(request, pk):
    from apps.lab.models import ServiceRequestInvestigation
    wo = get_object_or_404(ServiceRequestInvestigation, id=pk)
    if wo.status != ServiceRequestInvestigation.StatusChoices.PENDING:
        return JsonResponse({'status': 'error', 'message': 'Work order is not pending'})
    
    wo.status = ServiceRequestInvestigation.StatusChoices.RECEIVED
    wo.received_date = timezone.now().date()
    wo.received_time = timezone.now().time()
    wo.received_by = request.user
    wo.save()
    
    return JsonResponse({'status': 'success'})

class WorkOrderResultEntryView(LoginRequiredMixin, GranularPermissionRequiredMixin, DetailView):
    permission_required = 'lab_orders.work_orders.view'
    template_name = 'lab/orders/work_order_result_entry.html'
    context_object_name = 'wo'

    def get_queryset(self):
        from apps.lab.models import ServiceRequestInvestigation
        return ServiceRequestInvestigation.objects.select_related(
            'service_request',
            'service_request__patient',
            'service_request__department',
            'service_request__consultant',
            'investigation',
            'investigation__sample_type',
            'verified_by'
        )

    def get_object(self, queryset=None):
        pk = self.kwargs.get(self.pk_url_kwarg)
        from apps.lab.models import ServiceRequestInvestigation, ServiceRequest
        inv = ServiceRequestInvestigation.objects.filter(id=pk).select_related(
            'service_request', 'service_request__patient', 'service_request__department',
            'service_request__consultant', 'investigation', 'investigation__sample_type', 'verified_by'
        ).first()
        if inv:
            return inv
        sr = ServiceRequest.objects.filter(id=pk).first()
        if sr:
            first_inv = sr.investigations.filter(is_removed=False).order_by('id').first()
            if first_inv:
                return first_inv
        return super().get_object(queryset)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        wo = self.object
        service_request = wo.service_request
        patient = service_request.patient

        from apps.lab.services.work_order_service import (
            get_overall_sample_status,
            get_investigation_parameter_details
        )
        from django.contrib.auth import get_user_model
        User = get_user_model()

        # All investigations for this patient / service request
        all_invs = list(service_request.investigations.filter(is_removed=False).select_related(
            'investigation', 'investigation__sample_type'
        ).order_by('id'))

        # Prepare parameter details for the selected investigation
        details = get_investigation_parameter_details(wo, patient)

        # Diagnoses text
        diagnosis_texts = []
        for d in service_request.diagnoses.all():
            if d.diagnosis:
                diagnosis_texts.append(d.diagnosis.name)
            elif d.chief_complaint:
                diagnosis_texts.append(d.chief_complaint.name)

        completed_count = sum(1 for inv in all_invs if inv.status == 'COMPLETED')
        received_count = sum(1 for inv in all_invs if inv.status == 'RECEIVED')
        overall_status = get_overall_sample_status(service_request)

        # Doctor display name
        doc_name = '-'
        if service_request.consultant:
            doc_name = service_request.consultant.get_full_name() or service_request.consultant.username
        elif service_request.consultant_name:
            doc_name = service_request.consultant_name
        elif patient.unit_doctor:
            doc_name = patient.unit_doctor

        from apps.lab.services.work_order_service import (
            get_investigation_test_code,
            get_investigation_grpi,
            get_investigation_grp,
            get_or_create_doc_no
        )

        # Enrich all_invs with display attributes for the 9-column hospital laboratory table
        for idx, inv in enumerate(all_invs, 1):
            inv.s_no = idx
            # Req Status: R = Received or Completed, P = Pending
            inv.req_status = 'R' if inv.status in ('RECEIVED', 'COMPLETED') else 'P'
            inv.doctor_display = doc_name
            inv.doc_no_display = get_or_create_doc_no(inv)
            inv.test_code_display = get_investigation_test_code(inv.investigation)
            inv.grpi_display = get_investigation_grpi(inv.investigation)
            inv.grp_display = get_investigation_grp(inv.investigation)
            inv.is_current = (inv.id == wo.id)

        diagnosis_val = ", ".join(diagnosis_texts) if diagnosis_texts else ""
        clinical_remarks_val = service_request.clinical_remarks or ""

        context.update({
            'service_request': service_request,
            'patient': patient,
            'all_investigations': all_invs,
            'selected_inv': wo,
            'overall_status': overall_status,
            'completed_count': completed_count,
            'received_count': received_count,
            'can_print_selected': (wo.status == 'COMPLETED'),
            'can_print_all': (completed_count > 0),
            'diagnosis_text': diagnosis_val,
            'clinical_remarks': clinical_remarks_val,
            'parameters': details['parameters'],
            'param_data': details['param_data'],
            'matched_age_group': details['matched_age_group'],
            'verified_by_users': User.objects.filter(is_active=True).order_by('first_name', 'username'),
            'doctor_name': doc_name
        })
        return context

@login_required
@require_POST
def api_work_order_save_result(request, pk):
    from apps.lab.models import (
        ServiceRequestInvestigation,
        ServiceRequestResult,
        InvestigationParameter,
        ServiceRequest,
        ServiceRequestResultAudit,
        ServiceRequestDiagnosis,
        Diagnosis
    )
    from apps.lab.services.work_order_service import (
        get_overall_sample_status,
        get_next_received_investigation,
        get_investigation_parameter_details,
        get_or_create_doc_no
    )
    wo = get_object_or_404(ServiceRequestInvestigation, id=pk)

    try:
        data = json.loads(request.body)
        results = data.get('results', [])
        verified_by_id = data.get('verified_by_id')
        new_diagnosis = data.get('diagnosis')
        new_clinical_remarks = data.get('clinical_remarks')

        # If client requests updating metadata only (Diagnosis or Clinical Remarks)
        if data.get('only_metadata'):
            with transaction.atomic():
                if new_clinical_remarks is not None:
                    wo.service_request.clinical_remarks = str(new_clinical_remarks).strip()
                    wo.service_request.save(update_fields=['clinical_remarks'])
                if new_diagnosis is not None:
                    diag_str = str(new_diagnosis).strip()
                    if diag_str:
                        diag_obj = Diagnosis.objects.filter(name__iexact=diag_str).first()
                        if not diag_obj:
                            diag_obj = Diagnosis.objects.create(name=diag_str)
                        srd = ServiceRequestDiagnosis.objects.filter(service_request=wo.service_request).first()
                        if srd:
                            srd.diagnosis = diag_obj
                            srd.chief_complaint = None
                            srd.save(update_fields=['diagnosis', 'chief_complaint'])
                        else:
                            ServiceRequestDiagnosis.objects.create(
                                service_request=wo.service_request,
                                diagnosis=diag_obj
                            )
                    else:
                        ServiceRequestDiagnosis.objects.filter(service_request=wo.service_request).delete()
            return JsonResponse({'status': 'success', 'message': 'Details updated successfully.'})

        # Parameter validation: check that all actual editable parameters have entered values
        details = get_investigation_parameter_details(wo)
        entered_map = {
            str(res.get('parameter_id')): (str(res.get('value', '')).strip() if res.get('value') is not None else '')
            for res in results
        }

        missing_names = []
        for p in details['param_data']:
            if p['is_editable']:
                param_id_str = str(p['ip'].id)
                val = entered_map.get(param_id_str, '')
                if val == '':
                    missing_names.append(p['test_name'])

        if missing_names:
            msg = "Cannot save.\n\nPlease enter results for:\n" + "\n".join([f"• {name}" for name in missing_names])
            return JsonResponse({
                'status': 'error',
                'message': msg,
                'missing_parameters': missing_names
            }, status=400)

        with transaction.atomic():
            # Update diagnosis and clinical remarks if provided
            if new_clinical_remarks is not None:
                wo.service_request.clinical_remarks = str(new_clinical_remarks).strip()
                wo.service_request.save(update_fields=['clinical_remarks'])
            if new_diagnosis is not None:
                diag_str = str(new_diagnosis).strip()
                if diag_str:
                    diag_obj = Diagnosis.objects.filter(name__iexact=diag_str).first()
                    if not diag_obj:
                        diag_obj = Diagnosis.objects.create(name=diag_str)
                    srd = ServiceRequestDiagnosis.objects.filter(service_request=wo.service_request).first()
                    if srd:
                        srd.diagnosis = diag_obj
                        srd.chief_complaint = None
                        srd.save(update_fields=['diagnosis', 'chief_complaint'])
                    else:
                        ServiceRequestDiagnosis.objects.create(
                            service_request=wo.service_request,
                            diagnosis=diag_obj
                        )
                else:
                    ServiceRequestDiagnosis.objects.filter(service_request=wo.service_request).delete()

            for res in results:
                param_id = res.get('parameter_id')
                val = str(res.get('value', '')).strip()
                remarks = str(res.get('remarks', '')).strip()
                if val == '':
                    continue

                ip = InvestigationParameter.objects.get(id=param_id, investigation=wo.investigation)

                existing_res = ServiceRequestResult.objects.filter(
                    sr_investigation=wo,
                    investigation_parameter=ip
                ).first()

                if existing_res:
                    if existing_res.result_value != val:
                        # Record modification in audit log
                        ServiceRequestResultAudit.objects.create(
                            service_request=wo.service_request,
                            sr_investigation=wo,
                            investigation_parameter=ip,
                            old_value=existing_res.result_value,
                            new_value=val,
                            modified_by=request.user
                        )
                    existing_res.result_value = val
                    existing_res.remarks = remarks
                    existing_res.entered_by = request.user
                    existing_res.save()
                else:
                    ServiceRequestResult.objects.create(
                        sr_investigation=wo,
                        investigation_parameter=ip,
                        result_value=val,
                        remarks=remarks,
                        entered_by=request.user
                    )

            if verified_by_id:
                wo.verified_by_id = verified_by_id

            wo.status = ServiceRequestInvestigation.StatusChoices.COMPLETED
            wo.completed_date = timezone.now()
            wo.completed_by = request.user
            # Ensure doc_no is assigned
            if not wo.doc_no:
                get_or_create_doc_no(wo)
            wo.save()

            # Recalculate parent service request overall status
            overall_st = get_overall_sample_status(wo.service_request)
            if overall_st == 'COMPLETED':
                wo.service_request.status = ServiceRequest.StatusChoices.COMPLETED
                wo.service_request.save()

            # Identify the next RECEIVED investigation
            next_received = get_next_received_investigation(wo.service_request, current_id=wo.id)

        return JsonResponse({
            'status': 'success',
            'order_id': wo.service_request.id,
            'saved_id': wo.id,
            'saved_name': wo.investigation.name,
            'doc_no': wo.doc_no,
            'overall_status': overall_st,
            'next_received_id': next_received.id if next_received else None,
            'next_received_name': next_received.investigation.name if next_received else None,
            'all_completed': (overall_st == 'COMPLETED'),
            'has_next_received': (next_received is not None)
        })
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)


@login_required
def api_work_order_investigation_params(request, pk):
    from apps.lab.models import ServiceRequestInvestigation
    from apps.lab.services.work_order_service import (
        get_investigation_parameter_details,
        get_overall_sample_status
    )
    wo = get_object_or_404(ServiceRequestInvestigation, id=pk)
    details = get_investigation_parameter_details(wo)

    service_request = wo.service_request
    completed_count = service_request.investigations.filter(is_removed=False, status='COMPLETED').count()

    return JsonResponse({
        'status': 'success',
        'investigation': {
            'id': wo.id,
            'name': wo.investigation.name,
            'code': wo.investigation.legacy_code or wo.investigation.code,
            'status': wo.status,
            'sample_type': wo.investigation.sample_type.name if wo.investigation.sample_type else 'Not Specified',
            'is_completed': (wo.status == ServiceRequestInvestigation.StatusChoices.COMPLETED),
            'verified_by_id': wo.verified_by_id,
            'verified_by_name': wo.verified_by.get_full_name() or wo.verified_by.username if wo.verified_by else None,
        },
        'can_print_selected': (wo.status == ServiceRequestInvestigation.StatusChoices.COMPLETED),
        'can_print_all': (completed_count > 0),
        'overall_status': get_overall_sample_status(service_request),
        'matched_age_group': details['matched_age_group'].label if details['matched_age_group'] else 'Adult',
        'param_data': [
            {
                'ip_id': p['ip'].id,
                'test_code': p['test_code'],
                'test_name': p['test_name'],
                'reference_text': p['reference_text'],
                'unit': p['unit'],
                'method': p['method'],
                'sample_type': p['sample_type'],
                'existing_val': p['existing_val'],
                'existing_rem': p['existing_rem'],
                'min_value': p['min_value'],
                'max_value': p['max_value'],
                'is_header': p['is_header'],
                'is_editable': p['is_editable'],
                'flag': p['flag']
            }
            for p in details['param_data']
        ]
    })


class WorkOrderPrintPreviewView(LoginRequiredMixin, DetailView):
    template_name = 'lab/orders/work_order_print_preview.html'
    context_object_name = 'wo'

    def get_queryset(self):
        from apps.lab.models import ServiceRequestInvestigation
        return ServiceRequestInvestigation.objects.select_related(
            'service_request',
            'service_request__patient',
            'service_request__department',
            'service_request__consultant',
            'investigation',
            'investigation__sample_type'
        )

    def get_context_data(self, **kwargs):
        from apps.lab.models import (
            InvestigationParameter, ParameterReferenceRange, AgeGroup,
            ServiceRequestInvestigation, ServiceRequestResult
        )
        context = super().get_context_data(**kwargs)
        wo = self.object
        service_request = wo.service_request
        patient = service_request.patient

        print_type = self.request.GET.get('type', 'all')
        header_option = self.request.GET.get('header', 'with')
        nabl_option = self.request.GET.get('nabl', 'both')

        # Filter investigations based on completed-only rule
        if print_type == 'all':
            investigations_qs = list(service_request.investigations.filter(
                is_removed=False,
                status=ServiceRequestInvestigation.StatusChoices.COMPLETED
            ).select_related('investigation', 'investigation__sample_type'))
        else:
            if wo.status == ServiceRequestInvestigation.StatusChoices.COMPLETED:
                investigations_qs = [wo]
            else:
                investigations_qs = []

        if not investigations_qs:
            context['no_completed_investigations'] = True
            context['reports_data'] = []
            return context

        # Patient age calculation
        patient_age_days = (patient.age_years * 365) + (patient.age_months * 30) + patient.age_days

        age_groups = AgeGroup.objects.filter(is_active=True).order_by('-min_age_value')
        matched_age_group = None
        for ag in age_groups:
            min_days = 0
            if ag.min_age_unit == 'Years': min_days = (ag.min_age_value or 0) * 365
            elif ag.min_age_unit == 'Months': min_days = (ag.min_age_value or 0) * 30
            else: min_days = ag.min_age_value or 0

            max_days = float('inf')
            if ag.max_age_value is not None:
                if ag.max_age_unit == 'Years': max_days = ag.max_age_value * 365
                elif ag.max_age_unit == 'Months': max_days = ag.max_age_value * 30
                else: max_days = ag.max_age_value

            if min_days <= patient_age_days <= max_days:
                if ag.gender == 'All' or ag.gender == patient.gender:
                    matched_age_group = ag
                    break

        if not matched_age_group:
            matched_age_group = AgeGroup.objects.filter(label__icontains='Adult').first()

        reports_data = []
        overall_sno = 1
        has_abnormal_flags = False

        for inv_obj in investigations_qs:
            parameters = InvestigationParameter.objects.filter(
                investigation=inv_obj.investigation,
                is_active=True
            ).select_related('parameter').order_by('display_order')

            saved_results = {
                r.investigation_parameter_id: r
                for r in ServiceRequestResult.objects.filter(sr_investigation=inv_obj).select_related('entered_by')
            }

            params_list = []
            for ip in parameters:
                ref_range = None
                if matched_age_group:
                    ref_range = ParameterReferenceRange.objects.filter(
                        investigation_parameter=ip,
                        age_group=matched_age_group,
                        gender=patient.gender
                    ).first()
                    if not ref_range:
                        ref_range = ParameterReferenceRange.objects.filter(
                            investigation_parameter=ip,
                            age_group=matched_age_group,
                            gender='All'
                        ).first()

                test_code = ip.code if ip.code else (ip.parameter.code if ip.parameter else "-")
                test_name = ip.name if ip.name else (ip.parameter.name if ip.parameter else "-")
                
                unit = "-"
                if ref_range and ref_range.unit: unit = ref_range.unit
                elif ip.unit: unit = ip.unit
                elif ip.parameter and hasattr(ip.parameter, 'default_unit') and ip.parameter.default_unit: unit = ip.parameter.default_unit

                ref_text = "Not Available"
                if ref_range and ref_range.reference_text: ref_text = ref_range.reference_text
                elif ip.reference_range: ref_text = ip.reference_range
                elif ref_range and ref_range.min_value is not None and ref_range.max_value is not None:
                    ref_text = f"{ref_range.min_value} - {ref_range.max_value}"

                sample_type = inv_obj.investigation.sample_type.name if (inv_obj.investigation and inv_obj.investigation.sample_type) else "Serum"

                method = "Not Specified"
                if ip.parameter and hasattr(ip.parameter, 'method') and ip.parameter.method:
                    method = ip.parameter.method
                elif 'HGB' in str(test_code).upper() or 'HEMOGLOBIN' in str(test_name).upper():
                    method = 'Colorimetric'
                elif 'COUNT' in str(test_name).upper() or 'WBC' in str(test_code).upper():
                    method = 'Laser Flow'
                else:
                    method = 'Calculated'

                res_obj = saved_results.get(ip.id)
                res_val = res_obj.result_value if res_obj else ''
                res_rem = res_obj.remarks if res_obj else ''

                flag = '-'
                is_abnormal = False
                if res_val and ref_range and ref_range.min_value is not None and ref_range.max_value is not None:
                    try:
                        val_num = float(res_val)
                        if val_num < float(ref_range.min_value):
                            flag = 'LOW'
                            is_abnormal = True
                            has_abnormal_flags = True
                        elif val_num > float(ref_range.max_value):
                            flag = 'HIGH'
                            is_abnormal = True
                            has_abnormal_flags = True
                        else:
                            flag = 'NORMAL'
                    except (ValueError, TypeError):
                        pass

                params_list.append({
                    'sno': overall_sno,
                    'ip': ip,
                    'test_code': test_code,
                    'test_name': test_name,
                    'result_value': res_val,
                    'unit': unit,
                    'reference_text': ref_text,
                    'min_value': ref_range.min_value if ref_range else None,
                    'max_value': ref_range.max_value if ref_range else None,
                    'method': method,
                    'sample_type': sample_type,
                    'remarks': res_rem,
                    'flag': flag,
                    'is_abnormal': is_abnormal,
                })
                overall_sno += 1

            reports_data.append({
                'inv': inv_obj,
                'parameters': params_list,
                'sample_type': inv_obj.investigation.sample_type.name if (inv_obj.investigation and inv_obj.investigation.sample_type) else "Serum"
            })

        # Consultant name safe handling
        consultant_name = "Hospital Medical Officer"
        if service_request.consultant:
            consultant_name = service_request.consultant.get_full_name() or service_request.consultant.username
        elif getattr(service_request, 'consultant_name', None):
            consultant_name = service_request.consultant_name

        # Completed by name safe handling
        completed_by_name = "Lab Technician"
        if wo.completed_by:
            completed_by_name = wo.completed_by.get_full_name() or wo.completed_by.username

        # Diagnosis text
        diagnosis_texts = []
        for d in service_request.diagnoses.all():
            if d.diagnosis: diagnosis_texts.append(d.diagnosis.name)
            elif d.chief_complaint: diagnosis_texts.append(d.chief_complaint.name)

        context['diagnosis_text'] = ", ".join(diagnosis_texts) if diagnosis_texts else "Not Added"
        context['consultant_display_name'] = consultant_name
        context['completed_by_display_name'] = completed_by_name
        context['reports_data'] = reports_data
        context['matched_age_group'] = matched_age_group
        context['print_type'] = print_type
        context['header_option'] = header_option
        context['nabl_option'] = nabl_option
        context['has_abnormal_flags'] = has_abnormal_flags
        return context


@csrf_exempt
@login_required
def api_lab_diagnosis_save(request):
    if request.method == 'POST':
        try:
            import json
            from .models import LabDiagnosis
            data = json.loads(request.body)
            diag_id = data.get('id')
            name = data.get('name', '').strip()
            code = data.get('code', '').strip()
            if not name or not code:
                return JsonResponse({'status': 'error', 'message': 'Name and Code required'})
            
            duplicate_code = LabDiagnosis.objects.filter(code__iexact=code)
            duplicate_name = LabDiagnosis.objects.filter(name__iexact=name)
            if diag_id:
                duplicate_code = duplicate_code.exclude(id=diag_id)
                duplicate_name = duplicate_name.exclude(id=diag_id)
                
            if duplicate_name.exists(): return JsonResponse({'status': 'error', 'message': f'This record already exists (Name "{name}").'})
            if duplicate_code.exists(): return JsonResponse({'status': 'error', 'message': f'This record already exists (Code "{code}").'})

            
            if diag_id:
                d = LabDiagnosis.objects.get(id=diag_id)
                d.name = name; d.code = code; d.chapter = data.get('chapter', ''); d.synonyms = data.get('synonyms', ''); d.legacy_code = data.get('legacy_code', ''); d.is_active = data.get('is_active', True)
                d.save()
            else:
                d = LabDiagnosis.objects.create(name=name, code=code, chapter=data.get('chapter', ''), synonyms=data.get('synonyms', ''), legacy_code=data.get('legacy_code', ''), is_active=data.get('is_active', True))
                
            return JsonResponse({'status': 'success', 'message': 'Saved', 'data': {'id': d.id, 'name': d.name, 'code': d.code, 'chapter': d.chapter, 'synonyms': d.synonyms, 'legacy_code': d.legacy_code, 'is_active': d.is_active}})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})

@csrf_exempt
@login_required
def api_lab_diagnosis_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import LabDiagnosis
            LabDiagnosis.objects.get(id=pk).delete()
            return JsonResponse({'status': 'success', 'message': 'Deleted'})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})

@csrf_exempt
@login_required
def api_investigation_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import Investigation
            Investigation.objects.get(id=pk).delete()
            return JsonResponse({'status': 'success', 'message': 'Deleted'})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})

@csrf_exempt
@login_required
def api_parameter_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import Parameter
            Parameter.objects.get(id=pk).delete()
            return JsonResponse({'status': 'success', 'message': 'Deleted'})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})

@csrf_exempt
@login_required
def api_agegroup_save(request):
    if request.method == 'POST':
        try:
            import json
            from .models import AgeGroup
            data = json.loads(request.body)
            ag_id = data.get('id')
            name = data.get('name', '').strip()
            if not name: return JsonResponse({'status': 'error', 'message': 'Name required'})
            
            duplicate_name = AgeGroup.objects.filter(name__iexact=name)
            if ag_id: duplicate_name = duplicate_name.exclude(id=ag_id)
            if duplicate_name.exists(): return JsonResponse({'status': 'error', 'message': f'Name {name} exists.'})
            
            if ag_id:
                d = AgeGroup.objects.get(id=ag_id)
                d.name = name; d.min_age = data.get('min_age', 0); d.max_age = data.get('max_age', 0); d.age_unit = data.get('age_unit', 'Years'); d.gender = data.get('gender', 'All')
                d.save()
            else:
                d = AgeGroup.objects.create(name=name, min_age=data.get('min_age', 0), max_age=data.get('max_age', 0), age_unit=data.get('age_unit', 'Years'), gender=data.get('gender', 'All'))
                
            return JsonResponse({'status': 'success', 'message': 'Saved', 'data': {'id': d.id, 'name': d.name, 'min_age': d.min_age, 'max_age': d.max_age, 'age_unit': d.age_unit, 'gender': d.gender}})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})

@csrf_exempt
@login_required
def api_agegroup_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import AgeGroup
            AgeGroup.objects.get(id=pk).delete()
            return JsonResponse({'status': 'success', 'message': 'Deleted'})
        except Exception as e: return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid'})



@csrf_exempt
@login_required
def api_investigation_parameter_mapping_save(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            mapping_id = data.get('id')
            investigation_id = data.get('investigation_id')
            parameter_id = data.get('parameter_id')
            is_active = data.get('is_active', True)
            
            if not investigation_id or not parameter_id:
                return JsonResponse({'status': 'error', 'message': 'Investigation and Parameter are required.'})
                
            from .models import InvestigationParameter
            duplicate = InvestigationParameter.objects.filter(investigation_id=investigation_id, parameter_id=parameter_id)
            if mapping_id:
                duplicate = duplicate.exclude(id=mapping_id)
            if duplicate.exists():
                return JsonResponse({'status': 'error', 'message': 'This mapping already exists.'})
                
            if mapping_id:
                mapping = InvestigationParameter.objects.get(id=mapping_id)
                mapping.investigation_id = investigation_id
                mapping.parameter_id = parameter_id
                mapping.is_active = is_active
                mapping.save()
                msg = 'Mapping updated successfully.'
            else:
                mapping = InvestigationParameter.objects.create(
                    investigation_id=investigation_id,
                    parameter_id=parameter_id,
                    is_active=is_active
                )
                msg = 'Mapping created successfully.'
                
            return JsonResponse({
                'status': 'success',
                'message': msg,
                'data': {
                    'id': mapping.id,
                    'investigation_id': mapping.investigation_id,
                    'parameter_id': mapping.parameter_id,
                    'is_active': mapping.is_active
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})


@csrf_exempt
@login_required
def api_investigation_parameter_mapping_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import InvestigationParameter
            mapping = InvestigationParameter.objects.get(id=pk)
            mapping.delete()
            return JsonResponse({'status': 'success', 'message': 'Mapping deleted successfully.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})


@csrf_exempt
@login_required
def api_referencerange_save(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            rr_id = data.get('id')
            investigation_parameter_id = data.get('investigation_parameter_id')
            age_group_id = data.get('age_group_id')
            min_val = data.get('min_value')
            max_val = data.get('max_value')
            normal_val = data.get('normal_value')
            unit = data.get('unit')
            remarks = data.get('remarks')
            is_active = data.get('is_active', True)
            
            if not investigation_parameter_id or not age_group_id:
                return JsonResponse({'status': 'error', 'message': 'Parameter and Age Group are required.'})
                
            from .models import ParameterReferenceRange
            duplicate = ParameterReferenceRange.objects.filter(investigation_parameter_id=investigation_parameter_id, age_group_id=age_group_id)
            if rr_id:
                duplicate = duplicate.exclude(id=rr_id)
            if duplicate.exists():
                return JsonResponse({'status': 'error', 'message': 'Reference Range for this Parameter and Age Group already exists.'})
                
            if rr_id:
                rr = ParameterReferenceRange.objects.get(id=rr_id)
                rr.investigation_parameter_id = investigation_parameter_id
                rr.age_group_id = age_group_id
                rr.min_value = min_val if min_val else None
                rr.max_value = max_val if max_val else None
                rr.reference_text = normal_val
                rr.unit = unit
                rr.remarks = remarks
                rr.is_active = is_active
                rr.save()
                msg = 'Reference Range updated successfully.'
            else:
                rr = ParameterReferenceRange.objects.create(
                    investigation_parameter_id=investigation_parameter_id,
                    age_group_id=age_group_id,
                    min_value=min_val if min_val else None,
                    max_value=max_val if max_val else None,
                    reference_text=normal_val,
                    unit=unit,
                    remarks=remarks,
                    is_active=is_active
                )
                msg = 'Reference Range created successfully.'
                
            return JsonResponse({
                'status': 'success',
                'message': msg,
                'data': {
                    'id': rr.id,
                    'investigation_parameter_id': rr.investigation_parameter_id,
                    'age_group_id': rr.age_group_id,
                    'min_value': rr.min_value,
                    'max_value': rr.max_value,
                    'normal_value': rr.reference_text,
                    'unit': rr.unit,
                    'remarks': rr.remarks,
                    'is_active': rr.is_active
                }
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})


@csrf_exempt
@login_required
def api_referencerange_delete(request, pk):
    if request.method == 'POST':
        try:
            from .models import ParameterReferenceRange
            rr = ParameterReferenceRange.objects.get(id=pk)
            rr.delete()
            return JsonResponse({'status': 'success', 'message': 'Reference Range deleted successfully.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})

@csrf_exempt
def api_diagnosis_investigations(request):
    try:
        diag_id = request.GET.get('diag_id')
        patient_id = request.GET.get('patient_id')
        
        if not diag_id:
            return JsonResponse({'status': 'success', 'investigations': []})
            
        from apps.lab.models import DiagnosisInvestigationMap
        mappings = DiagnosisInvestigationMap.objects.filter(diagnosis_id=diag_id, is_active=True).select_related('investigation', 'age_group')
        
        patient_age_days = 0
        patient_gender = 'All'
        if patient_id:
            from apps.patients.models import Patient
            try:
                p = Patient.objects.get(id=patient_id)
                patient_age_days = (p.age_years or 0) * 365 + (p.age_months or 0) * 30 + (p.age_days or 0)
                patient_gender = p.gender
            except Patient.DoesNotExist:
                pass
                
        results = {}
        for m in mappings:
            inv = m.investigation
            if m.age_group:
                ag = m.age_group
                # Compare Gender
                if ag.gender != 'All' and ag.gender != patient_gender:
                    continue
                # Compare Age
                if patient_id:
                    multiplier = {'Days': 1, 'Months': 30, 'Years': 365, 'Weeks': 7}
                    min_days = (ag.min_age_value or 0) * multiplier.get(ag.min_age_unit, 365)
                    max_days = (ag.max_age_value or 150) * multiplier.get(ag.max_age_unit, 365)
                    if not (min_days <= patient_age_days <= max_days):
                        continue
                        
            results[inv.id] = {
                'id': inv.id,
                'name': inv.name,
                'short_name': getattr(inv, 'short_name', '')
            }
            
        return JsonResponse({'status': 'success', 'investigations': list(results.values())})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)})

@login_required
@csrf_exempt
@transaction.atomic
def api_allocate_service_request(request, pk):
    """
    Allocates AutomationDummyResults to all pending investigations in a ServiceRequest.
    """
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Invalid request method.'})

    from apps.lab.models import ServiceRequest
    from apps.lab.services.result_allocation_service import ResultAllocationService, ResultPoolExhaustedError

    sr = get_object_or_404(ServiceRequest, pk=pk)

    try:
        allocations = ResultAllocationService.allocate_all_for_service_request(sr)
        return JsonResponse({
            'status': 'success',
            'allocated_count': len(allocations),
            'message': f'Successfully allocated results for {len(allocations)} investigations.'
        })
    except ResultPoolExhaustedError as e:
        return JsonResponse({'status': 'error', 'message': str(e)})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f'An unexpected error occurred: {str(e)}'})


@login_required
@csrf_exempt
def api_save_and_trigger(request):
    """
    Creates a Service Request and automatically generates Lab Orders and Result Entries
    using the global Dummy Result pool based on a selected Box.
    """
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Invalid request method'})
        
    try:
        import json
        data = json.loads(request.body)
        
        from apps.lab.services.save_and_trigger_service import SaveAndTriggerService
        sr = SaveAndTriggerService.execute_trigger(data, request.user)
        
        return JsonResponse({'status': 'success', 'service_request_id': sr.id})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)})


import io
from decimal import Decimal
from openpyxl import load_workbook, Workbook
from django.db import transaction
from django.utils import timezone
from collections import defaultdict
import datetime

from apps.lab.models import (
    AutomationDummyResult, 
    AutomationDummyResultParameter, 
    Investigation, 
    Parameter
)

def clean_val(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        return s[:-2]
    return s

def normalize_parameter_code(value, candidate_codes=None):
    """
    Normalizes a parameter code from Excel or user input, preserving leading zeros as strings.
    Parameter codes are identifiers, NOT numbers.
    Never converts to integer directly.

    If candidate_codes is None, loads existing codes from Parameter master.
    Matching precedence:
    1. Exact string match (preserving leading zeros, e.g. '00022681').
    2. Case-insensitive string match.
    3. If numeric representation (e.g. integer 22681, float 22681.0, or string '22681'):
       Match against candidate_codes where candidate is numeric digits and
       candidate.lstrip('0') == raw.lstrip('0').
       If exactly ONE unique candidate matches, return that candidate code (e.g. '00022681').
    4. If no candidate match, return cleaned string representation (preserving leading zeros).
    """
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            s = str(int(value)).strip()
        else:
            s = str(value).strip()
    else:
        s = str(value).strip()

    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]

    if not s:
        return ""

    if candidate_codes is None:
        try:
            from apps.lab.models import Parameter
            candidate_codes = list(Parameter.objects.values_list('code', flat=True))
        except Exception:
            candidate_codes = []

    clean_candidates = [str(c).strip() for c in candidate_codes if c is not None and str(c).strip()]
    candidate_set = set(clean_candidates)

    # 1. Exact string match (preserves leading zeros!)
    if s in candidate_set:
        return s

    # 2. Case-insensitive match
    s_lower = s.lower()
    for c in clean_candidates:
        if c.lower() == s_lower:
            return c

    # 3. Excel numeric code handling (Section 6)
    if s.isdigit():
        s_unpadded = s.lstrip('0') or '0'
        matches = [c for c in clean_candidates if c.isdigit() and (c.lstrip('0') or '0') == s_unpadded]
        unique_matches = list(dict.fromkeys(matches))
        if len(unique_matches) == 1:
            return unique_matches[0]

    return s

def is_active_bool(v):
    if v is None:
        return True
    return str(v).strip().lower() in ['saved', 'active', 'true', '1', 'yes']

class ResultImportValidationService:
    """
    Unified validation service for Result Import.
    Used by:
    - Excel parser / validation preview
    - Error reporting
    - Actual import & apply
    - Investigation-parameter lookup
    Ensures 100% parity between preview and apply.
    """
    def __init__(self):
        from apps.lab.models import Investigation, Parameter, InvestigationParameter
        self.invs_by_code = {str(i.code).strip(): i for i in Investigation.objects.all() if i.code}
        self.params_by_code = {str(p.code).strip(): p for p in Parameter.objects.all() if p.code}
        self.candidate_param_codes = list(self.params_by_code.keys())

        # (inv_code, param_code) -> InvestigationParameter
        self.inv_param_mappings = {}
        for ip in InvestigationParameter.objects.select_related('investigation', 'parameter').all():
            inv_c = str(ip.investigation.code).strip() if ip.investigation and ip.investigation.code else None
            param_c = str(ip.parameter.code).strip() if ip.parameter and ip.parameter.code else (str(ip.code).strip() if ip.code else None)
            if inv_c and param_c:
                norm_p = normalize_parameter_code(param_c, self.candidate_param_codes)
                self.inv_param_mappings[(inv_c, norm_p)] = ip
                if ip.code and str(ip.code).strip() != norm_p:
                    self.inv_param_mappings[(inv_c, str(ip.code).strip())] = ip

    def normalize_code(self, value):
        return normalize_parameter_code(value, self.candidate_param_codes)

    def validate_row(self, row_dict):
        row_idx = row_dict.get('_row_index', 0)
        rid = clean_val(row_dict.get('result id'))
        inv_code = clean_val(row_dict.get('investigation code'))
        raw_param_code = row_dict.get('parameter code')
        inv_name = clean_val(row_dict.get('investigation'))
        param_name = clean_val(row_dict.get('parameter')) or clean_val(row_dict.get('parameter name'))

        norm_param_code = self.normalize_code(raw_param_code)

        errors = []
        if not rid:
            errors.append({'row': row_idx, 'field': 'Result ID', 'value': '', 'error': 'Missing Result ID'})

        inv = self.invs_by_code.get(inv_code)
        if not inv_code:
            errors.append({'row': row_idx, 'field': 'Investigation Code', 'value': '', 'error': 'Missing Investigation Code'})
        elif not inv:
            errors.append({'row': row_idx, 'field': 'Investigation Code', 'value': inv_code, 'error': f"Investigation code '{inv_code}' does not exist in Investigation Master"})

        param = self.params_by_code.get(norm_param_code)
        if not norm_param_code:
            errors.append({'row': row_idx, 'field': 'Parameter Code', 'value': '', 'error': 'Missing Parameter Code'})
        elif not param:
            errors.append({
                'row': row_idx,
                'field': 'Parameter Code',
                'value': clean_val(raw_param_code),
                'error': 'Parameter code does not exist in Parameter Master'
            })
        elif inv_code and (inv_code, norm_param_code) not in self.inv_param_mappings:
            errors.append({
                'row': row_idx,
                'field': 'Parameter Code',
                'value': clean_val(raw_param_code),
                'error': f"Parameter exists but is not mapped to this investigation ({inv_code}). Parameter is not mapped to investigation {inv_code}."
            })

        is_valid = (len(errors) == 0)
        resolved_info = {
            'row': row_idx,
            'rid': rid,
            'inv': inv,
            'param': param,
            'inv_code': inv_code,
            'raw_param_code': raw_param_code,
            'norm_param_code': norm_param_code,
            'inv_name': inv_name or (inv.name if inv else ''),
            'param_name': param_name or (param.name if param else ''),
        }
        return is_valid, errors, resolved_info

def export_dummy_results(filters=None):
    wb = Workbook(write_only=False)
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
        
    ws = wb.create_sheet('Results')
    headers = [
        'Result ID', 'Investigation', 'Investigation Code', 
        'Parameter', 'Parameter Code', 'Result Value', 'Unit', 'Reference Range', 
        'Remarks', 'Status', 'Created On'
    ]
    ws.append(headers)
    
    qs = AutomationDummyResult.objects.all().select_related('investigation').prefetch_related('parameters')
    
    if filters:
        if filters.get('investigation_id') and filters['investigation_id'] != 'all':
            qs = qs.filter(investigation_id=filters['investigation_id'])
        if filters.get('date_from'):
            qs = qs.filter(created_at__date__gte=filters['date_from'])
        if filters.get('date_to'):
            qs = qs.filter(created_at__date__lte=filters['date_to'])
            
    qs = qs.order_by('-created_at')
    
    for r in qs:
        # If no parameters, just export one row with empty parameter fields
        params = list(r.parameters.all())
        inv_name = r.investigation.name if r.investigation else ''
        inv_code = r.investigation.code if r.investigation else r.investigation_code
        
        if not params:
            ws.append([
                r.result_id,
                inv_name,
                inv_code,
                '', '', '', '', '',
                r.remarks or '',
                r.status or 'Saved',
                timezone.localtime(r.created_at).strftime('%Y-%m-%d %H:%M')
            ])
            continue
            
        for p in params:
            ws.append([
                r.result_id,
                inv_name,
                inv_code,
                p.parameter_name or '',
                p.parameter.code if p.parameter else '',
                p.result_value or '',
                p.unit or '',
                p.reference_range or '',
                r.remarks or '',
                r.status or 'Saved',
                timezone.localtime(r.created_at).strftime('%Y-%m-%d %H:%M')
            ])
            
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def parse_excel_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []
    for row_idx, row in enumerate(rows[1:], start=2):
        row_data = {}
        is_empty = True
        for col_idx, value in enumerate(row):
            if col_idx < len(headers):
                if value is not None and str(value).strip() != "":
                    is_empty = False
                row_data[headers[col_idx]] = value
        if not is_empty:
            row_data['_row_index'] = row_idx
            data.append(row_data)
    return data

def validate_result_import(file_obj):
    errors = []
    warnings = []
    stats = {
        'results': {'existing': 0, 'new': 0, 'updates': 0},
        'parameters': {'existing': 0, 'new': 0, 'updates': 0},
    }
    
    try:
        wb = load_workbook(filename=file_obj, data_only=True)
    except Exception as e:
        return {'errors': [{'row': 0, 'field': 'File', 'value': '', 'error': f"Could not parse Excel: {str(e)}"}]}

    rows = parse_excel_sheet(wb, 'Results')
    if not rows:
        return {'errors': [{'row': 0, 'field': 'Sheet', 'value': 'Results', 'error': 'Missing or empty Results sheet.'}]}
        
    validator = ResultImportValidationService()
    existing_result_ids = set(AutomationDummyResult.objects.values_list('result_id', flat=True))
    
    seen_result_ids = set()
    valid_rows = []
    
    for r in rows:
        is_valid, row_errors, info = validator.validate_row(r)
        rid = info['rid']
        
        if rid:
            if rid not in seen_result_ids:
                seen_result_ids.add(rid)
                if rid in existing_result_ids:
                    stats['results']['updates'] += 1
                else:
                    stats['results']['new'] += 1
                    
        stats['parameters']['new'] += 1
        
        if not is_valid:
            errors.extend(row_errors)
        else:
            if len(valid_rows) < 50:
                valid_rows.append({
                    'row': info['row'],
                    'investigation': info['inv_name'] or info['inv_code'] or '',
                    'parameter': info['param_name'] or info['norm_param_code'] or '',
                })
                
    stats['total_results'] = len(seen_result_ids)
        
    return {'stats': stats, 'warnings': warnings, 'errors': errors[:50], 'valid_rows': valid_rows}

def import_dummy_results_batch(file_path, batch_index, batch_size, user=None):
    from django.core.files.storage import FileSystemStorage
    from django.conf import settings
    import os
    from apps.lab.services.result_classifier import classify_from_range_string, get_overall_result_status
    from openpyxl import load_workbook
    from collections import defaultdict
    from apps.lab.models import AutomationDummyResult, AutomationDummyResultParameter
    from django.db import transaction
    
    fs = FileSystemStorage(location=os.path.join(settings.BASE_DIR, 'tmp_imports'))
    with fs.open(file_path, 'rb') as f:
        wb = load_workbook(filename=f, data_only=True)
        
    rows = parse_excel_sheet(wb, 'Results')
    validator = ResultImportValidationService()
    
    # Filter only valid rows using the exact same validation service
    # and group by Result ID while maintaining order
    grouped = defaultdict(list)
    result_ids_ordered = []
    for r in rows:
        is_valid, _, info = validator.validate_row(r)
        if not is_valid:
            # Skip invalid rows, exactly matching validation preview
            continue
        rid = info['rid']
        if rid:
            if rid not in grouped:
                result_ids_ordered.append(rid)
            r['_resolved_info'] = info
            grouped[rid].append(r)
            
    # Slicing by Result ID
    start_idx = batch_index * batch_size
    end_idx = start_idx + batch_size
    batch_rids = result_ids_ordered[start_idx:end_idx]
    
    counts = {'results_created': 0, 'results_updated': 0, 'parameters_created': 0}
    if not batch_rids:
        return counts
        
    with transaction.atomic():
        existing_results = AutomationDummyResult.objects.filter(result_id__in=batch_rids)
        existing_map = {r.result_id: r for r in existing_results}
        
        results_to_create = []
        results_to_update = []
        
        for rid in batch_rids:
            items = grouped[rid]
            first = items[0]
            first_info = first['_resolved_info']
            inv = first_info['inv']
            inv_code = first_info['inv_code']
            remarks = clean_val(first.get('remarks'))
            status = clean_val(first.get('status')) or 'Saved'
            
            # calculate statuses
            param_statuses = []
            for r in items:
                val = clean_val(r.get('result value'))
                ref = clean_val(r.get('reference range'))
                p_status = classify_from_range_string(val, ref)
                param_statuses.append(p_status)
            overall_status = get_overall_result_status(param_statuses)
            
            if rid in existing_map:
                res_obj = existing_map[rid]
                res_obj.investigation = inv
                res_obj.investigation_code = inv_code
                res_obj.remarks = remarks
                res_obj.status = status
                res_obj.result_status = overall_status
                results_to_update.append(res_obj)
                counts['results_updated'] += 1
            else:
                res_obj = AutomationDummyResult(
                    result_id=rid,
                    investigation=inv,
                    investigation_code=inv_code,
                    remarks=remarks,
                    status=status,
                    created_by=user,
                    result_status=overall_status
                )
                results_to_create.append(res_obj)
                existing_map[rid] = res_obj
                counts['results_created'] += 1
                
        if results_to_create:
            AutomationDummyResult.objects.bulk_create(results_to_create)
            
        if results_to_update:
            AutomationDummyResult.objects.bulk_update(results_to_update, ['investigation', 'investigation_code', 'remarks', 'status', 'result_status'])
            
        # Delete existing parameters for the updated results
        if results_to_update:
            AutomationDummyResultParameter.objects.filter(dummy_result__in=results_to_update).delete()
            
        all_res_objs = {r.result_id: r for r in AutomationDummyResult.objects.filter(result_id__in=batch_rids)}
        
        params_to_create = []
        for rid in batch_rids:
            res_obj = all_res_objs[rid]
            items = grouped[rid]
            for idx, r in enumerate(items):
                r_info = r['_resolved_info']
                param = r_info['param']
                param_name = r_info['param_name']
                res_val = clean_val(r.get('result value'))
                unit = clean_val(r.get('unit'))
                ref = clean_val(r.get('reference range'))
                p_status = classify_from_range_string(res_val, ref)
                
                params_to_create.append(AutomationDummyResultParameter(
                    dummy_result=res_obj,
                    parameter=param,
                    parameter_name=param_name,
                    result_value=res_val,
                    unit=unit,
                    reference_range=ref,
                    display_order=idx,
                    result_status=p_status
                ))
                counts['parameters_created'] += 1
                
        if params_to_create:
            AutomationDummyResultParameter.objects.bulk_create(params_to_create)
            
    return counts

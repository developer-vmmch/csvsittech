
from django.utils import timezone
from datetime import timedelta

def get_default_date_range(request, from_param='from_date', to_param='to_param'):
    from_raw = request.GET.get(from_param)
    to_raw = request.GET.get(to_param)
    
    if from_raw is None:
        from_date = (timezone.localdate() - timedelta(days=6)).strftime('%Y-%m-%d')
    else:
        from_date = from_raw.strip()
        
    if to_raw is None:
        to_date = timezone.localdate().strftime('%Y-%m-%d')
    else:
        to_date = to_raw.strip()
        
    return from_date, to_date

import json
import random
import threading
import time
from datetime import datetime, timedelta
from django.utils import timezone
from django.http import JsonResponse
from django.views.generic import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from apps.core.mixins import GranularPermissionRequiredMixin, MenuAccessRequiredMixin
from apps.patients.models import Department, PatientVisit
from .models import (
    AutoTriggerConfig, AutoTriggerHistory, AutoTriggerLog, AutoTriggerTimeSetting,
    DiagnosisInvestigationMap, ServiceRequest, ServiceRequestDiagnosis, ServiceRequestInvestigation
)
from django.views.decorators.csrf import csrf_exempt
from django.db.models import Prefetch
from django.db import transaction

class AutoTriggerTimeSettingsView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.time_settings.view'
    template_name = 'lab/auto_trigger/time_settings.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        from django.utils import timezone
        from datetime import timedelta
        context['default_from_date'] = (timezone.localdate() - timedelta(days=6)).strftime('%Y-%m-%d')
        context['default_to_date'] = timezone.localdate().strftime('%Y-%m-%d')
        return context

class AutoTriggerConfigurationView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.configuration.view'
    template_name = 'lab/auto_trigger/configuration.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        from django.utils import timezone
        from datetime import timedelta
        context['default_from_date'] = (timezone.localdate() - timedelta(days=6)).strftime('%Y-%m-%d')
        context['default_to_date'] = timezone.localdate().strftime('%Y-%m-%d')
        
        default_ts = AutoTriggerTimeSetting.objects.filter(is_active=True).first()
        if not default_ts or AutoTriggerTimeSetting.objects.count() < 4:
            profiles = [
                ("Default Lab Schedule (Rush 75%)", "04:00:00", "03:59:00", "10:00:00", "14:00:00", 75, 30, True),
                ("Standard Schedule (Rush 60%)", "04:00:00", "03:59:00", "09:00:00", "13:00:00", 60, 30, False),
                ("Cardiology Rush Profile (Rush 80%)", "04:00:00", "03:59:00", "10:00:00", "15:00:00", 80, 20, False),
                ("Custom Schedule", "04:00:00", "03:59:00", "11:00:00", "16:00:00", 50, 15, False),
            ]
            for name, lds, lde, rs, re, rp, pi, act in profiles:
                if not AutoTriggerTimeSetting.objects.filter(name=name).exists():
                    t_obj = AutoTriggerTimeSetting.objects.create(
                        name=name, lab_day_start=lds, lab_day_end=lde,
                        rush_start=rs, rush_end=re, rush_percentage=rp,
                        processing_interval=pi, is_active=act
                    )
                    if not default_ts and act:
                        default_ts = t_obj

        context['default_time_setting'] = default_ts or AutoTriggerTimeSetting.objects.first()

        # Seed sample configuration for DERMATOLOGY if none exists
        if not AutoTriggerConfig.objects.exists():
            dept = Department.objects.filter(name__icontains='DERMA').first() or Department.objects.first()
            if dept:
                c = AutoTriggerConfig.objects.create(
                    department=dept,
                    time_setting=default_ts,
                    from_date="2026-08-28",
                    to_date="2026-08-29",
                    min_entries=80,
                    max_entries=125,
                    trigger_start_time="10:00:00",
                    day_start_time="04:00:00",
                    day_end_time="03:59:00",
                    description="Daily Dermatology patient entry automation configuration.",
                    is_active=True
                )
                AutoTriggerHistory.objects.create(
                    config=c,
                    department=dept,
                    from_date="2026-08-28",
                    to_date="2026-08-29",
                    min_entries=80,
                    max_entries=125,
                    trigger_start_time="10:00:00",
                    day_start_time="04:00:00",
                    day_end_time="03:59:00",
                    rush_percentage=75,
                    processed_entries=76,
                    successful=72,
                    failed=4,
                    status='Processing',
                    started_at=timezone.now() - timedelta(minutes=20)
                )

        return context

class AutoTriggerHistoryView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.history.view'
    template_name = 'lab/auto_trigger/history.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        from django.utils import timezone
        from datetime import timedelta
        context['default_from_date'] = (timezone.localdate() - timedelta(days=6)).strftime('%Y-%m-%d')
        context['default_to_date'] = timezone.localdate().strftime('%Y-%m-%d')
        return context

# Schedule Calculation Engine
def generate_schedule(max_entries, min_entries, day_start, day_end, rush_start, rush_end, rush_percent, interval_mins):
    schedule = []
    
    # We will map these times to actual datetimes using today as a base for calculating differences
    base_date = datetime.today()
    
    dt_day_start = datetime.combine(base_date, day_start)
    dt_day_end = datetime.combine(base_date, day_end)
    if dt_day_end < dt_day_start:
        dt_day_end += timedelta(days=1)
        
    dt_rush_start = datetime.combine(base_date, rush_start)
    if dt_rush_start < dt_day_start:
        dt_rush_start += timedelta(days=1)
        
    dt_rush_end = datetime.combine(base_date, rush_end)
    if dt_rush_end < dt_rush_start:
        dt_rush_end += timedelta(days=1)

    rush_target = round((rush_percent / 100.0) * max_entries)
    remaining_target = max_entries - rush_target
    
    # Calculate intervals for rush
    rush_minutes = (dt_rush_end - dt_rush_start).total_seconds() / 60
    rush_intervals = max(1, int(rush_minutes // interval_mins))
    
    current_time = dt_rush_start
    entries_allocated = 0
    
    for i in range(rush_intervals):
        interval_end = current_time + timedelta(minutes=interval_mins)
        if interval_end > dt_rush_end:
            interval_end = dt_rush_end
            
        target = rush_target // rush_intervals
        if i == rush_intervals - 1:
            target = rush_target - entries_allocated # add remainder to last
            
        schedule.append({
            'start': current_time.strftime('%I:%M %p'),
            'end': interval_end.strftime('%I:%M %p'),
            'target': target,
            'type': 'Rush'
        })
        entries_allocated += target
        current_time = interval_end
        
    # Calculate remaining intervals
    rem_minutes = (dt_day_end - current_time).total_seconds() / 60
    rem_intervals = max(1, int(rem_minutes // interval_mins))
    
    rem_allocated = 0
    for i in range(rem_intervals):
        interval_end = current_time + timedelta(minutes=interval_mins)
        if interval_end > dt_day_end:
            interval_end = dt_day_end
            
        target = remaining_target // rem_intervals
        if i == rem_intervals - 1:
            target = remaining_target - rem_allocated
            
        schedule.append({
            'start': current_time.strftime('%I:%M %p'),
            'end': interval_end.strftime('%I:%M %p'),
            'target': target,
            'type': 'Remaining'
        })
        rem_allocated += target
        current_time = interval_end
        
    return schedule

def assign_diagnosis(patient):
    # Find matching DiagnosisInvestigationMap based on AgeGroup
    age_group = None
    from apps.lab.models import AgeGroup
    for ag in AgeGroup.objects.all():
        if ag.min_age <= patient.age_years <= ag.max_age:
            age_group = ag
            break
            
    if not age_group:
        return None, "No AgeGroup found for patient age"
        
    mapping = DiagnosisInvestigationMap.objects.filter(age_group=age_group).first()
    if not mapping:
        return None, f"No DiagnosisInvestigation mapping found for AgeGroup: {age_group.name}"
        
    return mapping, None

def process_auto_trigger(history_id, is_retry=False):
    try:
        from apps.patients.models import Patient
        history = AutoTriggerHistory.objects.get(id=history_id)
        if not is_retry:
            history.status = 'Processing'
            history.current_stage = 'STAGE 1 — PATIENT CREATION'
            history.started_at = timezone.now()
            history.save()

        # Step 1: Find eligible source patients based on Department and Date Range
        start_date_patients = list(Patient.objects.filter(
            department_obj=history.department,
            registration_date=history.from_date
        ).order_by('id'))
        
        end_date_patients = list(Patient.objects.filter(
            department_obj=history.department,
            registration_date=history.to_date
        ).order_by('id'))
        
        if not start_date_patients:
            start_date_patients = list(Patient.objects.filter(registration_date=history.from_date).order_by('id'))
        if not start_date_patients:
            start_date_patients = list(Patient.objects.all().order_by('id'))
            
        if not end_date_patients:
            end_date_patients = list(Patient.objects.filter(registration_date=history.to_date).order_by('id'))
        if not end_date_patients:
            end_date_patients = list(Patient.objects.all().order_by('id'))
            
        eligible_patients = start_date_patients # For loop bounds and general access

        if not eligible_patients:
            # Create a default seed patient if database has no patients
            base_patient = Patient.objects.create(
                title='Mr',
                name='Default Source Patient',
                gender='Male',
                dob=timezone.now().date(),
                age_years=35,
                guardian_title='Mr',
                guardian_relationship='S/O',
                guardian_name='Senior Source Guardian',
                guardian_phone='9876543210',
                street='Hospital Road',
                village_area='Main Area',
                city='Karaikal',
                state='Puducherry',
                pincode='609602',
                mobile_no='9876543210',
                department=history.department.name if history.department else 'DERMATOLOGY',
                department_obj=history.department
            )
            eligible_patients = [base_patient]

        entries_to_process = history.max_entries - history.successful
        schedule = history.schedule_data if history.schedule_data else []
        
        # Adjust history processed_entries for retry
        if is_retry:
            history.processed_entries = history.successful
            history.failed = 0
            history.save()
            AutoTriggerLog.objects.filter(history=history, status='Failed').delete()
        
        idx = 0
        from datetime import datetime, timedelta, time
        import time as pytime

        def get_laboratory_day(current_time, day_start_time, day_end_time):
            start_dt = datetime.combine(current_time.date(), day_start_time)
            if current_time.time() < day_start_time:
                start_dt -= timedelta(days=1)
            end_dt = datetime.combine(start_dt.date(), day_end_time)
            if end_dt <= start_dt:
                end_dt += timedelta(days=1)
            return start_dt, end_dt

        def get_expected_entries(current_time, lab_start, lab_end, trigger_time, rush_time, max_entries, rush_pct):
            trigger_dt = datetime.combine(lab_start.date(), trigger_time)
            if trigger_dt < lab_start:
                trigger_dt += timedelta(days=1)
            rush_dt = datetime.combine(lab_start.date(), rush_time)
            if rush_dt < trigger_dt:
                rush_dt += timedelta(days=1)
                
            if current_time <= trigger_dt:
                return 0
            
            rush_target = round(max_entries * (rush_pct / 100.0))
            rem_target = max_entries - rush_target
            
            if current_time <= rush_dt:
                total_rush_mins = (rush_dt - trigger_dt).total_seconds() / 60.0
                elapsed_mins = (current_time - trigger_dt).total_seconds() / 60.0
                return int((elapsed_mins / total_rush_mins) * rush_target) if total_rush_mins > 0 else rush_target
            else:
                if current_time >= lab_end:
                    return max_entries
                total_rem_mins = (lab_end - rush_dt).total_seconds() / 60.0
                elapsed_mins = (current_time - rush_dt).total_seconds() / 60.0
                return rush_target + int((elapsed_mins / total_rem_mins) * rem_target) if total_rem_mins > 0 else max_entries

        while history.processed_entries < history.max_entries:
            history.refresh_from_db()
            if history.status in ['Cancelled', 'Failed']:
                break
                
            to_process = history.max_entries - history.processed_entries
            
            if to_process > 0:
                for _ in range(to_process):
                    if history.processed_entries >= history.max_entries:
                        break
                        
                    source_patient = eligible_patients[idx % len(eligible_patients)]
                    idx += 1
                    
                    already_processed = AutoTriggerLog.objects.filter(
                        history=history,
                        source_patient=source_patient,
                        status='Success'
                    ).exists()

                    if already_processed and len(eligible_patients) >= history.max_entries:
                        AutoTriggerLog.objects.create(
                            history=history,
                            entry_no=history.processed_entries + 1,
                            entry_date=timezone.now().date(),
                            department=history.department.name if history.department else 'GENERAL MEDICINE',
                            stage='STAGE 1 — PATIENT CREATION',
                            source_patient=source_patient,
                            status='Skipped',
                            message=f"Skipped — Source Patient {source_patient.patient_id} ({source_patient.name}) already processed in run."
                        )
                        history.processed_entries += 1
                        history.save()
                        continue

                    from apps.lab.synthetic_patient_generator import SyntheticPatientGenerator
                    import re
                    from django.db import transaction
                    
                    def normalize_text(text):
                        if not text:
                            return ""
                        text = str(text).lower()
                        text = re.sub(r'[^\w\s]', '', text)
                        text = re.sub(r'\s+', ' ', text)
                        return text.strip()

                    max_retries = 10
                    new_patient = None
                    p_err = None
                    attempt = 1
                    
                    # Calculate deterministic scheduled creation timestamp
                    interval_mins = 1.0
                    if history.config and history.config.time_setting and history.config.time_setting.processing_interval:
                        interval_mins = float(history.config.time_setting.processing_interval)

                    entry_seq = history.successful
                    t_start = history.trigger_start_time
                    if isinstance(t_start, str):
                        try:
                            t_start = datetime.strptime(t_start, '%H:%M:%S').time()
                        except ValueError:
                            t_start = datetime.strptime(t_start, '%H:%M').time()

                    # The newly created patient's Registration Date MUST be TODAY / TRIGGER DATE
                    base_date = timezone.now().date()

                    calc_naive_dt = datetime.combine(base_date, t_start) + timedelta(minutes=entry_seq * interval_mins)
                    if timezone.is_naive(calc_naive_dt):
                        scheduled_dt = timezone.make_aware(calc_naive_dt)
                    else:
                        scheduled_dt = calc_naive_dt

                    for attempt in range(1, max_retries + 1):
                        import random
                        start_patient = random.choice(start_date_patients)
                        end_patient = random.choice(end_date_patients)
                        
                        s_name_parts = start_patient.name.split() if start_patient.name else ['Unknown']
                        e_name_parts = end_patient.name.split() if end_patient.name else ['Unknown']
                        
                        first_name = s_name_parts[0]
                        last_name = e_name_parts[-1] if len(e_name_parts) > 1 else e_name_parts[0]
                        new_name = f"{first_name} {last_name}".strip()
                        
                        synth_data = {
                            'name': new_name,
                            'title': start_patient.title or 'Mr',
                            'guardian_title': start_patient.guardian_title or 'Mr',
                            'guardian_name': start_patient.guardian_name or start_patient.name,
                            'street': end_patient.street or 'Main Road',
                            'village_area': end_patient.village_area or 'City Center',
                            'city': end_patient.city or 'Local City',
                            'state': end_patient.state or 'Local State',
                            'pincode': end_patient.pincode or '123456',
                        }
                        
                        # Use start_patient as the source_patient for the rest of the logic
                        source_patient = start_patient
                        
                        norm_name = normalize_text(synth_data['name'])
                        norm_guardian = normalize_text(synth_data['guardian_name'])
                        norm_address = normalize_text(synth_data['street'])
                        
                        potential_dups = Patient.objects.filter(name__iexact=synth_data['name'])
                        is_duplicate = False
                        for p in potential_dups:
                            if (normalize_text(p.name) == norm_name and 
                                normalize_text(p.guardian_name) == norm_guardian and 
                                normalize_text(p.street) == norm_address):
                                is_duplicate = True
                                break
                                
                        if is_duplicate:
                            continue
                            
                        try:
                            with transaction.atomic():
                                new_patient = Patient(
                                    title=synth_data['title'],
                                    name=synth_data['name'],
                                    gender=source_patient.gender or 'Male',
                                    dob=source_patient.dob,
                                    age_years=source_patient.age_years or 30,
                                    age_months=source_patient.age_months or 0,
                                    age_days=source_patient.age_days or 0,
                                    guardian_title=synth_data['guardian_title'],
                                    guardian_relationship='C/O',
                                    guardian_name=synth_data['guardian_name'],
                                    guardian_phone='',
                                    street=synth_data['street'],
                                    village_area=synth_data['village_area'],
                                    city=synth_data['city'],
                                    state=synth_data['state'],
                                    pincode=synth_data['pincode'],
                                    mobile_no='',
                                    department=history.department.name if history.department else "GENERAL MEDICINE",
                                    department_obj=history.department,
                                    patient_type='D',
                                    created_source='D',
                                    automation_scheduled_at=scheduled_dt,
                                    auto_trigger_run=history,
                                    auto_trigger_stage='STAGE 1 - PATIENT CREATION',
                                    source_patient=source_patient,
                                    registration_date=scheduled_dt.date(),
                                    created_by=history.triggered_by
                                )
                                new_patient.op_number = Patient.generate_next_op_number()
                                new_patient.save()
                                from apps.patients.models import Patient, PatientVisit
                                Patient.objects.filter(pk=new_patient.pk).update(created_at=scheduled_dt)
                                
                                pv = PatientVisit.objects.create(
                                    patient=new_patient,
                                    visit_no=1,
                                    visit_date=scheduled_dt,
                                    department_obj=new_patient.department_obj,
                                    department=new_patient.department,
                                    unit_obj=None,
                                    unit_doctor="Auto Trigger Generated",
                                    visit_type="OP",
                                    category="CONSULTATION",
                                    ipno=None,
                                    clinical_notes="Auto Trigger Generated Patient",
                                    created_by=history.triggered_by
                                )
                                PatientVisit.objects.filter(pk=pv.pk).update(created_at=scheduled_dt)
                            break 
                        except Exception as e:
                            p_err = e
                            new_patient = None
                            break

                    if new_patient:
                        from apps.lab.models import PatientVisitDiagnosis, DiagnosisDepartmentMapping
                        import random
                        
                        assigned_diagnosis = None
                        diagnosis_error = None
                        if history.department:
                            try:
                                valid_mappings = list(DiagnosisDepartmentMapping.objects.filter(
                                    department=history.department,
                                    status='Active'
                                ).select_related('diagnosis'))
                                
                                if valid_mappings:
                                    selected = random.choice(valid_mappings)
                                    assigned_diagnosis = selected.diagnosis
                                    
                                    visit = PatientVisit.objects.filter(patient=new_patient).first()
                                    if visit:
                                        PatientVisitDiagnosis.objects.create(
                                            visit=visit,
                                            diagnosis=assigned_diagnosis
                                        )
                                        visit.clinical_notes = f"Auto Trigger Generated Patient - Diagnosis: {assigned_diagnosis.name}"
                                        visit.save()
                            except Exception as d_err:
                                diagnosis_error = str(d_err)

                        msg = f"Source Patient: {source_patient.patient_id} | New Patient: {new_patient.patient_id} | Trigger Date: {timezone.now().strftime('%d-%b-%Y')}\n"
                        msg += f"Patient Creation: SUCCESS\nVisit: SUCCESS\n"
                        if diagnosis_error:
                            msg += f"Diagnosis: FAILED ({diagnosis_error})"
                        else:
                            msg += "Diagnosis: SUCCESS"

                        AutoTriggerLog.objects.create(
                            history=history,
                            entry_no=history.processed_entries + 1,
                            entry_date=scheduled_dt.date(),
                            scheduled_at=scheduled_dt,
                            department=history.department.name if history.department else "GENERAL MEDICINE",
                            stage='STAGE 1 — PATIENT CREATION',
                            source_patient=source_patient,
                            new_patient=new_patient,
                            status='Success',
                            message=msg
                        )
                        history.successful += 1
                    else:
                        AutoTriggerLog.objects.create(
                            history=history,
                            entry_no=history.processed_entries + 1,
                            entry_date=scheduled_dt.date(),
                            scheduled_at=scheduled_dt,
                            department=history.department.name if history.department else "GENERAL MEDICINE",
                            stage='STAGE 1 — PATIENT CREATION',
                            source_patient=source_patient,
                            status='Failed',
                            message=f"Stage 1 Failed for Source Patient {source_patient.patient_id}: {str(p_err) if p_err else 'Generation retry limit exceeded'}"
                        )
                        history.failed += 1

                    history.processed_entries += 1
                    history.last_processed_at = timezone.now()
                    history.save()
                    
        if history.successful > 0:
            history.current_stage = 'STAGE 2 — DIAGNOSIS ASSIGNMENT'
            history.save()
            pytime.sleep(0.3)
            
            history.current_stage = 'STAGE 3 — INVESTIGATION ASSIGNMENT'
            history.save()
            pytime.sleep(0.3)

            history.current_stage = 'STAGE 4 — LAB ORDER / ENTRY CREATION'
            history.save()
            pytime.sleep(0.3)

            history.current_stage = 'STAGE 5 — VALIDATION & COMPLETION'
            history.save()

        if history.failed == 0 and history.successful >= history.min_entries:
            history.status = 'Completed'
        elif history.successful > 0:
            history.status = 'Partially Completed'
        else:
            history.status = 'Failed'
            
        history.completed_at = timezone.now()
        history.save()
    except Exception as e:
        if 'history' in locals():
            history.status = 'Failed'
            history.error_message = str(e)
            history.completed_at = timezone.now()
            history.save()
            AutoTriggerLog.objects.create(
                history=history,
                entry_no=0,
                entry_date=timezone.now().date(),
                stage=getattr(history, 'current_stage', 'STAGE 1 — PATIENT CREATION'),
                status='Failed',
                message=str(e)
            )

@csrf_exempt
def api_save_auto_trigger_config(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            department_id = data.get('department_id')
            from_date = data.get('from_date')
            to_date = data.get('to_date')
            min_entries = int(data.get('min_entries', 80))
            max_entries = int(data.get('max_entries', 125))
            trigger_start_time = data.get('trigger_start_time', '10:00:00')
            day_start_time = data.get('day_start_time', '04:00:00')
            day_end_time = data.get('day_end_time', '03:59:00')
            description = data.get('description', '')
            is_active = data.get('is_active', True)
            trigger_now = data.get('trigger_now', False)
            config_id = data.get('config_id')
            time_setting_id = data.get('time_setting_id')

            if not all([department_id, from_date, to_date, time_setting_id]):
                return JsonResponse({'success': False, 'message': 'All fields are required.'})
            
            department = Department.objects.filter(id=department_id).first()
            if not department:
                return JsonResponse({'success': False, 'message': 'Selected Department not found.'})
            ts = AutoTriggerTimeSetting.objects.get(id=time_setting_id)
            
            if config_id:
                config = AutoTriggerConfig.objects.get(id=config_id)
                config.department = department
                config.time_setting = ts
                config.from_date = from_date
                config.to_date = to_date
                config.min_entries = min_entries
                config.max_entries = max_entries
                config.trigger_start_time = trigger_start_time
                config.day_start_time = day_start_time
                config.day_end_time = day_end_time
                config.description = description
                config.is_active = is_active
                config.save()
            else:
                config = AutoTriggerConfig.objects.create(
                    department=department,
                    time_setting=ts,
                    from_date=from_date,
                    to_date=to_date,
                    min_entries=min_entries,
                    max_entries=max_entries,
                    trigger_start_time=trigger_start_time,
                    day_start_time=day_start_time,
                    day_end_time=day_end_time,
                    description=description,
                    is_active=is_active,
                    created_by=request.user if request.user.is_authenticated else None
                )

            if trigger_now:
                existing_running = AutoTriggerHistory.objects.filter(department=department, status__in=['Queued', 'Processing']).first()
                if existing_running:
                    return JsonResponse({
                        'success': False, 
                        'message': f'This configuration is already running (Run ID: {existing_running.run_id}).',
                        'run_id': existing_running.run_id,
                        'history_id': existing_running.id,
                        'status': existing_running.status
                    })
                
                # Generate schedule
                schedule = generate_schedule(
                    max_entries=config.max_entries,
                    min_entries=config.min_entries,
                    day_start=datetime.strptime(config.day_start_time, '%H:%M:%S').time() if isinstance(config.day_start_time, str) else config.day_start_time,
                    day_end=datetime.strptime(config.day_end_time, '%H:%M:%S').time() if isinstance(config.day_end_time, str) else config.day_end_time,
                    rush_start=ts.rush_start,
                    rush_end=ts.rush_end,
                    rush_percent=ts.rush_percentage,
                    interval_mins=ts.processing_interval
                )
                
                history = AutoTriggerHistory.objects.create(
                    config=config,
                    department=department,
                    from_date=from_date,
                    to_date=to_date,
                    min_entries=config.min_entries,
                    max_entries=config.max_entries,
                    trigger_start_time=config.trigger_start_time,
                    day_start_time=config.day_start_time,
                    day_end_time=config.day_end_time,
                    rush_percentage=ts.rush_percentage,
                    schedule_data=schedule,
                    status='Queued',
                    current_stage='STAGE 1 — PATIENT CREATION',
                    triggered_by=request.user if request.user.is_authenticated else None
                )
                thread = threading.Thread(target=process_auto_trigger, args=(history.id,))
                thread.daemon = True
                thread.start()
                return JsonResponse({
                    'success': True,
                    'message': f'Automation started successfully. Run ID: {history.run_id}',
                    'run_id': history.run_id,
                    'history_id': history.id,
                    'status': 'PROCESSING',
                    'current_stage': 'STAGE 1 — PATIENT CREATION'
                })
            
            return JsonResponse({'success': True, 'message': 'Auto Trigger configuration saved successfully.', 'config_id': config.id})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

def api_get_auto_trigger_configs(request):
    configs = AutoTriggerConfig.objects.all().select_related('department', 'time_setting').order_by('-created_at')
    
    search = request.GET.get('search', '')
    if search:
        configs = configs.filter(department__name__icontains=search)
    
    data = []
    for c in configs:
        last_history = c.autotriggerhistory_set.order_by('-created_at').first()
        ts = c.time_setting
        
        rush_start = ts.rush_start.strftime('%I:%M %p') if ts and ts.rush_start else '10:00 AM'
        rush_end = ts.rush_end.strftime('%I:%M %p') if ts and ts.rush_end else '02:00 PM'
        rush_pct = ts.rush_percentage if ts else 75
        interval_mins = ts.processing_interval if ts else 30
        
        progress_pct = 60.8
        if last_history:
            if last_history.max_entries > 0:
                progress_pct = round((last_history.processed_entries / last_history.max_entries) * 100, 1)
            
        data.append({
            'id': c.id,
            'department': c.department.name if c.department else '',
            'department_id': c.department_id,
            'time_setting_id': c.time_setting_id,
            'time_setting_name': c.time_setting.name if c.time_setting else 'Default Lab Schedule',
            'from_date': c.from_date.strftime('%Y-%m-%d') if c.from_date else '',
            'to_date': c.to_date.strftime('%Y-%m-%d') if c.to_date else '',
            'min_entries': c.min_entries,
            'max_entries': c.max_entries,
            'trigger_start_time': c.trigger_start_time.strftime('%H:%M:%S') if c.trigger_start_time else '10:00:00',
            'day_start_time': c.day_start_time.strftime('%H:%M:%S') if c.day_start_time else '04:00:00',
            'day_end_time': c.day_end_time.strftime('%H:%M:%S') if c.day_end_time else '03:59:00',
            'description': c.description or '',
            'is_active': c.is_active,
            'created_by': c.created_by.username if c.created_by else 'System',
            'created_at': c.created_at.strftime('%Y-%m-%d %H:%M'),
            'last_triggered': last_history.started_at.strftime('%d %b %Y %I:%M %p') if last_history and last_history.started_at else '28 Aug 2026 09:40 AM',
            'last_history_id': last_history.id if last_history else None,
            'rush_period': f"{rush_start} - {rush_end}",
            'rush_percentage': rush_pct,
            'interval': f"{int(round(interval_mins * 60))} sec" if interval_mins < 1 else (f"{int(interval_mins)} min" if interval_mins % 1 == 0 else f"{round(interval_mins, 1)} min"),
            'progress': progress_pct
        })
    return JsonResponse({'success': True, 'data': data})

@csrf_exempt
def api_execute_auto_trigger(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            config_id = data.get('config_id')
            if not config_id:
                return JsonResponse({'success': False, 'message': 'Config ID is required.'})
            
            config = AutoTriggerConfig.objects.get(id=config_id)
            ts = config.time_setting
            
            existing_running = AutoTriggerHistory.objects.filter(config=config, status__in=['Queued', 'Processing']).first()
            if existing_running:
                return JsonResponse({
                    'success': False, 
                    'message': f'This configuration is already running (Run ID: {existing_running.run_id}).',
                    'run_id': existing_running.run_id,
                    'history_id': existing_running.id,
                    'status': existing_running.status
                })
            
            schedule = generate_schedule(
                max_entries=config.max_entries,
                min_entries=config.min_entries,
                day_start=config.day_start_time,
                day_end=config.day_end_time,
                rush_start=ts.rush_start,
                rush_end=ts.rush_end,
                rush_percent=ts.rush_percentage,
                interval_mins=ts.processing_interval
            )
            
            history = AutoTriggerHistory.objects.create(
                config=config,
                department=config.department,
                from_date=config.from_date,
                to_date=config.to_date,
                min_entries=config.min_entries,
                max_entries=config.max_entries,
                trigger_start_time=config.trigger_start_time,
                day_start_time=config.day_start_time,
                day_end_time=config.day_end_time,
                rush_percentage=ts.rush_percentage,
                schedule_data=schedule,
                status='Queued',
                current_stage='STAGE 1 — PATIENT CREATION',
                triggered_by=request.user if request.user.is_authenticated else None
            )
            thread = threading.Thread(target=process_auto_trigger, args=(history.id,))
            thread.daemon = True
            thread.start()
            
            return JsonResponse({
                'success': True, 
                'message': f'Automation started successfully. Run ID: {history.run_id}',
                'run_id': history.run_id,
                'history_id': history.id,
                'status': 'PROCESSING',
                'current_stage': 'STAGE 1 — PATIENT CREATION'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

def api_get_auto_trigger_history(request):
    try:
        history_list = AutoTriggerHistory.objects.all().select_related(
            'department', 'config', 'config__time_setting', 'triggered_by'
        ).order_by('-created_at')
        
        department_id = request.GET.get('department_id', '').strip()
        if department_id and department_id.isdigit():
            history_list = history_list.filter(department_id=int(department_id))
            
        status = request.GET.get('status', '').strip()
        if status and status.lower() not in ('', 'all'):
            # Map display statuses to DB statuses
            status_map = {
                'running': ['Processing', 'Queued'],
                'processing': ['Processing', 'Queued'],
                'completed': ['Completed'],
                'partially completed': ['Partially Completed'],
                'failed': ['Failed'],
            }
            db_statuses = status_map.get(status.lower(), [status])
            history_list = history_list.filter(status__in=db_statuses)
            
        from_date, to_date = get_default_date_range(request, 'from_date', 'to_date')
        
        if from_date:
            try:
                history_list = history_list.filter(created_at__date__gte=from_date)
            except Exception:
                pass
        if to_date:
            try:
                history_list = history_list.filter(created_at__date__lte=to_date)
            except Exception:
                pass
            
        search = request.GET.get('search', '').strip()
        if search:
            # Support full AT-YYYYMMDD-NNN format or partial numeric ID
            import re as _re
            # Try to extract the trailing numeric ID from the Run ID format
            m = _re.search(r'(\d+)$', search)
            if m:
                try:
                    history_list = history_list.filter(id=int(m.group(1)))
                except Exception:
                    history_list = history_list.none()
            else:
                history_list = history_list.none()
            
        data = []
        for h in history_list:
            try:
                ts = h.config.time_setting if h.config else None
                time_profile = ts.name if ts else (h.config.description if h.config and h.config.description else 'Custom Schedule')
                
                # Determine display status: Queued shows as 'Processing' in the UI
                display_status = h.status
                if display_status == 'Queued':
                    display_status = 'Processing'
                
                # Last updated: prefer last_processed_at > completed_at > started_at > created_at
                last_updated_dt = getattr(h, 'last_processed_at', None) or h.completed_at or h.started_at or h.created_at
                last_updated_str = last_updated_dt.strftime('%d %b %Y %I:%M %p') if last_updated_dt else ''
                
                data.append({
                    'id': h.id,
                    'trigger_id': f"AT-{h.created_at.strftime('%Y%m%d')}-{h.id:03d}",
                    'department': h.department.name if h.department else '',
                    'department_id': h.department_id,
                    'from_date': h.from_date.strftime('%d %b %Y') if h.from_date else '',
                    'to_date': h.to_date.strftime('%d %b %Y') if h.to_date else '',
                    'min_entries': h.min_entries,
                    'max_entries': h.max_entries,
                    'time_profile': time_profile,
                    'day_start_time': h.day_start_time.strftime('%I:%M %p') if h.day_start_time else '04:00 AM',
                    'day_end_time': (h.day_end_time.strftime('%I:%M %p') + ' (Next Day)') if h.day_end_time else '03:59 AM (Next Day)',
                    'trigger_start_time': h.trigger_start_time.strftime('%I:%M %p') if h.trigger_start_time else '',
                    'processed_entries': h.processed_entries,
                    'successful': h.successful,
                    'failed': h.failed,
                    'status': display_status,
                    'raw_status': h.status,
                    'triggered_by': h.triggered_by.username if h.triggered_by else 'System',
                    'started_at': h.started_at.strftime('%d %b %Y %I:%M %p') if h.started_at else h.created_at.strftime('%d %b %Y %I:%M %p'),
                    'completed_at': h.completed_at.strftime('%d %b %Y %I:%M %p') if h.completed_at else '',
                    'last_updated': last_updated_str,
                    'progress': int((h.processed_entries / h.max_entries) * 100) if h.max_entries > 0 else 0,
                })
            except Exception as row_err:
                # Include errored row with minimal info so we don't drop it silently
                data.append({
                    'id': h.id,
                    'trigger_id': f"AT-{h.created_at.strftime('%Y%m%d') if h.created_at else '00000000'}-{h.id:03d}",
                    'department': h.department.name if h.department else '',
                    'department_id': h.department_id,
                    'from_date': '', 'to_date': '',
                    'min_entries': h.min_entries, 'max_entries': h.max_entries,
                    'time_profile': 'Error loading profile',
                    'day_start_time': '', 'day_end_time': '',
                    'trigger_start_time': '',
                    'processed_entries': h.processed_entries,
                    'successful': h.successful, 'failed': h.failed,
                    'status': h.status if h.status != 'Queued' else 'Processing',
                    'raw_status': h.status,
                    'triggered_by': h.triggered_by.username if h.triggered_by else 'System',
                    'started_at': h.created_at.strftime('%d %b %Y %I:%M %p') if h.created_at else '',
                    'completed_at': '',
                    'last_updated': '',
                    'progress': int((h.processed_entries / h.max_entries) * 100) if h.max_entries > 0 else 0,
                })
                
        return JsonResponse({'success': True, 'data': data, 'total': len(data)})
    except Exception as e:
        import traceback
        return JsonResponse({'success': False, 'message': f'History API error: {str(e)}', 'trace': traceback.format_exc()})

def api_get_auto_trigger_history_detail(request, pk):
    try:
        h = AutoTriggerHistory.objects.get(id=pk)
        
        time_profile = h.config.time_setting.name if h.config and h.config.time_setting else 'Custom Schedule'
        
        details = {
            'id': h.id,
            'trigger_id': f"AT-{h.created_at.strftime('%Y%m%d')}-{h.id:03d}",
            'department': h.department.name if h.department else '',
            'date_range': f"{h.from_date.strftime('%d %b %Y')} \u2192 {h.to_date.strftime('%d %b %Y')}",
            'time_profile': time_profile,
            'min_entries': h.min_entries,
            'max_entries': h.max_entries,
            'processed_entries': h.processed_entries,
            'successful': h.successful,
            'failed': h.failed,
            'status': h.status,
            'progress': int((h.processed_entries / h.max_entries) * 100) if h.max_entries > 0 else 0,
            'triggered_by': h.triggered_by.username if h.triggered_by else 'System',
            'started_at': h.started_at.strftime('%d %b %Y %I:%M %p') if h.started_at else h.created_at.strftime('%d %b %Y %I:%M %p'),
            'completed_at': h.completed_at.strftime('%d %b %Y %I:%M %p') if h.completed_at else '',
            'schedule_data': h.schedule_data
        }
        
        from datetime import datetime, timedelta
        rush_start = h.trigger_start_time.strftime('%I:%M %p')
        rush_end = h.config.time_setting.rush_end.strftime('%I:%M %p') if (h.config and h.config.time_setting) else '02:00 PM'
        
        rush_target = round(h.max_entries * (h.rush_percentage / 100.0))
        details['window_str'] = f"{rush_start} \u2192 {rush_end}"
        details['window_target'] = rush_target
        details['window_completed'] = min(h.successful, rush_target)
        details['window_remaining'] = max(0, rush_target - h.successful)
        
        rush_dt1 = datetime.strptime(rush_start, '%I:%M %p')
        rush_dt2 = datetime.strptime(rush_end, '%I:%M %p')
        if rush_dt2 < rush_dt1: rush_dt2 += timedelta(days=1)
        hours = (rush_dt2 - rush_dt1).total_seconds() / 3600.0
        details['req_rate'] = round(rush_target / hours, 1) if hours > 0 else 0
        
        if h.started_at:
            elapsed_hours = (timezone.now() - h.started_at).total_seconds() / 3600.0
            details['cur_rate'] = round(h.successful / elapsed_hours, 1) if elapsed_hours > 0.05 else details['req_rate']
        else:
            details['cur_rate'] = 0
            
        details['rate_status'] = 'On Schedule' if details['cur_rate'] >= details['req_rate'] else 'Falling Behind'
        if h.status in ['Completed', 'Partially Completed']:
            details['rate_status'] = 'Completed'
        
        logs = h.logs.all().select_related('source_patient', 'new_patient').order_by('entry_no')
        log_data = []
        for l in logs:
            src_id = l.source_patient.patient_id if l.source_patient else 'N/A'
            src_name = l.source_patient.name if l.source_patient else ''
            new_id = l.new_patient.patient_id if l.new_patient else 'N/A'
            new_name = l.new_patient.name if l.new_patient else ''
            op_num = l.new_patient.op_number if l.new_patient and l.new_patient.op_number else '-'
            
            sched_str = ''
            if l.scheduled_at:
                sched_str = l.scheduled_at.strftime('%d %b %Y %I:%M %p')
            elif l.new_patient and l.new_patient.automation_scheduled_at:
                sched_str = l.new_patient.automation_scheduled_at.strftime('%d %b %Y %I:%M %p')
            elif l.created_at:
                sched_str = l.created_at.strftime('%d %b %Y %I:%M %p')

            log_data.append({
                'entry_no': l.entry_no,
                'entry_date': l.entry_date.strftime('%d %b %Y') if l.entry_date else '',
                'department': l.department,
                'stage': l.stage,
                'source_patient_id': src_id,
                'source_patient_name': src_name,
                'new_patient_id': new_id,
                'new_patient_name': new_name,
                'op_number': op_num,
                'gender': l.new_patient.gender if l.new_patient else '-',
                'dob': l.new_patient.dob.strftime('%d/%m/%Y') if l.new_patient and l.new_patient.dob else '-',
                'age': l.new_patient.age_years if l.new_patient else '-',
                'guardian_title': l.new_patient.guardian_title if l.new_patient else '-',
                'guardian_name': l.new_patient.guardian_name if l.new_patient else '-',
                'guardian_rel': l.new_patient.guardian_relationship if l.new_patient else '-',
                'address': f"{l.new_patient.street}, {l.new_patient.city} - {l.new_patient.pincode}" if l.new_patient else '-',
                'phone': l.new_patient.mobile_no if l.new_patient else '-',
                'status': 'Created' if l.status == 'Success' else l.status,
                'message': l.message,
                'scheduled_at': sched_str,
                'created_at': l.created_at.strftime('%d %b %Y %I:%M:%S %p') if l.created_at else '',
                'created_source': l.new_patient.created_source if l.new_patient else 'D',
                'created_source_display': l.new_patient.get_created_source_display() if l.new_patient else 'D - Auto Trigger'
            })
            
        return JsonResponse({'success': True, 'details': details, 'logs': log_data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@csrf_exempt
def api_delete_auto_trigger_config(request, pk):
    if request.method == 'POST':
        try:
            config = AutoTriggerConfig.objects.get(id=pk)
            config.delete()
            return JsonResponse({'success': True, 'message': 'Configuration deleted successfully.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

@csrf_exempt
def api_retry_failed_entries(request, pk):
    if request.method == 'POST':
        try:
            history = AutoTriggerHistory.objects.get(id=pk)
            if history.failed == 0:
                return JsonResponse({'success': False, 'message': 'No failed entries to retry.'})
                
            history.status = 'Processing'
            history.save()
            thread = threading.Thread(target=process_auto_trigger, args=(history.id, True))
            thread.daemon = True
            thread.start()
            return JsonResponse({'success': True, 'message': 'Retrying failed entries. Processing started.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

# Time Settings API
@csrf_exempt
def api_save_time_setting(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            ts_id = data.get('id')
            is_active = data.get('is_active', True)
            
            if ts_id:
                ts = AutoTriggerTimeSetting.objects.get(id=ts_id)
                ts.name = data['name']
                ts.lab_day_start = data['lab_day_start']
                ts.lab_day_end = data['lab_day_end']
                ts.rush_start = data['rush_start']
                ts.rush_end = data['rush_end']
                ts.rush_percentage = int(data['rush_percentage'])
                ts.processing_interval = float(data['processing_interval'])
                ts.is_active = is_active
                ts.save()
            else:
                ts = AutoTriggerTimeSetting.objects.create(
                    name=data['name'],
                    lab_day_start=data['lab_day_start'],
                    lab_day_end=data['lab_day_end'],
                    rush_start=data['rush_start'],
                    rush_end=data['rush_end'],
                    rush_percentage=int(data['rush_percentage']),
                    processing_interval=float(data['processing_interval']),
                    is_active=is_active
                )
            return JsonResponse({'success': True, 'message': 'Time Profile saved successfully.', 'id': ts.id})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

def api_get_time_settings(request):
    settings = AutoTriggerTimeSetting.objects.all().order_by('-is_active', '-created_at')
    data = []
    for s in settings:
        data.append({
            'id': s.id,
            'name': s.name,
            'lab_day_start': s.lab_day_start.strftime('%H:%M') if s.lab_day_start else '04:00',
            'lab_day_end': s.lab_day_end.strftime('%H:%M') if s.lab_day_end else '03:59',
            'rush_start': s.rush_start.strftime('%H:%M') if s.rush_start else '10:00',
            'rush_end': s.rush_end.strftime('%H:%M') if s.rush_end else '14:00',
            'rush_percentage': s.rush_percentage,
            'processing_interval': s.processing_interval,
            'is_active': s.is_active
        })
    return JsonResponse({'success': True, 'data': data})

@csrf_exempt
def api_delete_time_setting(request, pk):
    if request.method == 'POST':
        try:
            ts = AutoTriggerTimeSetting.objects.get(id=pk)
            ts.delete()
            return JsonResponse({'success': True, 'message': 'Time Setting deleted.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid method.'})

def api_get_auto_trigger_status(request, pk):
    try:
        if isinstance(pk, str) and str(pk).startswith('AT-'):
            try:
                parts = str(pk).split('-')
                hid = int(parts[-1])
                h = AutoTriggerHistory.objects.get(id=hid)
            except Exception:
                h = AutoTriggerHistory.objects.get(id=int(str(pk).replace('AT-', '')))
        else:
            h = AutoTriggerHistory.objects.get(id=int(pk))

        last_log = h.logs.order_by('-id').first()
        cur_sched_time = '-'
        next_sched_time = '-'
        if last_log and last_log.scheduled_at:
            cur_sched_time = timezone.localtime(last_log.scheduled_at).strftime('%I:%M %p')
            interval = float(h.config.time_setting.processing_interval) if (h.config and h.config.time_setting and h.config.time_setting.processing_interval) else 1.0
            next_sched = last_log.scheduled_at + timedelta(minutes=interval)
            next_sched_time = timezone.localtime(next_sched).strftime('%I:%M %p')

        data = {
            'success': True,
            'run_id': h.run_id,
            'history_id': h.id,
            'department': h.department.name if h.department else 'GENERAL',
            'status': h.status,
            'current_stage': h.current_stage or 'STAGE 1 — PATIENT CREATION',
            'total_created': h.successful,
            'total_failed': h.failed,
            'total_processed': h.processed_entries,
            'total_remaining': max(0, h.max_entries - h.processed_entries),
            'min_entries': h.min_entries,
            'max_entries': h.max_entries,
            'progress_percentage': round((h.processed_entries / h.max_entries) * 100, 1) if h.max_entries > 0 else 0,
            'started_at': h.started_at.strftime('%d %b %Y %I:%M %p') if h.started_at else h.created_at.strftime('%d %b %Y %I:%M %p'),
            'last_updated_at': h.last_processed_at.strftime('%I:%M:%S %p') if h.last_processed_at else timezone.now().strftime('%I:%M:%S %p'),
            'completed_at': h.completed_at.strftime('%d %b %Y %I:%M %p') if h.completed_at else '',
            'current_scheduled_time': cur_sched_time,
            'next_scheduled_time': next_sched_time,
            'error_message': h.error_message or ''
        }
        return JsonResponse(data)
    except AutoTriggerHistory.DoesNotExist:
        return JsonResponse({'success': False, 'message': f'Auto Trigger Run {pk} not found.'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

def api_get_active_auto_trigger_run(request):
    try:
        active_run = AutoTriggerHistory.objects.filter(
            status__in=['Queued', 'Processing']
        ).select_related('department', 'config').order_by('-id').first()

        if not active_run:
            return JsonResponse({'success': True, 'has_active': False, 'data': None})

        last_log = active_run.logs.order_by('-id').first()
        cur_sched_time = '-'
        next_sched_time = '-'
        if last_log and last_log.scheduled_at:
            cur_sched_time = timezone.localtime(last_log.scheduled_at).strftime('%I:%M %p')
            interval = float(active_run.config.time_setting.processing_interval) if (active_run.config and active_run.config.time_setting and active_run.config.time_setting.processing_interval) else 1.0
            next_sched = last_log.scheduled_at + timedelta(minutes=interval)
            next_sched_time = timezone.localtime(next_sched).strftime('%I:%M %p')

        data = {
            'has_active': True,
            'run_id': active_run.run_id,
            'history_id': active_run.id,
            'department': active_run.department.name if active_run.department else 'GENERAL',
            'status': active_run.status,
            'current_stage': active_run.current_stage or 'STAGE 1 — PATIENT CREATION',
            'total_created': active_run.successful,
            'total_failed': active_run.failed,
            'total_processed': active_run.processed_entries,
            'total_remaining': max(0, active_run.max_entries - active_run.processed_entries),
            'min_entries': active_run.min_entries,
            'max_entries': active_run.max_entries,
            'progress_percentage': round((active_run.processed_entries / active_run.max_entries) * 100, 1) if active_run.max_entries > 0 else 0,
            'started_at': active_run.started_at.strftime('%d %b %Y %I:%M %p') if active_run.started_at else active_run.created_at.strftime('%d %b %Y %I:%M %p'),
            'last_updated_at': active_run.last_processed_at.strftime('%I:%M:%S %p') if active_run.last_processed_at else timezone.now().strftime('%I:%M:%S %p'),
            'current_scheduled_time': cur_sched_time,
            'next_scheduled_time': next_sched_time,
            'error_message': active_run.error_message or ''
        }
        return JsonResponse({'success': True, 'has_active': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


class AutoTriggerMonthlyCreateView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.configuration.view'
    template_name = 'lab/auto_trigger/monthly_trigger_create.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.patients.models import Department
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        return context

class AutoTriggerMonthlyView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.configuration.view'
    template_name = 'lab/auto_trigger/monthly_trigger.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.patients.models import Department
        from apps.lab.models import MonthlyTriggerPlan
        context['departments'] = Department.objects.filter(is_active=True).order_by('name')
        context['plans'] = MonthlyTriggerPlan.objects.all().order_by('-id')
        return context


class AutoTriggerMonthlyCensusView(LoginRequiredMixin, GranularPermissionRequiredMixin, TemplateView):
    permission_required = 'auto_trigger.history.view'
    template_name = 'lab/auto_trigger/monthly_census.html'

import json
import datetime
from django.utils import timezone
from apps.lab.models import MonthlyTriggerPlan, MonthlyTriggerDailyTarget

def api_get_monthly_trigger(request):
    dept_id = request.GET.get('department_id')
    month = request.GET.get('month')
    year = request.GET.get('year')
    if not all([dept_id, month, year]):
        return JsonResponse({'success': False, 'message': 'Missing parameters'})
    
    plan = MonthlyTriggerPlan.objects.filter(department_id=dept_id, month=month, year=year).first()
    if not plan:
        return JsonResponse({'success': True, 'exists': False})
        
    targets = MonthlyTriggerDailyTarget.objects.filter(plan=plan).order_by('target_date')
    daily_targets = []
    for t in targets:
        daily_targets.append({
            'date': t.target_date.strftime('%Y-%m-%d'),
            'min': t.min_entries,
            'max': t.max_entries,
            'status': t.status
        })
        
    data = {
        'source_from_date': plan.source_from_date.strftime('%Y-%m-%d'),
        'source_to_date': plan.source_to_date.strftime('%Y-%m-%d'),
        'trigger_start_time': plan.trigger_start_time.strftime('%H:%M'),
        'trigger_stop_time': plan.trigger_stop_time.strftime('%H:%M'),
        'interval_mins': plan.interval_mins,
        'status': plan.status,
        'is_saved': plan.is_saved,
        'daily_targets': daily_targets
    }
    return JsonResponse({'success': True, 'exists': True, 'data': data})

def api_save_monthly_trigger(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid method'})
    try:
        data = json.loads(request.body)
        dept_id = data.get('department_id')
        month = data.get('month')
        year = data.get('year')
        
        plan, created = MonthlyTriggerPlan.objects.update_or_create(
            department_id=dept_id, month=month, year=year,
            defaults={
                'source_from_date': data.get('source_from_date'),
                'source_to_date': data.get('source_to_date'),
                'trigger_start_time': data.get('trigger_start_time'),
                'trigger_stop_time': data.get('trigger_stop_time'),
                'interval_mins': data.get('interval_mins'),
                'status': data.get('status', 'Active'),
                'is_saved': True,
                'created_by': request.user if request.user.is_authenticated else None
            }
        )
        
        daily_targets = data.get('daily_targets', [])
        for target in daily_targets:
            MonthlyTriggerDailyTarget.objects.update_or_create(
                plan=plan,
                target_date=target.get('date'),
                defaults={
                    'min_entries': target.get('min'),
                    'max_entries': target.get('max'),
                    'status': target.get('status', 'Active')
                }
            )
            
        return JsonResponse({'success': True, 'message': 'Configuration saved successfully'})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

def api_stop_monthly_automation(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid method'})
    if not (request.user.is_superuser or (hasattr(request.user, 'role') and request.user.role == 'Developer')):
        return JsonResponse({'success': False, 'message': 'Permission denied. Developer only.'})
    try:
        plans = MonthlyTriggerPlan.objects.filter(status='Active')
        for plan in plans:
            plan.status = 'Stopped'
            plan.save()
            
        targets = MonthlyTriggerDailyTarget.objects.filter(status__in=['Active', 'Not Started', 'Running'])
        for target in targets:
            target.status = 'Stopped'
            target.stopped_by = request.user
            target.stopped_at = timezone.now()
            target.save()
            
        from apps.lab.models import MonthlyTriggerExecutionLog
        # Create a generic log or attach to plans
        if plans.exists():
            MonthlyTriggerExecutionLog.objects.create(
                plan=plans.first(), # Attach to first plan just to satisfy FK
                action='GLOBAL AUTOMATION STOP',
                status='STOPPED',
                stopped_by=request.user,
                reason='Manual emergency stop'
            )
            
        return JsonResponse({'success': True, 'message': 'Monthly automation stopped successfully.'})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

def api_stop_daily_execution(request):
    if request.method != 'POST': return JsonResponse({'success': False})
    if not (request.user.is_superuser or (hasattr(request.user, 'role') and request.user.role == 'Developer')):
        return JsonResponse({'success': False, 'message': 'Developer only.'})
    try:
        import json
        data = json.loads(request.body)
        plan_id = data.get('plan_id')
        target_date = data.get('target_date')
        if not plan_id or not target_date:
            return JsonResponse({'success': False, 'message': 'Missing plan ID or target date'})
            
        targets = MonthlyTriggerDailyTarget.objects.filter(plan_id=plan_id, target_date=target_date)
        if not targets.exists():
            return JsonResponse({'success': False, 'message': 'No targets found for this date'})
            
        plan = targets.first().plan
        
        # Mark targets as stopped if they are running or not started
        targets.filter(status__in=['Not Started', 'Running']).update(
            status='Stopped',
            stopped_by=request.user,
            stopped_at=timezone.now()
        )
        
        total_created = sum(t.created_count for t in targets)
        total_dup = sum(t.duplicates_count for t in targets)
        total_fail = sum(t.failed_count for t in targets)
        
        from apps.lab.models import MonthlyTriggerExecutionLog
        MonthlyTriggerExecutionLog.objects.create(
            plan=plan,
            target_date=target_date,
            action='AUTOMATION STOPPED',
            status='Stopped',
            created_count=total_created,
            duplicates_count=total_dup,
            failed_count=total_fail,
            stopped_by=request.user,
            reason='Manual stop for the day'
        )
        return JsonResponse({'success': True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

from apps.lab.models import MonthlyTriggerExecutionLog
from django.db.models import Sum
from datetime import datetime

def api_get_monthly_history(request):
    dept_id = request.GET.get('department_id')
    plan_id = request.GET.get('plan_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    status = request.GET.get('status')
    
    targets = MonthlyTriggerDailyTarget.objects.all().select_related('plan', 'department')
    logs = MonthlyTriggerExecutionLog.objects.all().select_related('plan', 'stopped_by')
    
    if dept_id:
        targets = targets.filter(department_id=dept_id)
        # logs are per plan/date mostly, but we can filter targets
        
    if plan_id:
        targets = targets.filter(plan_id=plan_id)
        logs = logs.filter(plan_id=plan_id)
    else:
        # Default to the most recent running/active/scheduled plan to avoid duplicate rows
        from django.utils import timezone
        today = timezone.localtime().date()
        date_filter = today
        if from_date:
            date_filter = datetime.strptime(from_date, '%Y-%m-%d').date()
            
        latest_plan = MonthlyTriggerPlan.objects.filter(
            is_saved=True,
            automation_start_date__lte=date_filter,
            automation_end_date__gte=date_filter
        ).order_by('-id').first()
        
        if latest_plan:
            targets = targets.filter(plan_id=latest_plan.id)
        else:
            targets = targets.none()
        
    if from_date:
        targets = targets.filter(target_date__gte=from_date)
        logs = logs.filter(target_date__gte=from_date)
        
    if to_date:
        targets = targets.filter(target_date__lte=to_date)
        logs = logs.filter(target_date__lte=to_date)
        
    if status and status != 'All':
        targets = targets.filter(status=status)
        logs = logs.filter(status=status)
        
    targets = targets.order_by('target_date', 'department__name')
    logs = logs.order_by('-created_at')[:50] # Last 50 logs
    
    # 1. Daily Targets (for the "Daily Trigger Execution" table)
    daily_targets_data = []
    for t in targets:
        progress = 0
        if t.max_entries > 0:
            progress = (t.created_count / t.max_entries) * 100
            
        rem_min = max(0, t.min_entries - t.created_count)
        rem_max = max(0, t.max_entries - t.created_count)
        
        # Display status logic
        display_status = t.status
        from django.utils import timezone
        today = timezone.localtime().date()
        
        if t.target_date == today:
            if t.plan.status == 'Running' and t.status in ['Not Started', 'Running']:
                display_status = 'Active in Queue' if t.status == 'Not Started' else 'Running'
        elif t.target_date > today and display_status not in ['Completed', 'Stopped']:
            display_status = 'Not Started'
            
        daily_targets_data.append({
            'id': t.id,
            'date': t.target_date.strftime('%d-%b-%Y'),
            'day': t.target_date.strftime('%a'),
            'department': t.department.name,
            'min': t.min_entries,
            'max': t.max_entries,
            'new_op_target': t.new_op_target,
            'review_target': t.review_target,
            'new_op_created': t.new_op_created,
            'review_created': t.review_created,
            'created': t.created_count,
            'duplicates': t.duplicates_count,
            'failed': t.failed_count,
            'remaining_min': rem_min,
            'remaining_max': rem_max,
            'progress': round(progress, 1),
            'status': display_status
        })
        
    # 2. Daily Execution History (Grouped by Plan and Date)
    # Get unique (plan_id, target_date) combinations from targets
    logs_data = []
    history_grouped = targets.values('plan_id', 'target_date').distinct().order_by('-target_date')[:100]
    
    for group in history_grouped:
        p_id = group['plan_id']
        t_date = group['target_date']
        
        # Get all targets for this plan and date to aggregate stats
        day_targets = targets.filter(plan_id=p_id, target_date=t_date)
        if not day_targets.exists():
            continue
            
        plan = day_targets.first().plan
        plan_desc = plan.description if plan.description else f"Plan #{plan.id}"
        
        total_min = sum(dt.min_entries for dt in day_targets)
        total_max = sum(dt.max_entries for dt in day_targets)
        total_new_op_target = sum(dt.new_op_target for dt in day_targets)
        total_review_target = sum(dt.review_target for dt in day_targets)
        total_new_op_created = sum(dt.new_op_created for dt in day_targets)
        total_review_completed = sum(dt.review_created for dt in day_targets)
        total_created = sum(dt.created_count for dt in day_targets)
        total_dup = sum(dt.duplicates_count for dt in day_targets)
        total_failed = sum(dt.failed_count for dt in day_targets)
        
        # Determine overall status for the day based on targets or plan
        # If plan is running and date is today, it's running. Otherwise determine from targets.
        from django.utils import timezone
        today = timezone.localtime().date()
        
        if plan.status == 'Running' and t_date == today:
            day_status = 'Running'
        else:
            # Check targets
            if all(dt.status == 'Completed' for dt in day_targets):
                day_status = 'Completed'
            elif any(dt.status == 'Stopped' for dt in day_targets):
                day_status = 'Stopped'
            elif any(dt.status == 'Failed' for dt in day_targets):
                day_status = 'Failed'
            else:
                day_status = 'Completed' if t_date < today else 'Scheduled'
                
        logs_data.append({
            'plan_id': p_id,
            'date_raw': t_date.strftime('%Y-%m-%d'),
            'date': t_date.strftime('%d-%b-%Y'),
            'plan': plan_desc,
            'department': 'Global Daily Execution',
            'start_time': plan.trigger_start_time.strftime('%I:%M %p') if plan.trigger_start_time else '-',
            'stop_time': plan.trigger_stop_time.strftime('%I:%M %p') if plan.trigger_stop_time else '-',
            'min': total_min,
            'max': total_max,
            'new_op_target': total_new_op_target,
            'review_target': total_review_target,
            'new_op_created': total_new_op_created,
            'review_completed': total_review_completed,
            'created': total_created,
            'duplicates': total_dup,
            'failed': total_failed,
            'status': day_status
        })
    # Plan Summary (only if 1 plan is filtered, or aggregate)
    summary = None
    if plan_id:
        plan = MonthlyTriggerPlan.objects.filter(id=plan_id).first()
        if plan:
            aggr = MonthlyTriggerDailyTarget.objects.filter(plan=plan).aggregate(
                total_created=Sum('created_count'),
                total_duplicates=Sum('duplicates_count'),
                total_failed=Sum('failed_count')
            )
            summary = {
                'department': 'Multiple Departments',
                'source_from': plan.source_from_date.strftime('%d-%b-%Y') if plan.source_from_date else '—',
                'source_to': plan.source_to_date.strftime('%d-%b-%Y') if plan.source_to_date else '—',
                'start_date': plan.automation_start_date.strftime('%d-%b-%Y') if plan.automation_start_date else '—',
                'end_date': plan.automation_end_date.strftime('%d-%b-%Y') if plan.automation_end_date else '—',
                'start_time': plan.trigger_start_time.strftime('%I:%M %p') if plan.trigger_start_time else '—',
                'stop_time': plan.trigger_stop_time.strftime('%I:%M %p') if plan.trigger_stop_time else '—',
                'interval': f"{plan.interval_mins} mins",
                'status': plan.status,
                'total_days': MonthlyTriggerDailyTarget.objects.filter(plan=plan).values('target_date').distinct().count(),
                'total_created': aggr['total_created'] or 0,
                'total_duplicates': aggr['total_duplicates'] or 0,
                'total_failed': aggr['total_failed'] or 0
            }
            
    return JsonResponse({
        'success': True, 
        'daily_targets': daily_targets_data,
        'logs': logs_data,
        'summary': summary
    })

def api_save_monthly_trigger_v2(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid method'})
    try:
        import json
        data = json.loads(request.body)
        
        plan = MonthlyTriggerPlan.objects.create(
            automation_start_date=data.get('automation_start_date'),
            automation_end_date=data.get('automation_end_date'),
            source_from_date=data.get('source_from_date'),
            source_to_date=data.get('source_to_date'),
            review_source_month_year=data.get('review_source_month_year'),
            trigger_start_time=data.get('trigger_start_time'),
            trigger_stop_time=data.get('trigger_stop_time'),
            interval_mins=data.get('interval_mins', 1.5),
            description=data.get('description'),
            status='Scheduled',
            is_saved=True,
            created_by=request.user if request.user.is_authenticated else None
        )
        
        daily_targets = data.get('daily_targets', [])
        for target in daily_targets:
            min_entries = int(target.get('min', 0))
            new_op_target = int(min_entries * 0.8)
            review_target = min_entries - new_op_target
            MonthlyTriggerDailyTarget.objects.create(
                plan=plan,
                department_id=target.get('department_id'),
                target_date=target.get('date'),
                min_entries=min_entries,
                max_entries=int(target.get('max', 0)),
                new_op_target=new_op_target,
                review_target=review_target,
                status='Not Started'
            )
            
        MonthlyTriggerExecutionLog.objects.create(
            plan=plan,
            action='PLAN SAVED',
            status='Scheduled',
            stopped_by=request.user if request.user.is_authenticated else None,
            reason='Plan created and targets scheduled'
        )
            
        return JsonResponse({'success': True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

def api_get_daily_created_patients(request):
    from apps.lab.models import MonthlyTriggerGeneratedPatient, MonthlyTriggerDailyTarget
    from django.utils import timezone
    from datetime import datetime
    try:
        plan_id = request.GET.get('plan_id')
        date_str = request.GET.get('date')
        
        if not plan_id or not date_str:
            return JsonResponse({'success': False, 'message': 'Missing plan or date'})
            
        date_obj = datetime.strptime(date_str, '%Y-%m-%d').date()
        
        targets = MonthlyTriggerDailyTarget.objects.filter(plan_id=plan_id, target_date=date_obj).select_related('department', 'plan')
        if not targets.exists():
            return JsonResponse({'success': False, 'message': 'No data found for this execution'})
            
        plan = targets.first().plan
        
        # 1. Global Details
        total_created = sum(t.created_count for t in targets)
        total_dup = sum(t.duplicates_count for t in targets)
        total_fail = sum(t.failed_count for t in targets)
        
        # Determine status
        today = timezone.localtime().date()
        if plan.status == 'Running' and date_obj == today:
            day_status = 'Running'
        else:
            if all(t.status == 'Completed' for t in targets): day_status = 'Completed'
            elif any(t.status == 'Stopped' for t in targets): day_status = 'Stopped'
            elif any(t.status == 'Failed' for t in targets): day_status = 'Failed'
            else: day_status = 'Completed' if date_obj < today else 'Scheduled'
            
        details = {
            'date': date_obj.strftime('%d-%b-%Y'),
            'plan': plan.description if plan.description else f"Plan #{plan.id}",
            'status': day_status,
            'total_created': total_created,
            'total_new_op': sum(t.new_op_created for t in targets),
            'total_review': sum(t.review_created for t in targets),
            'total_duplicates': total_dup,
            'total_failed': total_fail
        }
        
        # 2. Department Statistics
        dept_stats = []
        for t in targets:
            progress = (t.created_count / t.max_entries * 100) if t.max_entries > 0 else 0
            dept_stats.append({
                'department': t.department.name,
                'created': t.created_count,
                'new_op_created': t.new_op_created,
                'review_created': t.review_created,
                'min': t.min_entries,
                'max': t.max_entries,
                'progress': round(progress, 1)
            })
            
        # 3. Created Patients
        generated = MonthlyTriggerGeneratedPatient.objects.filter(
            target__in=targets, status='Success'
        ).select_related('patient', 'target__department').order_by('created_at')
        
        patients_data = []
        for g in generated:
            p = g.patient
            patients_data.append({
                'id': p.id,
                'patient_id': p.patient_id,
                'op_number': p.op_number,
                'name': p.name,
                'title': p.title,
                'gender': p.gender,
                'age': p.age_years,
                'age_group': getattr(g, 'age_group_snapshot', None) or '-',
                'diagnosis': getattr(g, 'diagnosis_snapshot', None) or '-',
                'department': p.department,
                'guardian': p.guardian_name,
                'registration_date': p.registration_date.strftime('%d-%b-%Y') if p.registration_date else '',
                'created_time': timezone.localtime(p.created_at).strftime('%I:%M %p'),
                'patient_type': p.patient_type,
                'op_type': getattr(g, 'op_type', 'NEW OP'),
                'status': 'Success'
            })
            
        return JsonResponse({
            'success': True, 
            'details': details,
            'departments': dept_stats,
            'patients': patients_data
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

class AutoTriggerMonthlyReviewsView(LoginRequiredMixin, MenuAccessRequiredMixin, TemplateView):
    menu_key = 'auto_trigger'
    template_name = 'lab/auto_trigger/monthly_trigger_reviews.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.lab.models import Department, MonthlyTriggerPlan
        context['departments'] = Department.objects.all().order_by('name')
        context['plans'] = MonthlyTriggerPlan.objects.filter(is_saved=True).order_by('-id')[:20]
        return context

def api_get_monthly_reviews(request):
    from apps.lab.models import MonthlyTriggerGeneratedPatient
    
    dept_id = request.GET.get('department_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    status = request.GET.get('status')
    
    reviews = MonthlyTriggerGeneratedPatient.objects.filter(op_type='REVIEW').select_related(
        'patient', 'target__plan', 'target__department'
    ).order_by('-created_at')
    
    if dept_id:
        reviews = reviews.filter(target__department_id=dept_id)
        
    if from_date:
        reviews = reviews.filter(target__target_date__gte=from_date)
        
    if to_date:
        reviews = reviews.filter(target__target_date__lte=to_date)
        
    if status and status != 'All':
        reviews = reviews.filter(status=status)
        
    reviews = reviews[:100] # Limit for performance
    
    data = []
    for r in reviews:
        p = r.patient
        plan = r.target.plan
        source_month = plan.review_source_month_year if plan else '-'
        
        data.append({
            'review_date': r.target.target_date.strftime('%d-%b-%Y'),
            'review_time': timezone.localtime(r.created_at).strftime('%I:%M %p'),
            'patient_id': p.patient_id,
            'op_number': p.op_number,
            'name': p.name,
            'age_gender': f"{p.age_years} / {p.gender}",
            'department': r.target.department.name,
            'original_op_date': p.registration_date.strftime('%d-%b-%Y') if p.registration_date else '-',
            'source_month': source_month,
            'patient_type': p.patient_type,
            'review_type': 'MONTHLY REVIEW',
            'status': r.status,
            'automation_source': 'MONTHLY_TRIGGER'
        })
        
    return JsonResponse({'success': True, 'data': data})

class AutoTriggerMonthlyExecutionView(LoginRequiredMixin, MenuAccessRequiredMixin, TemplateView):
    menu_key = 'auto_trigger'
    template_name = 'lab/auto_trigger/monthly_trigger_execution.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plan_id = self.kwargs.get('plan_id')
        date_str = self.kwargs.get('date_str')
        context['plan_id'] = plan_id
        context['date_str'] = date_str
        return context

# ==========================================
# AUTOMATE TEST & RESULT VIEW
# ==========================================

class AutoTriggerAutomateTestView(LoginRequiredMixin, MenuAccessRequiredMixin, TemplateView):
    menu_key = 'auto_trigger'
    template_name = 'lab/auto_trigger/automate_test.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.lab.models import Investigation
        context['investigations'] = Investigation.objects.filter(is_active=True).order_by('name')
        return context

class AutoTriggerResultView(LoginRequiredMixin, MenuAccessRequiredMixin, TemplateView):
    menu_key = 'auto_trigger'
    template_name = 'lab/auto_trigger/result_view.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.lab.models import AutomationDummyResult
        # Only fetch distinct investigations that have dummy results
        inv_ids = AutomationDummyResult.objects.values_list('investigation_id', flat=True).distinct()
        from apps.lab.models import Investigation
        context['investigations'] = Investigation.objects.filter(id__in=inv_ids).order_by('name')
        return context

def format_single_range(r):
    if r.range_type == 'Numeric':
        if r.min_value is None and r.max_value is None:
            return r.reference_text or "-"
        min_v = f"{r.min_value.normalize():f}" if r.min_value is not None else ""
        max_v = f"{r.max_value.normalize():f}" if r.max_value is not None else ""
        return f"{min_v}-{max_v}".strip("-")
    elif r.range_type == 'Text':
        return r.reference_text or "-"
    return "-"

def format_ranges_for_dummy(inv_param):
    ranges = inv_param.reference_ranges.filter(is_active=True).order_by('age_group__min_age_value', 'gender')
    if not ranges.exists():
        return "-"
    
    if ranges.count() == 1:
        r = ranges.first()
        return format_single_range(r)
        
    lines = []
    for r in ranges:
        prefix = ""
        if r.age_group:
            prefix += f"{r.age_group.label}"
        if r.gender and r.gender != 'All':
            if not r.age_group or r.gender.lower() not in r.age_group.label.lower():
                prefix += f" {r.gender}"
        
        prefix = f"{prefix.strip()}: " if prefix.strip() else ""
        lines.append(f"{prefix}{format_single_range(r)}")
        
    return "<br>".join(lines)

def api_get_investigation_parameters(request):
    try:
        inv_id = request.GET.get('investigation_id')
        if not inv_id:
            return JsonResponse({'success': False, 'message': 'Missing investigation_id'})
            
        from apps.lab.models import InvestigationParameter
        params = InvestigationParameter.objects.filter(investigation_id=inv_id, is_active=True).select_related('parameter').order_by('display_order', 'id')
        
        data = []
        for p in params:
            # 1. Resolve unit
            unit_val = getattr(p, 'unit', None)
            if not unit_val and p.parameter:
                unit_val = getattr(p.parameter, 'default_unit', None)
                
            # 2. Resolve reference range
            ref_str = format_ranges_for_dummy(p)
            
            data.append({
                'id': p.parameter.id if p.parameter else p.id,
                'name': p.parameter.name if p.parameter else p.name,
                'unit': unit_val or '',
                'reference_range': ref_str,
            })
            
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

def api_generate_dummy_result_values(request):
    try:
        inv_id = request.GET.get('investigation_id')
        if not inv_id:
            return JsonResponse({'success': False, 'message': 'Missing investigation_id'})
            
        from apps.lab.models import Investigation
        from apps.lab.services.automation_test_service import AutomationTestService
        
        inv = Investigation.objects.get(id=inv_id)
        dummy_data = AutomationTestService.generate_dummy_values_for_investigation(inv)
        
        return JsonResponse({'success': True, 'data': dummy_data})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@csrf_exempt
def api_save_dummy_result(request):
    import json
    from datetime import datetime
    from apps.lab.models import AutomationDummyResult, AutomationDummyResultParameter, Investigation
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid method'})
        
    try:
        data = json.loads(request.body)
        inv_id = data.get('investigation_id')
        remarks = data.get('remarks', '')
        parameters = data.get('parameters', [])
        
        if not inv_id or not parameters:
            return JsonResponse({'success': False, 'message': 'Missing required data'})
            
        inv = Investigation.objects.get(id=inv_id)
        
        dummy_id = data.get('dummy_id')
        
        if dummy_id:
            with transaction.atomic():
                dummy = AutomationDummyResult.objects.get(id=dummy_id)
                dummy.remarks = remarks
                dummy.save()
                
                # Delete old parameters and recreate, or update existing
                # Deleting and recreating is simpler and ensures display order is kept
                dummy.parameters.all().delete()
                
                for i, p in enumerate(parameters):
                    AutomationDummyResultParameter.objects.create(
                        dummy_result=dummy,
                        parameter_id=p.get('parameter_id'),
                        parameter_name=p.get('parameter_name'),
                        result_value=p.get('result_value'),
                        unit=p.get('unit'),
                        reference_range=p.get('reference_range'),
                        display_order=i
                    )
            return JsonResponse({'success': True, 'message': 'Result updated successfully', 'result_id': dummy.result_id})
        else:
            # Generate Result ID
            now = datetime.now()
            prefix = f"RES-{now.strftime('%Y-%m')}-"
            last_result = AutomationDummyResult.objects.filter(result_id__startswith=prefix).order_by('-result_id').first()
            if last_result:
                last_num = int(last_result.result_id.split('-')[-1])
                new_num = last_num + 1
            else:
                new_num = 1
            result_id = f"{prefix}{new_num:05d}"
            
            # Find dummy name (e.g., CBC Dummy 1)
            count = AutomationDummyResult.objects.filter(investigation=inv).count() + 1
            dummy_name = f"{inv.short_name or inv.name} Dummy {count}"
            
            with transaction.atomic():
                dummy = AutomationDummyResult.objects.create(
                    result_id=result_id,
                    investigation=inv,
                    investigation_code=inv.code or '',
                    dummy_name=dummy_name,
                    created_by=request.user,
                    remarks=remarks
                )
                
                for i, p in enumerate(parameters):
                    AutomationDummyResultParameter.objects.create(
                        dummy_result=dummy,
                        parameter_id=p.get('parameter_id'),
                        parameter_name=p.get('parameter_name'),
                        result_value=p.get('result_value'),
                        unit=p.get('unit'),
                        reference_range=p.get('reference_range'),
                        display_order=i
                    )
                    
            return JsonResponse({'success': True, 'message': 'Result saved successfully', 'result_id': result_id})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'success': False, 'message': str(e)})

def api_get_dummy_results(request):
    from apps.lab.models import AutomationDummyResult
    from django.utils import timezone
    
    inv_id = request.GET.get('investigation_id')
    
    qs = AutomationDummyResult.objects.all().select_related('investigation', 'created_by').prefetch_related('parameters')
    if inv_id and inv_id != 'all':
        qs = qs.filter(investigation_id=inv_id)
        
    qs = qs.order_by('investigation__name', 'id')
    
    # Group by investigation
    groups = {}
    for r in qs:
        inv_name = f"{r.investigation.name}"
        if r.investigation.short_name and r.investigation.short_name != r.investigation.name:
            inv_name = f"{r.investigation.short_name} - {inv_name}"
            
        if inv_name not in groups:
            groups[inv_name] = {
                'id': r.investigation.id,
                'name': inv_name,
                'short_name': r.investigation.short_name or r.investigation.name[:3].upper(),
                'results': []
            }
            
        params = []
        for p in r.parameters.all():
            params.append({
                'parameter_id': p.parameter_id,
                'name': p.parameter_name,
                'value': p.result_value,
                'unit': p.unit,
                'reference_range': p.reference_range,
                'result_status': p.result_status or 'NON_NUMERIC',
            })
            
        # Overall status: use DB-stored value
        overall_status = r.result_status or 'NON_NUMERIC'
        is_consumed = r.status == 'Consumed'

        groups[inv_name]['results'].append({
            'id': r.id,
            'result_id': r.result_id,
            'dummy_name': r.dummy_name,
            'created_on': timezone.localtime(r.created_at).strftime('%d-%b-%Y %I:%M %p'),
            'created_by': r.created_by.username if r.created_by else 'System',
            'status': r.status,
            'overall_status': overall_status,
            'is_consumed': is_consumed,
            'param_count': len(params),
            'parameters': params,
            'remarks': r.remarks,
            # Patient context
            'patient_name': r.patient_name or '',
            'patient_age': r.patient_age_display or '',
            'patient_gender': r.patient_gender or '',
        })
        
    return JsonResponse({'success': True, 'groups': list(groups.values())})

class AutoTriggerDashboardView(LoginRequiredMixin, TemplateView):
    template_name = 'lab/auto_trigger/dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.lab.models import Investigation, InvestigationParameter, AutomationDummyResult, AutoTriggerHistory
        from django.utils import timezone
        from datetime import timedelta
        from django.db.models import Count

        # Basic Stats
        context['total_investigations'] = Investigation.objects.count()
        context['total_parameters'] = InvestigationParameter.objects.count()
        context['total_dummy_results'] = AutomationDummyResult.objects.count()
        
        today = timezone.localtime().date()
        context['results_today'] = AutomationDummyResult.objects.filter(created_at__date=today).count()
        
        start_of_month = today.replace(day=1)
        context['monthly_tests'] = AutomationDummyResult.objects.filter(created_at__date__gte=start_of_month).count()

        # Results by investigation (Donut chart data)
        inv_counts = AutomationDummyResult.objects.values(
            'investigation__name', 'investigation__short_name'
        ).annotate(count=Count('id')).order_by('-count')
        
        donut_labels = []
        donut_data = []
        donut_details = []
        
        for item in inv_counts:
            short = item['investigation__short_name']
            name = item['investigation__name']
            label = f"{short} - {name}" if short and short != name else name
            donut_labels.append(label)
            donut_data.append(item['count'])
            donut_details.append({'label': label, 'count': item['count']})
            
        context['donut_labels_json'] = json.dumps(donut_labels)
        context['donut_data_json'] = json.dumps(donut_data)
        context['donut_details'] = donut_details

        # Results Trend (Last 7 Days)
        trend_labels = []
        trend_data = []
        for i in range(6, -1, -1):
            d = today - timedelta(days=i)
            trend_labels.append(d.strftime('%d %b'))
            c = AutomationDummyResult.objects.filter(created_at__date=d).count()
            trend_data.append(c)
            
        context['trend_labels_json'] = json.dumps(trend_labels)
        context['trend_data_json'] = json.dumps(trend_data)
        
        context['trend_total'] = sum(trend_data)
        context['trend_avg'] = round(context['trend_total'] / 7, 2)
        if trend_data:
            max_idx = trend_data.index(max(trend_data))
            min_idx = trend_data.index(min(trend_data))
            context['trend_highest'] = trend_data[max_idx]
            context['trend_highest_date'] = (today - timedelta(days=6 - max_idx)).strftime('%d %b %Y')
            context['trend_lowest'] = trend_data[min_idx]
            context['trend_lowest_date'] = (today - timedelta(days=6 - min_idx)).strftime('%d %b %Y')
        else:
            context['trend_highest'] = 0
            context['trend_highest_date'] = '-'
            context['trend_lowest'] = 0
            context['trend_lowest_date'] = '-'

        # Recent Dummy Results (last 5)
        recent = AutomationDummyResult.objects.select_related('investigation').prefetch_related('parameters').order_by('-created_at')[:5]
        recent_list = []
        for r in recent:
            recent_list.append({
                'inv_short': r.investigation.short_name or r.investigation.name[:3],
                'name': r.dummy_name,
                'result_id': r.result_id,
                'created_on': timezone.localtime(r.created_at).strftime('%d-%b-%Y %I:%M %p'),
                'param_count': r.parameters.count(),
                'id': r.id
            })
        context['recent_results'] = recent_list

        # Automation Summary
        context['successful_triggers'] = AutoTriggerHistory.objects.filter(status='Completed').count()
        context['pending_triggers'] = AutoTriggerHistory.objects.filter(status__in=['Pending', 'Running']).count()
        context['failed_triggers'] = AutoTriggerHistory.objects.filter(status='Failed').count()
        last_execution = AutoTriggerHistory.objects.order_by('-last_processed_at').first()
        if last_execution and last_execution.last_processed_at:
            context['last_trigger_run'] = timezone.localtime(last_execution.last_processed_at).strftime('%d-%b-%Y %I:%M %p')
        else:
            context['last_trigger_run'] = '-'

        return context


def api_result_import_export(request):
    from apps.lab.result_importer import export_dummy_results, validate_result_import, import_dummy_results
    import json
    from django.http import JsonResponse, HttpResponse
    
    if request.method == 'POST':
        action = request.POST.get('action')
        
        if action == 'validate':
            if 'file' not in request.FILES:
                return JsonResponse({'status': 'error', 'message': 'No file uploaded'})
            file_obj = request.FILES['file']
            res = validate_result_import(file_obj)
            return JsonResponse({'status': 'success', 'validation': res})
            
        elif action == 'import':
            if 'file' not in request.FILES:
                return JsonResponse({'status': 'error', 'message': 'No file uploaded'})
            file_obj = request.FILES['file']
            try:
                counts = import_dummy_results(file_obj, request.user)
                return JsonResponse({'status': 'success', 'counts': counts})
            except Exception as e:
                import traceback
                traceback.print_exc()
                return JsonResponse({'status': 'error', 'message': str(e)})
                
    elif request.method == 'GET':
        # Blank import template
        if request.GET.get('template') == 'true':
            from openpyxl import Workbook
            import io
            wb = Workbook()
            ws = wb.active
            ws.title = 'Results'
            headers = [
                'Result ID', 'Investigation', 'Investigation Code',
                'Parameter', 'Parameter Code', 'Result Value',
                'Unit', 'Reference Range', 'Remarks', 'Status', 'Created On',
            ]
            ws.append(headers)
            # Style header row
            from openpyxl.styles import Font, PatternFill, Alignment
            header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="result_import_template.xlsx"'
            return response

        if request.GET.get('export') == 'excel':
            filters = {
                'investigation_id': request.GET.get('investigation_id'),
                'date_from': request.GET.get('date_from'),
                'date_to': request.GET.get('date_to'),
            }
            output = export_dummy_results(filters)
            response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="result_export.xlsx"'
            return response
            
    return JsonResponse({'status': 'error', 'message': 'Invalid request'})


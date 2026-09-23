"""
Universal Import / Export engine for VMMC ERP Lab Master.

Handles a multi-sheet workbook with sheets:
  1. Departments
  2. Diagnoses
  3. Investigations
  4. Parameters
  5. Age Groups
  6. Reference Ranges
  7. Investigation Parameters
  8. Diagnosis Departments
  9. Diagnosis Investigations

Provides:
  - validate_import(file_obj): Pre-import inspection returning comprehensive sheet-by-sheet
    stats (detected, valid, invalid, duplicates, new, existing), column validation, data type
    checks, FK & sheet-to-sheet reference checks, and row-level error reporting.
  - import_universal_master(file_obj, admin_user=None): Transaction-wrapped safe upsert.
  - export_universal_master(): Exports current master configuration into ONE Excel workbook
    preserving exact schema for 100% round-trip compatibility.
  - generate_blank_template(): Generates formatted empty template with all 9 sheets.
"""

import io
import logging
from decimal import Decimal

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db import transaction, models, IntegrityError
import re

from apps.lab.models import (
    LabDepartment, Diagnosis, Investigation, Parameter, AgeGroup, SampleType,
    InvestigationParameter, ParameterReferenceRange, DiagnosisInvestigationMap,
    DiagnosisDepartmentMapping,
)
from apps.patients.models import Department

logger = logging.getLogger('apps.lab.universal_importer')

# ---------------------------------------------------------------------------
# Single source of truth: sheet names and their column definitions
# ---------------------------------------------------------------------------
SHEET_DEFS = {
    'Departments': {
        'required': True,
        'columns': ['Department Name', 'Department Code', 'Status'],
        'key_fields': ['Department Code'],
    },
    'Diagnoses': {
        'required': True,
        'columns': ['Diagnosis Name', 'Diagnosis Code', 'Status'],
        'key_fields': ['Diagnosis Code'],
    },
    'Investigations': {
        'required': True,
        'columns': ['Investigation Name', 'Investigation Code', 'Department',
                    'Sample Type', 'Is Panel', 'Status'],
        'key_fields': ['Investigation Code'],
    },
    'Parameters': {
        'required': True,
        'columns': ['Parameter Name', 'Parameter Code', 'Data Type',
                    'Default Unit', 'Status'],
        'key_fields': ['Parameter Code'],
    },
    'Age Groups': {
        'required': False,
        'columns': ['Age Group Code', 'Age Group Name', 'Age From',
                    'Age From Unit', 'Age To', 'Age To Unit', 'Gender', 'Status'],
        'key_fields': ['Age Group Code'],
    },
    'Investigation Parameters': {
        'required': False,
        'columns': ['Investigation Code', 'Parameter Code', 'Display Order', 'Status'],
        'key_fields': ['Investigation Code', 'Parameter Code'],
    },
    'Reference Ranges': {
        'required': False,
        'columns': ['Investigation Code', 'Parameter Code', 'Age Group Code',
                    'Gender', 'Diagnosis Code', 'Min Value', 'Max Value',
                    'Normal Value', 'Unit', 'Method', 'Remarks', 'Status'],
        'key_fields': ['Investigation Code', 'Parameter Code', 'Age Group Code', 'Gender', 'Diagnosis Code'],
    },
    'Diagnosis Departments': {
        'required': False,
        'columns': ['Diagnosis Code', 'Department Code', 'Age Group Code', 'Status'],
        'key_fields': ['Diagnosis Code', 'Department Code', 'Age Group Code'],
    },
    'Diagnosis Investigations': {
        'required': False,
        'columns': ['Diagnosis Code', 'Investigation Code', 'Age Group Code', 'Status'],
        'key_fields': ['Diagnosis Code', 'Investigation Code', 'Age Group Code'],
    },
}

SHEET_ORDER = [
    'Departments',
    'Age Groups',
    'Diagnoses',
    'Investigations',
    'Parameters',
    'Diagnosis Departments',
    'Investigation Parameters',
    'Reference Ranges',
    'Diagnosis Investigations',
]


# ---------------------------------------------------------------------------
# Value Sanitation & Normalization Helpers
# ---------------------------------------------------------------------------

def clean_code(v, candidate_codes=None):
    """
    Cleans an identifier code (Department, Age Group, Diagnosis, Investigation, Parameter).
    Treats codes STRICTLY as string values, preserving leading zeros (e.g. 00022681, 00020346).
    Never coerces codes to integers.
    """
    if v is None:
        return ''
    if isinstance(v, float):
        if v.is_integer():
            s = str(int(v)).strip()
        else:
            s = str(v).strip()
    else:
        s = str(v).strip()

    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]

    if not s:
        return ''

    # If candidate codes are available, match against them:
    if candidate_codes:
        s_lower = s.lower()
        for cand in candidate_codes:
            if cand and str(cand).strip().lower() == s_lower:
                return str(cand).strip()
        # If numeric, match against candidate_codes ignoring leading zeros
        if s.isdigit():
            s_lstrip = s.lstrip('0')
            matches = [c for c in candidate_codes if c and str(c).strip().isdigit() and str(c).strip().lstrip('0') == s_lstrip]
            if len(matches) == 1:
                return str(matches[0]).strip()
            # Try standard 8-digit zfill
            z8 = s.zfill(8)
            for cand in candidate_codes:
                if cand and str(cand).strip() == z8:
                    return z8

    return s


def clean_val(v):
    if v is None:
        return ''
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        return s[:-2]
    return s


def normalize_text(v):
    """Normalize string: trim leading/trailing whitespace and collapse multiple internal spaces."""
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def get_decimal(v):
    if v is None or str(v).strip() == '':
        return None
    try:
        return Decimal(str(v).strip())
    except Exception:
        return None


def is_active_bool(v):
    if v is None:
        return True
    return str(v).strip().lower() in ['active', 'true', '1', 'yes', 'y']


class UniversalImportError(Exception):
    """Structured error raised when Universal Import encounters an issue."""
    def __init__(self, message, entity='Department', entity_name='', row_number=None, reason='', action=''):
        super().__init__(message)
        self.message = message
        self.entity = entity
        self.entity_name = entity_name
        self.row_number = row_number
        self.reason = reason
        self.action = action

    def to_dict(self):
        return {
            'entity': self.entity,
            'entity_name': self.entity_name,
            'row_number': self.row_number,
            'reason': self.reason or self.message,
            'action': self.action or 'Existing record should be updated instead of created.',
            'message': self.message,
        }


# ---------------------------------------------------------------------------
# Shared Entity Resolvers (Single source of matching logic for Preview & Import)
# ---------------------------------------------------------------------------

class DepartmentResolver:
    """
    Shared Department resolver for both Preview and Import.
    Matches primarily by department code, and secondarily by normalized department name.
    Guarantees true upsert without ever failing patients_department.name or code UNIQUE constraint.
    """
    def __init__(self):
        self.by_code = {}        # lower(code) -> Department instance
        self.by_name = {}        # lower(normalize(name)) -> Department instance
        self.seen_codes = set()  # lower(code) seen in workbook for duplicate detection
        self.seen_names = set()  # lower(normalize(name)) seen in workbook
        self.preload()

    def preload(self):
        for d in Department.objects.all():
            self.register(d)

    def register(self, d):
        if d.code:
            self.by_code[clean_code(d.code).lower()] = d
        if d.name:
            self.by_name[normalize_text(d.name).lower()] = d

    def register_placeholder(self, code, name):
        """Register newly resolved/created department keys so later workbook rows update it."""
        key_c = clean_code(code).lower()
        key_n = normalize_text(name).lower()
        dummy = {'code': clean_code(code), 'name': normalize_text(name)}
        if key_c and key_c not in self.by_code:
            self.by_code[key_c] = dummy
        if key_n and key_n not in self.by_name:
            self.by_name[key_n] = dummy

    def resolve(self, code, name):
        """
        Returns:
            status: 'UPDATE' | 'CREATE'
            target: existing Department or dict or None
            is_duplicate: bool (True if duplicate within the current workbook)
        """
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        key_c = c_clean.lower() if c_clean else ''
        key_n = n_clean.lower() if n_clean else ''

        is_dup = False
        if key_c and key_c in self.seen_codes:
            is_dup = True
        if not key_c and key_n and key_n in self.seen_names:
            is_dup = True

        if key_c:
            self.seen_codes.add(key_c)
        if key_n:
            self.seen_names.add(key_n)

        # 1. Prefer match by code
        if key_c and key_c in self.by_code:
            return 'UPDATE', self.by_code[key_c], is_dup

        # 2. Match by normalized name ONLY if code is not given
        if not key_c and key_n and key_n in self.by_name:
            return 'UPDATE', self.by_name[key_n], is_dup

        # 3. Not found -> CREATE
        return 'CREATE', None, is_dup

    def get_department(self, code_or_name):
        """Helper for FK lookups in subsequent sheets (e.g. Diagnosis Departments)."""
        if not code_or_name:
            return None
        c_val = clean_code(code_or_name, self.by_code.keys()).lower()
        if c_val in self.by_code:
            obj = self.by_code[c_val]
            if isinstance(obj, Department):
                return obj
        key = normalize_text(code_or_name).lower()
        if key in self.by_code:
            obj = self.by_code[key]
            if isinstance(obj, Department):
                return obj
        if key in self.by_name:
            obj = self.by_name[key]
            if isinstance(obj, Department):
                return obj
        return Department.objects.filter(
            models.Q(code__iexact=clean_code(code_or_name)) | models.Q(name__iexact=normalize_text(code_or_name))
        ).first()

    def upsert(self, code, name, is_active=True, row_index=None):
        """
        Executes safe upsert in the database.
        Returns: (instance, created_bool)
        """
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        if not c_clean or not n_clean:
            return None, False

        status, target, _ = self.resolve(c_clean, n_clean)

        if status == 'UPDATE' and isinstance(target, Department):
            dept = target
            # Update name if changed and not conflicting with another DB department
            if normalize_text(dept.name).lower() != n_clean.lower():
                other = Department.objects.filter(name__iexact=n_clean).exclude(id=dept.id).first()
                if other:
                    dept = other
                else:
                    dept.name = n_clean

            # Update code if changed and not conflicting
            if c_clean and dept.code.lower() != c_clean.lower():
                other_code = Department.objects.filter(code__iexact=c_clean).exclude(id=dept.id).first()
                if not other_code:
                    dept.code = c_clean

            dept.is_active = is_active
            dept.save()
            self.register(dept)
            return dept, False

        # Otherwise CREATE, with absolute guard against race conditions or name collision
        existing = Department.objects.filter(
            models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
        ).first()

        if existing:
            if normalize_text(existing.name).lower() != n_clean.lower():
                other_name = Department.objects.filter(name__iexact=n_clean).exclude(id=existing.id).first()
                if not other_name:
                    existing.name = n_clean
            if c_clean and existing.code.lower() != c_clean.lower():
                if not Department.objects.filter(code__iexact=c_clean).exclude(id=existing.id).exists():
                    existing.code = c_clean
            existing.is_active = is_active
            existing.save()
            self.register(existing)
            return existing, False

        try:
            dept = Department.objects.create(
                code=c_clean,
                name=n_clean,
                is_active=is_active
            )
            self.register(dept)
            return dept, True
        except Exception as e:
            # Fallback if constraint triggers
            fallback = Department.objects.filter(
                models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
            ).first()
            if fallback:
                fallback.is_active = is_active
                fallback.save()
                self.register(fallback)
                return fallback, False

            raise UniversalImportError(
                message=f"Failed to upsert department '{n_clean}': {str(e)}",
                entity='Department',
                entity_name=n_clean,
                row_number=row_index,
                reason="A department with this unique identity already exists.",
                action="Existing department should be updated instead of created."
            )


class DiagnosisResolver:
    """
    Shared Diagnosis resolver for both Preview and Import.
    Matches primarily by diagnosis code, secondarily by normalized diagnosis name.
    """
    def __init__(self):
        self.by_code = {}
        self.by_name = {}
        self.seen_codes = set()
        self.seen_names = set()
        self.preload()

    def preload(self):
        for d in Diagnosis.objects.all():
            self.register(d)

    def register(self, d):
        if d.code:
            self.by_code[clean_code(d.code).lower()] = d
        if d.name:
            self.by_name[normalize_text(d.name).lower()] = d

    def register_placeholder(self, code, name):
        key_c = clean_code(code).lower()
        key_n = normalize_text(name).lower()
        dummy = {'code': clean_code(code), 'name': normalize_text(name)}
        if key_c and key_c not in self.by_code:
            self.by_code[key_c] = dummy
        if key_n and key_n not in self.by_name:
            self.by_name[key_n] = dummy

    def resolve(self, code, name):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        key_c = c_clean.lower() if c_clean else ''
        key_n = n_clean.lower() if n_clean else ''

        is_dup = False
        if key_c and key_c in self.seen_codes:
            is_dup = True
        if not key_c and key_n and key_n in self.seen_names:
            is_dup = True

        if key_c:
            self.seen_codes.add(key_c)
        if key_n:
            self.seen_names.add(key_n)

        # Primary: by code
        if key_c and key_c in self.by_code:
            return 'UPDATE', self.by_code[key_c], is_dup
        # Secondary: by name if code not given
        if not key_c and key_n and key_n in self.by_name:
            return 'UPDATE', self.by_name[key_n], is_dup

        return 'CREATE', None, is_dup

    def get_diagnosis(self, code_or_name):
        if not code_or_name:
            return None
        c_val = clean_code(code_or_name, self.by_code.keys()).lower()
        if c_val in self.by_code and isinstance(self.by_code[c_val], Diagnosis):
            return self.by_code[c_val]
        key = normalize_text(code_or_name).lower()
        if key in self.by_code and isinstance(self.by_code[key], Diagnosis):
            return self.by_code[key]
        if key in self.by_name and isinstance(self.by_name[key], Diagnosis):
            return self.by_name[key]
        return Diagnosis.objects.filter(
            models.Q(code__iexact=clean_code(code_or_name)) | models.Q(name__iexact=normalize_text(code_or_name))
        ).first()

    def upsert(self, code, name, is_active=True, row_index=None):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        if not c_clean or not n_clean:
            return None, False

        status, target, _ = self.resolve(c_clean, n_clean)
        if status == 'UPDATE' and isinstance(target, Diagnosis):
            diag = target
            diag.name = n_clean
            if c_clean and diag.code != c_clean:
                if not Diagnosis.objects.filter(code__iexact=c_clean).exclude(id=diag.id).exists():
                    diag.code = c_clean
            diag.is_active = is_active
            diag.save()
            self.register(diag)
            return diag, False

        existing = Diagnosis.objects.filter(
            models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
        ).first()
        if existing:
            existing.name = n_clean
            if c_clean and existing.code != c_clean:
                if not Diagnosis.objects.filter(code__iexact=c_clean).exclude(id=existing.id).exists():
                    existing.code = c_clean
            existing.is_active = is_active
            existing.save()
            self.register(existing)
            return existing, False

        try:
            diag = Diagnosis.objects.create(
                code=c_clean,
                name=n_clean,
                is_active=is_active
            )
            self.register(diag)
            return diag, True
        except Exception as e:
            fallback = Diagnosis.objects.filter(
                models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
            ).first()
            if fallback:
                fallback.is_active = is_active
                fallback.save()
                self.register(fallback)
                return fallback, False

            raise UniversalImportError(
                message=f"Failed to upsert diagnosis '{n_clean}': {str(e)}",
                entity='Diagnosis',
                entity_name=n_clean,
                row_number=row_index,
                reason="A diagnosis with this unique identity already exists.",
                action="Existing diagnosis should be updated instead of created."
            )


class InvestigationResolver:
    """
    Shared Investigation resolver for both Preview and Import.
    Matches primarily by investigation code, secondarily by normalized name.
    """
    def __init__(self):
        self.by_code = {}
        self.by_name = {}
        self.seen_codes = set()
        self.seen_names = set()
        self.preload()

    def preload(self):
        for inv in Investigation.objects.all():
            self.register(inv)

    def register(self, inv):
        if inv.code:
            self.by_code[clean_code(inv.code).lower()] = inv
        if inv.name:
            self.by_name[normalize_text(inv.name).lower()] = inv

    def register_placeholder(self, code, name):
        key_c = clean_code(code).lower()
        key_n = normalize_text(name).lower()
        dummy = {'code': clean_code(code), 'name': normalize_text(name)}
        if key_c and key_c not in self.by_code:
            self.by_code[key_c] = dummy
        if key_n and key_n not in self.by_name:
            self.by_name[key_n] = dummy

    def resolve(self, code, name):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        key_c = c_clean.lower() if c_clean else ''
        key_n = n_clean.lower() if n_clean else ''

        is_dup = False
        if key_c and key_c in self.seen_codes:
            is_dup = True
        if not key_c and key_n and key_n in self.seen_names:
            is_dup = True

        if key_c:
            self.seen_codes.add(key_c)
        if key_n:
            self.seen_names.add(key_n)

        # Primary: by code
        if key_c and key_c in self.by_code:
            return 'UPDATE', self.by_code[key_c], is_dup
        # Secondary: by name if code not given
        if not key_c and key_n and key_n in self.by_name:
            return 'UPDATE', self.by_name[key_n], is_dup

        return 'CREATE', None, is_dup

    def get_investigation(self, code_or_name):
        if not code_or_name:
            return None
        c_val = clean_code(code_or_name, self.by_code.keys()).lower()
        if c_val in self.by_code and isinstance(self.by_code[c_val], Investigation):
            return self.by_code[c_val]
        key = normalize_text(code_or_name).lower()
        if key in self.by_code and isinstance(self.by_code[key], Investigation):
            return self.by_code[key]
        if key in self.by_name and isinstance(self.by_name[key], Investigation):
            return self.by_name[key]
        return Investigation.objects.filter(
            models.Q(code__iexact=clean_code(code_or_name)) | models.Q(name__iexact=normalize_text(code_or_name))
        ).first()

    def upsert(self, code, name, dept_name='', samp_name='', is_panel=False, is_active=True, row_index=None):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        if not c_clean or not n_clean:
            return None, False

        ldept = LabDepartment.objects.get_or_create(name=normalize_text(dept_name))[0] if dept_name else None
        samp = SampleType.objects.get_or_create(name=normalize_text(samp_name))[0] if samp_name else None

        status, target, _ = self.resolve(c_clean, n_clean)
        if status == 'UPDATE' and isinstance(target, Investigation):
            inv = target
            inv.name = n_clean
            if c_clean and inv.code != c_clean:
                if not Investigation.objects.filter(code__iexact=c_clean).exclude(id=inv.id).exists():
                    inv.code = c_clean
            if ldept:
                inv.department = ldept
            if samp:
                inv.sample_type = samp
            inv.is_panel = is_panel
            inv.is_active = is_active
            inv.save()
            self.register(inv)
            return inv, False

        existing = Investigation.objects.filter(
            models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
        ).first()
        if existing:
            existing.name = n_clean
            if c_clean and existing.code != c_clean:
                if not Investigation.objects.filter(code__iexact=c_clean).exclude(id=existing.id).exists():
                    existing.code = c_clean
            if ldept:
                existing.department = ldept
            if samp:
                existing.sample_type = samp
            existing.is_panel = is_panel
            existing.is_active = is_active
            existing.save()
            self.register(existing)
            return existing, False

        try:
            inv = Investigation.objects.create(
                code=c_clean,
                name=n_clean,
                department=ldept,
                sample_type=samp,
                is_panel=is_panel,
                is_active=is_active
            )
            self.register(inv)
            return inv, True
        except Exception as e:
            fallback = Investigation.objects.filter(
                models.Q(code__iexact=c_clean) | models.Q(name__iexact=n_clean)
            ).first()
            if fallback:
                fallback.is_active = is_active
                fallback.save()
                self.register(fallback)
                return fallback, False

            raise UniversalImportError(
                message=f"Failed to upsert investigation '{n_clean}': {str(e)}",
                entity='Investigation',
                entity_name=n_clean,
                row_number=row_index,
                reason="An investigation with this unique identity already exists.",
                action="Existing investigation should be updated instead of created."
            )


class ParameterResolver:
    """
    Shared Parameter resolver for both Preview and Import.
    Matches primarily by parameter code, secondarily by normalized name.
    """
    def __init__(self):
        self.by_code = {}
        self.by_name = {}
        self.seen_codes = set()
        self.seen_names = set()
        self.preload()

    def preload(self):
        for p in Parameter.objects.all():
            self.register(p)

    def register(self, p):
        if p.code:
            self.by_code[clean_code(p.code).lower()] = p
        if p.name:
            self.by_name[normalize_text(p.name).lower()] = p

    def register_placeholder(self, code, name):
        key_c = clean_code(code).lower()
        key_n = normalize_text(name).lower()
        dummy = {'code': clean_code(code), 'name': normalize_text(name)}
        if key_c and key_c not in self.by_code:
            self.by_code[key_c] = dummy
        if key_n and key_n not in self.by_name:
            self.by_name[key_n] = dummy

    def resolve(self, code, name):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        key_c = c_clean.lower() if c_clean else ''
        key_n = n_clean.lower() if n_clean else ''

        is_dup = False
        if key_c and key_c in self.seen_codes:
            is_dup = True
        if not key_c and key_n and key_n in self.seen_names:
            is_dup = True

        if key_c:
            self.seen_codes.add(key_c)
        if key_n:
            self.seen_names.add(key_n)

        # Primary match: by code
        if key_c and key_c in self.by_code:
            return 'UPDATE', self.by_code[key_c], is_dup
            
        # Secondary match: by name ONLY if code is NOT provided
        if not key_c and key_n and key_n in self.by_name:
            return 'UPDATE', self.by_name[key_n], is_dup

        return 'CREATE', None, is_dup

    def get_parameter(self, code_or_name):
        if not code_or_name:
            return None
        c_val = clean_code(code_or_name, self.by_code.keys()).lower()
        if c_val in self.by_code and isinstance(self.by_code[c_val], Parameter):
            return self.by_code[c_val]
        from apps.lab.result_importer import normalize_parameter_code
        norm_code = normalize_parameter_code(code_or_name, self.by_code.keys())
        key_c = (norm_code or clean_code(code_or_name)).lower()
        if key_c in self.by_code and isinstance(self.by_code[key_c], Parameter):
            return self.by_code[key_c]
        key = normalize_text(code_or_name).lower()
        if key in self.by_code and isinstance(self.by_code[key], Parameter):
            return self.by_code[key]
        if key in self.by_name and isinstance(self.by_name[key], Parameter):
            return self.by_name[key]
        return Parameter.objects.filter(
            models.Q(code__iexact=norm_code or clean_code(code_or_name)) | models.Q(name__iexact=normalize_text(code_or_name))
        ).first()

    def upsert(self, code, name, data_type='NUMERIC', default_unit='', is_active=True, row_index=None):
        c_clean = clean_code(code, self.by_code.keys())
        n_clean = normalize_text(name)
        if not c_clean or not n_clean:
            return None, False

        status, target, _ = self.resolve(c_clean, n_clean)
        dtype = data_type.upper() if data_type else 'NUMERIC'
        if status == 'UPDATE' and isinstance(target, Parameter):
            param = target
            param.name = n_clean
            if c_clean and param.code != c_clean:
                if not Parameter.objects.filter(code__iexact=c_clean).exclude(id=param.id).exists():
                    param.code = c_clean
            param.data_type = dtype
            if default_unit or not param.default_unit:
                param.default_unit = default_unit
            param.is_active = is_active
            param.save()
            self.register(param)
            return param, False

        existing = None
        if c_clean:
            existing = Parameter.objects.filter(code__iexact=c_clean).first()
        elif n_clean:
            existing = Parameter.objects.filter(name__iexact=n_clean).first()

        if existing:
            existing.name = n_clean
            if c_clean and existing.code != c_clean:
                if not Parameter.objects.filter(code__iexact=c_clean).exclude(id=existing.id).exists():
                    existing.code = c_clean
            existing.data_type = dtype
            if default_unit or not existing.default_unit:
                existing.default_unit = default_unit
            existing.is_active = is_active
            existing.save()
            self.register(existing)
            return existing, False

        try:
            param = Parameter.objects.create(
                code=c_clean,
                name=n_clean,
                data_type=dtype,
                default_unit=default_unit,
                is_active=is_active
            )
            self.register(param)
            return param, True
        except Exception as e:
            fallback = Parameter.objects.filter(code__iexact=c_clean).first() if c_clean else Parameter.objects.filter(name__iexact=n_clean).first()
            if fallback:
                fallback.is_active = is_active
                fallback.save()
                self.register(fallback)
                return fallback, False

            raise UniversalImportError(
                message=f"Failed to upsert parameter '{n_clean}': {str(e)}",
                entity='Parameter',
                entity_name=n_clean,
                row_number=row_index,
                reason="A parameter with this unique identity already exists.",
                action="Existing parameter should be updated instead of created."
            )


class AgeGroupResolver:
    """
    Shared AgeGroup resolver for both Preview and Import.
    Matches primarily by age group code, secondarily by normalized label.
    """
    def __init__(self):
        self.by_code = {}
        self.by_label = {}
        self.seen_codes = set()
        self.seen_labels = set()
        self.preload()

    def preload(self):
        for ag in AgeGroup.objects.all():
            self.register(ag)

    def register(self, ag):
        if ag.code:
            self.by_code[clean_code(ag.code).lower()] = ag
        if ag.label:
            self.by_label[normalize_text(ag.label).lower()] = ag

    def register_placeholder(self, code, label):
        key_c = clean_code(code).lower()
        key_l = normalize_text(label).lower()
        dummy = {'code': clean_code(code), 'label': normalize_text(label)}
        if key_c and key_c not in self.by_code:
            self.by_code[key_c] = dummy
        if key_l and key_l not in self.by_label:
            self.by_label[key_l] = dummy

    def resolve(self, code, label):
        c_clean = clean_code(code, self.by_code.keys())
        l_clean = normalize_text(label)
        key_c = c_clean.lower() if c_clean else ''
        key_l = l_clean.lower() if l_clean else ''

        is_dup = False
        if key_c and key_c in self.seen_codes:
            is_dup = True
        if not key_c and key_l and key_l in self.seen_labels:
            is_dup = True

        if key_c:
            self.seen_codes.add(key_c)
        if key_l:
            self.seen_labels.add(key_l)

        # Primary match: by code
        if key_c and key_c in self.by_code:
            return 'UPDATE', self.by_code[key_c], is_dup
        # Secondary match: by label ONLY if code is NOT provided
        if not key_c and key_l and key_l in self.by_label:
            return 'UPDATE', self.by_label[key_l], is_dup

        return 'CREATE', None, is_dup

    def get_age_group(self, code_or_label):
        if not code_or_label:
            return None
        c_val = clean_code(code_or_label, self.by_code.keys()).lower()
        if c_val in self.by_code and isinstance(self.by_code[c_val], AgeGroup):
            return self.by_code[c_val]
        key = normalize_text(code_or_label).lower()
        if key in self.by_code and isinstance(self.by_code[key], AgeGroup):
            return self.by_code[key]
        if key in self.by_label and isinstance(self.by_label[key], AgeGroup):
            return self.by_label[key]
        return AgeGroup.objects.filter(
            models.Q(code__iexact=clean_code(code_or_label)) | models.Q(label__iexact=normalize_text(code_or_label))
        ).first()

    def upsert(self, code, label, a_from=None, f_unit='Years', a_to=None, t_unit='Years', gender='All', is_active=True, row_index=None):
        c_clean = clean_code(code, self.by_code.keys())
        l_clean = normalize_text(label)
        if not c_clean or not l_clean:
            return None, False

        min_val = int(a_from) if a_from and str(a_from).isdigit() else None
        max_val = int(a_to) if a_to and str(a_to).isdigit() else None

        status, target, _ = self.resolve(c_clean, l_clean)
        if status == 'UPDATE' and isinstance(target, AgeGroup):
            ag = target
            ag.label = l_clean
            if c_clean and ag.code != c_clean:
                if not AgeGroup.objects.filter(code__iexact=c_clean).exclude(id=ag.id).exists():
                    ag.code = c_clean
            ag.min_age_value = min_val
            ag.min_age_unit = f_unit or 'Years'
            ag.max_age_value = max_val
            ag.max_age_unit = t_unit or 'Years'
            ag.gender = gender or 'All'
            ag.is_active = is_active
            ag.save()
            self.register(ag)
            return ag, False

        existing = None
        if c_clean:
            existing = AgeGroup.objects.filter(code__iexact=c_clean).first()
        elif l_clean:
            existing = AgeGroup.objects.filter(label__iexact=l_clean).first()

        if existing:
            existing.label = l_clean
            if c_clean and existing.code != c_clean:
                if not AgeGroup.objects.filter(code__iexact=c_clean).exclude(id=existing.id).exists():
                    existing.code = c_clean
            existing.min_age_value = min_val
            existing.min_age_unit = f_unit or 'Years'
            existing.max_age_value = max_val
            existing.max_age_unit = t_unit or 'Years'
            existing.gender = gender or 'All'
            existing.is_active = is_active
            existing.save()
            self.register(existing)
            return existing, False

        try:
            ag = AgeGroup.objects.create(
                code=c_clean,
                label=l_clean,
                min_age_value=min_val,
                min_age_unit=f_unit or 'Years',
                max_age_value=max_val,
                max_age_unit=t_unit or 'Years',
                gender=gender or 'All',
                is_active=is_active
            )
            self.register(ag)
            return ag, True
        except Exception as e:
            fallback = AgeGroup.objects.filter(code__iexact=c_clean).first() if c_clean else AgeGroup.objects.filter(label__iexact=l_clean).first()
            if fallback:
                fallback.is_active = is_active
                fallback.save()
                self.register(fallback)
                return fallback, False

            raise UniversalImportError(
                message=f"Failed to upsert age group '{l_clean}': {str(e)}",
                entity='Age Group',
                entity_name=l_clean,
                row_number=row_index,
                reason="An age group with this unique identity already exists.",
                action="Existing age group should be updated instead of created."
            )


def normalize_sheet_name(raw_name):
    """Strip whitespace; return canonical SHEET_DEFS key or the raw name."""
    name = str(raw_name).strip()
    if name in SHEET_DEFS:
        return name
    for k in SHEET_DEFS:
        if k.lower() == name.lower():
            return k
    return name



def get_workbook_sheet_map(wb):
    """Return {canonical_name: ws} with whitespace/case normalisation."""
    result = {}
    for raw in wb.sheetnames:
        result[normalize_sheet_name(raw)] = wb[raw]
    return result


def parse_sheet(ws):
    """Parse worksheet into list of row-dicts (lowercased keys, skips blanks)."""
    if ws is None:
        return []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [
        str(h).strip().lower() if (h is not None and str(h).strip()) else '_col_%d' % i
        for i, h in enumerate(rows[0])
    ]
    data = []
    for row_idx, row in enumerate(rows[1:], start=2):
        row_data = {}
        is_empty = True
        for col_idx, value in enumerate(row):
            if col_idx < len(headers):
                if value is not None and str(value).strip():
                    is_empty = False
                row_data[headers[col_idx]] = value
        if not is_empty:
            row_data['_row_index'] = row_idx
            data.append(row_data)
    return data


def detect_format(wb, sheet_map):
    """Returns 'multi', 'legacy', or 'unknown'."""
    matched = set(SHEET_DEFS.keys()) & set(sheet_map.keys())
    if len(matched) >= 2:
        return 'multi'
    if 'Universal Master' in wb.sheetnames:
        return 'legacy'
    return 'unknown'


def _style_header_row(ws):
    fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    font = Font(bold=True, color='FFFFFF', size=10)
    border = Border(
        bottom=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
    )
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border


# ---------------------------------------------------------------------------
# VALIDATION (read-only — no DB writes)
# ---------------------------------------------------------------------------

def validate_import(file_obj):
    """
    Validate a Universal Master workbook.

    Returns a dict with keys:
        format, detected_sheets, sheet_stats, total_stats, stats (legacy compat),
        errors, warnings, valid

    Every error dict has:
        {sheet, row, field, value, message}
    """
    errors = []
    warnings = []

    # Sheet-by-sheet detailed metrics
    sheet_stats = {}
    for s_name in SHEET_ORDER:
        sheet_stats[s_name] = {
            'detected_rows': 0,
            'valid_rows': 0,
            'invalid_rows': 0,
            'duplicates': 0,
            'new': 0,
            'existing': 0,
        }

    def err(sheet, row, field, value, message):
        errors.append({
            'sheet': sheet,
            'row': row,
            'field': field,
            'value': str(value) if value is not None else '',
            'message': message,
            'error': message  # compatibility with frontend key
        })

    # -- Open workbook -------------------------------------------------------
    try:
        wb = load_workbook(filename=file_obj, data_only=True, read_only=False)
    except Exception as exc:
        logger.exception('universal_importer.validate_import: failed to open workbook')
        return {
            'format': 'unknown',
            'detected_sheets': [],
            'sheet_stats': sheet_stats,
            'total_stats': {'detected_rows': 0, 'valid_rows': 0, 'invalid_rows': 1, 'duplicates': 0, 'new': 0, 'existing': 0},
            'stats': {},
            'warnings': [],
            'errors': [{'sheet': 'File', 'row': 0, 'field': 'File', 'value': '',
                        'message': 'Could not open Excel file: %s' % exc,
                        'error': 'Could not open Excel file: %s' % exc}],
            'valid': False,
        }

    detected_sheets = list(wb.sheetnames)
    sheet_map = get_workbook_sheet_map(wb)
    fmt = detect_format(wb, sheet_map)
    logger.info('universal_importer.validate_import: format=%s sheets=%s', fmt, detected_sheets)

    # -- Legacy format handling ----------------------------------------------
    if fmt == 'legacy':
        warnings.append(
            "This workbook uses the older single-sheet 'Universal Master' format and does not "
            "contain complete mapping sheets. Please download the latest Universal Import Template "
            "or use Export Complete Master to get the current multi-sheet format."
        )
        return {
            'format': 'legacy',
            'detected_sheets': detected_sheets,
            'sheet_stats': sheet_stats,
            'total_stats': {'detected_rows': 0, 'valid_rows': 0, 'invalid_rows': 0, 'duplicates': 0, 'new': 0, 'existing': 0},
            'stats': {},
            'errors': [],
            'warnings': warnings,
            'valid': False,
        }

    if fmt == 'unknown':
        return {
            'format': 'unknown',
            'detected_sheets': detected_sheets,
            'sheet_stats': sheet_stats,
            'total_stats': {'detected_rows': 0, 'valid_rows': 0, 'invalid_rows': 1, 'duplicates': 0, 'new': 0, 'existing': 0},
            'stats': {},
            'warnings': [],
            'errors': [{
                'sheet': 'File', 'row': 0, 'field': 'File', 'value': '',
                'message': f'Unrecognised workbook format. Detected sheets: {detected_sheets}. Expected a multi-sheet Universal Master workbook.',
                'error': f'Unrecognised workbook format. Detected sheets: {detected_sheets}.'
            }],
            'valid': False,
        }

    # -- Check required sheets exist -----------------------------------------
    for s_name, sdef in SHEET_DEFS.items():
        if sdef['required'] and s_name not in sheet_map:
            err(s_name, 0, 'Sheet', s_name, f"Required sheet '{s_name}' is missing from workbook.")

    # -- Check columns for present sheets ------------------------------------
    for s_name, ws in sheet_map.items():
        if s_name in SHEET_DEFS:
            expected_cols = [c.strip().lower() for c in SHEET_DEFS[s_name]['columns']]
            first_row = list(ws.iter_rows(values_only=True, max_row=1))
            if first_row and first_row[0]:
                actual_cols = [str(c).strip().lower() for c in first_row[0] if c is not None]
                missing_cols = [c for c in expected_cols if c not in actual_cols]
                if missing_cols:
                    err(s_name, 1, 'Columns', ', '.join(missing_cols),
                        f"Missing required columns in sheet '{s_name}': {', '.join(missing_cols)}")

    if errors:
        return {
            'format': fmt,
            'detected_sheets': detected_sheets,
            'sheet_stats': sheet_stats,
            'total_stats': {'detected_rows': 0, 'valid_rows': 0, 'invalid_rows': len(errors), 'duplicates': 0, 'new': 0, 'existing': 0},
            'stats': {},
            'errors': errors,
            'warnings': warnings,
            'valid': False,
        }

    # -- Parse all sheets ----------------------------------------------------
    def load(name):
        ws = sheet_map.get(name)
        if ws is None:
            return []
        try:
            return parse_sheet(ws)
        except Exception as exc:
            logger.exception('universal_importer: failed to parse sheet %s', name)
            err(name, 0, 'Sheet', name, f"Could not read sheet '{name}': {exc}")
            return []

    depts      = load('Departments')
    ages       = load('Age Groups')
    diags      = load('Diagnoses')
    invs       = load('Investigations')
    params     = load('Parameters')
    diag_depts = load('Diagnosis Departments')
    inv_params = load('Investigation Parameters')
    refs       = load('Reference Ranges')
    diag_invs  = load('Diagnosis Investigations')

    # Update detected row counts
    sheet_stats['Departments']['detected_rows'] = len(depts)
    sheet_stats['Age Groups']['detected_rows'] = len(ages)
    sheet_stats['Diagnoses']['detected_rows'] = len(diags)
    sheet_stats['Investigations']['detected_rows'] = len(invs)
    sheet_stats['Parameters']['detected_rows'] = len(params)
    sheet_stats['Diagnosis Departments']['detected_rows'] = len(diag_depts)
    sheet_stats['Investigation Parameters']['detected_rows'] = len(inv_params)
    sheet_stats['Reference Ranges']['detected_rows'] = len(refs)
    sheet_stats['Diagnosis Investigations']['detected_rows'] = len(diag_invs)

    # -- DB lookup sets & shared resolvers -----------------------------------
    try:
        dept_resolver = DepartmentResolver()
        diag_resolver = DiagnosisResolver()
        inv_resolver  = InvestigationResolver()
        param_resolver = ParameterResolver()
        age_resolver  = AgeGroupResolver()

        db_dept_codes  = set(dept_resolver.by_code.keys())
        db_diag_codes  = set(diag_resolver.by_code.keys())
        db_inv_codes   = set(inv_resolver.by_code.keys())
        db_param_codes = set(param_resolver.by_code.keys())
        db_age_codes   = set(age_resolver.by_code.keys())
    except Exception as exc:
        logger.exception('universal_importer: DB lookup failed during validation')
        return {
            'format': fmt,
            'detected_sheets': detected_sheets,
            'sheet_stats': sheet_stats,
            'total_stats': {},
            'stats': {},
            'warnings': [],
            'errors': [{'sheet': 'Database', 'row': 0, 'field': '', 'value': '',
                        'message': 'Database lookup failed: %s' % exc,
                        'error': 'Database lookup failed: %s' % exc}],
            'valid': False,
        }

    # Working sets include DB codes + codes present in workbook (forward references)
    file_dept_codes  = set(db_dept_codes)
    file_age_codes   = set(db_age_codes)
    file_diag_codes  = set(db_diag_codes)
    file_inv_codes   = set(db_inv_codes)
    file_param_codes = set(db_param_codes)

    seen_inv_param_keys = set()
    seen_ref_keys = set()
    seen_diag_dept_keys = set()
    seen_diag_inv_keys = set()

    # -- 1. Departments ------------------------------------------------------
    for r in depts:
        row = r['_row_index']
        code = clean_code(r.get('department code'), dept_resolver.by_code.keys())
        name = clean_val(r.get('department name'))
        row_has_error = False

        if not code:
            err('Departments', row, 'Department Code', '', 'Department Code is required')
            row_has_error = True
        if not name:
            err('Departments', row, 'Department Name', '', 'Department Name is required')
            row_has_error = True

        if row_has_error:
            sheet_stats['Departments']['invalid_rows'] += 1
        else:
            status, target, is_dup = dept_resolver.resolve(code, name)
            if is_dup:
                warnings.append(f"Departments (Row {row}): Duplicate Department '{code or name}' in workbook; subsequent row will update.")
                sheet_stats['Departments']['duplicates'] += 1

            sheet_stats['Departments']['valid_rows'] += 1
            if status == 'UPDATE':
                sheet_stats['Departments']['existing'] += 1
            else:
                sheet_stats['Departments']['new'] += 1
                dept_resolver.register_placeholder(code, name)

            file_dept_codes.add(code.lower())
            if target and hasattr(target, 'code') and target.code:
                file_dept_codes.add(target.code.lower())

    # -- 2. Age Groups -------------------------------------------------------
    for r in ages:
        row = r['_row_index']
        code = clean_code(r.get('age group code'), age_resolver.by_code.keys())
        name = clean_val(r.get('age group name'))
        a_from = clean_val(r.get('age from'))
        a_to = clean_val(r.get('age to'))
        row_has_error = False

        if not code:
            err('Age Groups', row, 'Age Group Code', '', 'Age Group Code is required')
            row_has_error = True
        if not name:
            err('Age Groups', row, 'Age Group Name', '', 'Age Group Name is required')
            row_has_error = True

        if a_from and not a_from.isdigit():
            err('Age Groups', row, 'Age From', a_from, 'Age From must be a numeric integer')
            row_has_error = True

        if a_to and not a_to.isdigit():
            err('Age Groups', row, 'Age To', a_to, 'Age To must be a numeric integer')
            row_has_error = True

        if row_has_error:
            sheet_stats['Age Groups']['invalid_rows'] += 1
        else:
            status, target, is_dup = age_resolver.resolve(code, name)
            if is_dup:
                warnings.append(f"Age Groups (Row {row}): Duplicate Age Group '{code or name}' in workbook.")
                sheet_stats['Age Groups']['duplicates'] += 1

            sheet_stats['Age Groups']['valid_rows'] += 1
            if status == 'UPDATE':
                sheet_stats['Age Groups']['existing'] += 1
            else:
                sheet_stats['Age Groups']['new'] += 1
                age_resolver.register_placeholder(code, name)

            file_age_codes.add(code.lower())
            if target and hasattr(target, 'code') and target.code:
                file_age_codes.add(target.code.lower())

    # -- 3. Diagnoses --------------------------------------------------------
    for r in diags:
        row = r['_row_index']
        code = clean_code(r.get('diagnosis code'), diag_resolver.by_code.keys())
        name = clean_val(r.get('diagnosis name'))
        row_has_error = False

        if not code:
            err('Diagnoses', row, 'Diagnosis Code', '', 'Diagnosis Code is required')
            row_has_error = True
        if not name:
            err('Diagnoses', row, 'Diagnosis Name', '', 'Diagnosis Name is required')
            row_has_error = True

        if row_has_error:
            sheet_stats['Diagnoses']['invalid_rows'] += 1
        else:
            status, target, is_dup = diag_resolver.resolve(code, name)
            if is_dup:
                warnings.append(f"Diagnoses (Row {row}): Duplicate Diagnosis '{code or name}' in workbook; subsequent row will update.")
                sheet_stats['Diagnoses']['duplicates'] += 1

            sheet_stats['Diagnoses']['valid_rows'] += 1
            if status == 'UPDATE':
                sheet_stats['Diagnoses']['existing'] += 1
            else:
                sheet_stats['Diagnoses']['new'] += 1
                diag_resolver.register_placeholder(code, name)

            file_diag_codes.add(code.lower())
            if target and hasattr(target, 'code') and target.code:
                file_diag_codes.add(target.code.lower())

    # -- 4. Investigations ---------------------------------------------------
    for r in invs:
        row = r['_row_index']
        code = clean_code(r.get('investigation code'), inv_resolver.by_code.keys())
        name = clean_val(r.get('investigation name'))
        row_has_error = False

        if not code:
            err('Investigations', row, 'Investigation Code', '', 'Investigation Code is required')
            row_has_error = True
        if not name:
            err('Investigations', row, 'Investigation Name', '', 'Investigation Name is required')
            row_has_error = True

        if row_has_error:
            sheet_stats['Investigations']['invalid_rows'] += 1
        else:
            status, target, is_dup = inv_resolver.resolve(code, name)
            if is_dup:
                warnings.append(f"Investigations (Row {row}): Duplicate Investigation '{code or name}' in workbook; subsequent row will update.")
                sheet_stats['Investigations']['duplicates'] += 1

            sheet_stats['Investigations']['valid_rows'] += 1
            if status == 'UPDATE':
                sheet_stats['Investigations']['existing'] += 1
            else:
                sheet_stats['Investigations']['new'] += 1
                inv_resolver.register_placeholder(code, name)

            file_inv_codes.add(code.lower())
            if target and hasattr(target, 'code') and target.code:
                file_inv_codes.add(target.code.lower())

    # -- 5. Parameters -------------------------------------------------------
    valid_data_types = {'NUMERIC', 'TEXT', 'CALCULATED', 'SELECT'}
    for r in params:
        row = r['_row_index']
        code = clean_code(r.get('parameter code'), param_resolver.by_code.keys())
        name = clean_val(r.get('parameter name'))
        dtype = clean_val(r.get('data type')).upper() or 'NUMERIC'
        row_has_error = False

        if not code:
            err('Parameters', row, 'Parameter Code', '', 'Parameter Code is required')
            row_has_error = True
        if not name:
            err('Parameters', row, 'Parameter Name', '', 'Parameter Name is required')
            row_has_error = True

        if dtype and dtype not in valid_data_types:
            warnings.append(f"Parameters (Row {row}): Data type '{dtype}' is unusual, default NUMERIC will be applied if needed.")

        if row_has_error:
            sheet_stats['Parameters']['invalid_rows'] += 1
        else:
            status, target, is_dup = param_resolver.resolve(code, name)
            if is_dup:
                warnings.append(f"Parameters (Row {row}): Duplicate Parameter '{code or name}' in workbook; subsequent row will update.")
                sheet_stats['Parameters']['duplicates'] += 1

            sheet_stats['Parameters']['valid_rows'] += 1
            if status == 'UPDATE':
                sheet_stats['Parameters']['existing'] += 1
            else:
                sheet_stats['Parameters']['new'] += 1
                param_resolver.register_placeholder(code, name)

            file_param_codes.add(code.lower())
            if target and hasattr(target, 'code') and target.code:
                file_param_codes.add(target.code.lower())

    # -- 6. Diagnosis Departments --------------------------------------------
    for r in diag_depts:
        row = r['_row_index']
        diag_code = clean_code(r.get('diagnosis code'), file_diag_codes)
        dept_code = clean_code(r.get('department code'), file_dept_codes)
        age_code  = clean_code(r.get('age group code'), file_age_codes)
        row_has_error = False

        if not diag_code:
            err('Diagnosis Departments', row, 'Diagnosis Code', '', 'Diagnosis Code is required')
            row_has_error = True
        elif diag_code.lower() not in file_diag_codes and diag_resolver.get_diagnosis(diag_code) is None:
            err('Diagnosis Departments', row, 'Diagnosis Code', diag_code,
                f"Diagnosis Code '{diag_code}' does not exist")
            row_has_error = True

        if not dept_code:
            err('Diagnosis Departments', row, 'Department Code', '', 'Department Code is required')
            row_has_error = True
        elif dept_code.lower() not in file_dept_codes and dept_resolver.get_department(dept_code) is None:
            err('Diagnosis Departments', row, 'Department Code', dept_code,
                f"Department Code '{dept_code}' does not exist")
            row_has_error = True

        if age_code and age_code.lower() not in file_age_codes and age_resolver.get_age_group(age_code) is None:
            err('Diagnosis Departments', row, 'Age Group Code', age_code,
                f"Age Group Code '{age_code}' does not exist")
            row_has_error = True

        dd_key = (diag_code.lower(), dept_code.lower(), age_code.lower())
        if dd_key in seen_diag_dept_keys:
            warnings.append(f"Diagnosis Departments (Row {row}): Duplicate mapping for ({diag_code}, {dept_code}).")
            sheet_stats['Diagnosis Departments']['duplicates'] += 1
        else:
            seen_diag_dept_keys.add(dd_key)

        if row_has_error:
            sheet_stats['Diagnosis Departments']['invalid_rows'] += 1
        else:
            sheet_stats['Diagnosis Departments']['valid_rows'] += 1
            diag_obj = diag_resolver.get_diagnosis(diag_code)
            dept_obj = dept_resolver.get_department(dept_code)
            ag_obj   = age_resolver.get_age_group(age_code) if age_code else None
            is_existing = False
            if diag_obj and dept_obj:
                is_existing = DiagnosisDepartmentMapping.objects.filter(
                    diagnosis=diag_obj, department=dept_obj, age_group=ag_obj
                ).exists()
            if is_existing:
                sheet_stats['Diagnosis Departments']['existing'] += 1
            else:
                sheet_stats['Diagnosis Departments']['new'] += 1

    # -- 7. Investigation Parameters -----------------------------------------
    for r in inv_params:
        row = r['_row_index']
        inv_code   = clean_code(r.get('investigation code'), file_inv_codes)
        param_code = clean_code(r.get('parameter code'), file_param_codes)
        order_val  = clean_val(r.get('display order'))
        row_has_error = False

        if not inv_code:
            err('Investigation Parameters', row, 'Investigation Code', '', 'Investigation Code is required')
            row_has_error = True
        elif inv_code.lower() not in file_inv_codes and inv_resolver.get_investigation(inv_code) is None:
            err('Investigation Parameters', row, 'Investigation Code', inv_code,
                f"Investigation Code '{inv_code}' does not exist in Investigations or Database")
            row_has_error = True

        if not param_code:
            err('Investigation Parameters', row, 'Parameter Code', '', 'Parameter Code is required')
            row_has_error = True
        elif param_code.lower() not in file_param_codes and param_resolver.get_parameter(param_code) is None:
            err('Investigation Parameters', row, 'Parameter Code', param_code,
                f"Parameter Code '{param_code}' does not exist in Parameters or Database")
            row_has_error = True

        if order_val and not order_val.isdigit():
            err('Investigation Parameters', row, 'Display Order', order_val, 'Display Order must be a numeric integer')
            row_has_error = True

        pair_key = (inv_code.lower(), param_code.lower())
        if pair_key in seen_inv_param_keys:
            warnings.append(f"Investigation Parameters (Row {row}): Duplicate mapping for ({inv_code}, {param_code}).")
            sheet_stats['Investigation Parameters']['duplicates'] += 1
        else:
            seen_inv_param_keys.add(pair_key)

        if row_has_error:
            sheet_stats['Investigation Parameters']['invalid_rows'] += 1
        else:
            sheet_stats['Investigation Parameters']['valid_rows'] += 1
            inv_obj   = inv_resolver.get_investigation(inv_code)
            param_obj = param_resolver.get_parameter(param_code)
            is_existing = False
            if inv_obj and param_obj:
                is_existing = InvestigationParameter.objects.filter(
                    investigation=inv_obj, parameter=param_obj
                ).exists() or InvestigationParameter.objects.filter(
                    investigation=inv_obj, code=param_obj.code
                ).exists()
            if is_existing:
                sheet_stats['Investigation Parameters']['existing'] += 1
            else:
                sheet_stats['Investigation Parameters']['new'] += 1

    # -- 8. Reference Ranges -------------------------------------------------
    for r in refs:
        row = r['_row_index']
        inv_code   = clean_code(r.get('investigation code'), file_inv_codes)
        param_code = clean_code(r.get('parameter code'), file_param_codes)
        age_code   = clean_code(r.get('age group code'), file_age_codes)
        diag_code  = clean_code(r.get('diagnosis code'), file_diag_codes)
        gender     = clean_val(r.get('gender')) or 'All'
        min_v      = clean_val(r.get('min value'))
        max_v      = clean_val(r.get('max value'))
        n_val      = clean_val(r.get('normal value'))
        row_has_error = False

        if not inv_code:
            err('Reference Ranges', row, 'Investigation Code', '', 'Investigation Code is required')
            row_has_error = True
        elif inv_code.lower() not in file_inv_codes and inv_resolver.get_investigation(inv_code) is None:
            err('Reference Ranges', row, 'Investigation Code', inv_code,
                f"Investigation Code '{inv_code}' does not exist")
            row_has_error = True

        if not param_code:
            err('Reference Ranges', row, 'Parameter Code', '', 'Parameter Code is required')
            row_has_error = True
        elif param_code.lower() not in file_param_codes and param_resolver.get_parameter(param_code) is None:
            err('Reference Ranges', row, 'Parameter Code', param_code,
                f"Parameter Code '{param_code}' does not exist")
            row_has_error = True

        if age_code and age_code.lower() not in file_age_codes and age_resolver.get_age_group(age_code) is None:
            err('Reference Ranges', row, 'Age Group Code', age_code,
                f"Age Group Code '{age_code}' does not exist")
            row_has_error = True

        if diag_code and diag_code.lower() not in file_diag_codes and diag_resolver.get_diagnosis(diag_code) is None:
            err('Reference Ranges', row, 'Diagnosis Code', diag_code,
                f"Diagnosis Code '{diag_code}' does not exist")
            row_has_error = True

        if min_v and max_v:
            try:
                dec_min = Decimal(min_v)
                dec_max = Decimal(max_v)
                if dec_min > dec_max:
                    err('Reference Ranges', row, 'Min Value', min_v,
                        f'Min Value ({min_v}) cannot be greater than Max Value ({max_v})')
                    row_has_error = True
            except Exception:
                err('Reference Ranges', row, 'Min/Max Value', f"{min_v}/{max_v}", 'Numeric values required for Min and Max')
                row_has_error = True

        ref_key = (inv_code.lower(), param_code.lower(), age_code.lower(), gender.lower(), diag_code.lower())
        if ref_key in seen_ref_keys:
            warnings.append(f"Reference Ranges (Row {row}): Duplicate range key in workbook; will be updated safely.")
            sheet_stats['Reference Ranges']['duplicates'] += 1
        else:
            seen_ref_keys.add(ref_key)

        if row_has_error:
            sheet_stats['Reference Ranges']['invalid_rows'] += 1
        else:
            sheet_stats['Reference Ranges']['valid_rows'] += 1
            inv_obj   = inv_resolver.get_investigation(inv_code)
            param_obj = param_resolver.get_parameter(param_code)
            ag_obj    = age_resolver.get_age_group(age_code) if age_code else None
            diag_obj  = diag_resolver.get_diagnosis(diag_code) if diag_code else None
            is_existing = False
            if inv_obj and param_obj:
                ip = InvestigationParameter.objects.filter(
                    investigation=inv_obj, parameter=param_obj
                ).first() or InvestigationParameter.objects.filter(
                    investigation=inv_obj, code=param_obj.code
                ).first()
                if ip:
                    is_existing = ParameterReferenceRange.objects.filter(
                        investigation_parameter=ip, age_group=ag_obj,
                        gender=gender, diagnosis=diag_obj
                    ).exists()
            if is_existing:
                sheet_stats['Reference Ranges']['existing'] += 1
            else:
                sheet_stats['Reference Ranges']['new'] += 1

    # -- 9. Diagnosis Investigations -----------------------------------------
    for r in diag_invs:
        row = r['_row_index']
        diag_code = clean_code(r.get('diagnosis code'), file_diag_codes)
        inv_code  = clean_code(r.get('investigation code'), file_inv_codes)
        age_code  = clean_code(r.get('age group code'), file_age_codes)
        row_has_error = False

        if not diag_code:
            err('Diagnosis Investigations', row, 'Diagnosis Code', '', 'Diagnosis Code is required')
            row_has_error = True
        elif diag_code.lower() not in file_diag_codes and diag_resolver.get_diagnosis(diag_code) is None:
            err('Diagnosis Investigations', row, 'Diagnosis Code', diag_code,
                f"Diagnosis Code '{diag_code}' does not exist")
            row_has_error = True

        if not inv_code:
            err('Diagnosis Investigations', row, 'Investigation Code', '', 'Investigation Code is required')
            row_has_error = True
        elif inv_code.lower() not in file_inv_codes and inv_resolver.get_investigation(inv_code) is None:
            err('Diagnosis Investigations', row, 'Investigation Code', inv_code,
                f"Investigation Code '{inv_code}' does not exist")
            row_has_error = True

        if age_code and age_code.lower() not in file_age_codes and age_resolver.get_age_group(age_code) is None:
            err('Diagnosis Investigations', row, 'Age Group Code', age_code,
                f"Age Group Code '{age_code}' does not exist")
            row_has_error = True

        di_key = (diag_code.lower(), inv_code.lower(), age_code.lower())
        if di_key in seen_diag_inv_keys:
            warnings.append(f"Diagnosis Investigations (Row {row}): Duplicate mapping for ({diag_code}, {inv_code}).")
            sheet_stats['Diagnosis Investigations']['duplicates'] += 1
        else:
            seen_diag_inv_keys.add(di_key)

        if row_has_error:
            sheet_stats['Diagnosis Investigations']['invalid_rows'] += 1
        else:
            sheet_stats['Diagnosis Investigations']['valid_rows'] += 1
            diag_obj = diag_resolver.get_diagnosis(diag_code)
            inv_obj  = inv_resolver.get_investigation(inv_code)
            ag_obj   = age_resolver.get_age_group(age_code) if age_code else None
            is_existing = False
            if diag_obj and inv_obj:
                is_existing = DiagnosisInvestigationMap.objects.filter(
                    diagnosis=diag_obj, investigation=inv_obj, age_group=ag_obj
                ).exists()
            if is_existing:
                sheet_stats['Diagnosis Investigations']['existing'] += 1
            else:
                sheet_stats['Diagnosis Investigations']['new'] += 1

    # Aggregate overall stats
    total_stats = {
        'detected_rows': sum(s['detected_rows'] for s in sheet_stats.values()),
        'valid_rows': sum(s['valid_rows'] for s in sheet_stats.values()),
        'invalid_rows': sum(s['invalid_rows'] for s in sheet_stats.values()),
        'duplicates': sum(s['duplicates'] for s in sheet_stats.values()),
        'new': sum(s['new'] for s in sheet_stats.values()),
        'existing': sum(s['existing'] for s in sheet_stats.values()),
    }

    # Legacy stats dict key mapping for backward-compatibility with UI
    compat_stats = {
        'departments': {'existing': sheet_stats['Departments']['existing'], 'new': sheet_stats['Departments']['new']},
        'age_groups': {'existing': sheet_stats['Age Groups']['existing'], 'new': sheet_stats['Age Groups']['new']},
        'diagnoses': {'existing': sheet_stats['Diagnoses']['existing'], 'new': sheet_stats['Diagnoses']['new']},
        'investigations': {'existing': sheet_stats['Investigations']['existing'], 'new': sheet_stats['Investigations']['new']},
        'parameters': {'existing': sheet_stats['Parameters']['existing'], 'new': sheet_stats['Parameters']['new']},
        'diag_dept_mappings': {'existing': sheet_stats['Diagnosis Departments']['existing'], 'new': sheet_stats['Diagnosis Departments']['new']},
        'inv_param_mappings': {'existing': sheet_stats['Investigation Parameters']['existing'], 'new': sheet_stats['Investigation Parameters']['new']},
        'reference_ranges': {'existing': sheet_stats['Reference Ranges']['existing'], 'new': sheet_stats['Reference Ranges']['new']},
        'diag_inv_mappings': {'existing': sheet_stats['Diagnosis Investigations']['existing'], 'new': sheet_stats['Diagnosis Investigations']['new']},
    }

    valid = len(errors) == 0
    logger.info('universal_importer: validation complete valid=%s errors=%d detected=%d',
                valid, len(errors), total_stats['detected_rows'])

    return {
        'format': fmt,
        'detected_sheets': detected_sheets,
        'sheet_stats': sheet_stats,
        'total_stats': total_stats,
        'stats': compat_stats,
        'errors': errors[:200],
        'warnings': warnings[:50],
        'valid': valid,
    }


# ---------------------------------------------------------------------------
# IMPORT (DB writes inside transaction.atomic)
# ---------------------------------------------------------------------------

def import_universal_master(file_obj, admin_user=None):
    """
    Import a validated multi-sheet Universal Master workbook.
    Any error rolls back the entire transaction.
    """
    try:
        wb = load_workbook(filename=file_obj, data_only=True)
    except Exception as exc:
        logger.exception('universal_importer: failed to open workbook for import')
        raise ValueError('Could not open Excel file: %s' % exc)

    sheet_map = get_workbook_sheet_map(wb)

    def load(name):
        ws = sheet_map.get(name)
        if ws is None:
            return []
        return parse_sheet(ws)

    depts      = load('Departments')
    ages       = load('Age Groups')
    diags      = load('Diagnoses')
    invs       = load('Investigations')
    params     = load('Parameters')
    diag_depts = load('Diagnosis Departments')
    inv_params = load('Investigation Parameters')
    refs       = load('Reference Ranges')
    diag_invs  = load('Diagnosis Investigations')

    counts = {
        'departments_created': 0,        'departments_updated': 0,        'departments_skipped': 0,
        'age_groups_created': 0,         'age_groups_updated': 0,         'age_groups_skipped': 0,
        'diagnoses_created': 0,          'diagnoses_updated': 0,          'diagnoses_skipped': 0,
        'investigations_created': 0,     'investigations_updated': 0,     'investigations_skipped': 0,
        'parameters_created': 0,         'parameters_updated': 0,         'parameters_skipped': 0,
        'diag_dept_mappings_created': 0, 'diag_dept_mappings_updated': 0, 'diag_dept_mappings_skipped': 0,
        'inv_params_created': 0,         'inv_params_updated': 0,         'inv_params_skipped': 0,
        'ref_ranges_created': 0,         'ref_ranges_updated': 0,         'ref_ranges_skipped': 0,
        'diag_inv_mappings_created': 0,  'diag_inv_mappings_updated': 0,  'diag_inv_mappings_skipped': 0,
        'mappings_created': 0,           'mappings_updated': 0,
        'skipped': 0,
        'errors': 0,
    }

    with transaction.atomic():
        try:
            dept_resolver = DepartmentResolver()
            diag_resolver = DiagnosisResolver()
            inv_resolver  = InvestigationResolver()
            param_resolver = ParameterResolver()
            age_resolver  = AgeGroupResolver()

            # 1. Departments
            for r in depts:
                row_idx = r.get('_row_index', 0)
                code   = clean_code(r.get('department code'), dept_resolver.by_code.keys())
                name   = clean_val(r.get('department name'))
                status = is_active_bool(r.get('status'))
                if code and name:
                    dept, created = dept_resolver.upsert(code, name, is_active=status, row_index=row_idx)
                    if created:
                        counts['departments_created'] += 1
                    else:
                        counts['departments_updated'] += 1
                else:
                    counts['departments_skipped'] += 1
                    counts['skipped'] += 1

            # 2. Age Groups
            for r in ages:
                row_idx = r.get('_row_index', 0)
                code    = clean_code(r.get('age group code'), age_resolver.by_code.keys())
                ag_name = clean_val(r.get('age group name'))
                a_from  = clean_val(r.get('age from'))
                f_unit  = clean_val(r.get('age from unit')) or 'Years'
                a_to    = clean_val(r.get('age to'))
                t_unit  = clean_val(r.get('age to unit')) or 'Years'
                gender  = clean_val(r.get('gender')) or 'All'
                status  = is_active_bool(r.get('status'))
                if code and ag_name:
                    ag, created = age_resolver.upsert(
                        code, ag_name, a_from=a_from, f_unit=f_unit,
                        a_to=a_to, t_unit=t_unit, gender=gender,
                        is_active=status, row_index=row_idx
                    )
                    if created:
                        counts['age_groups_created'] += 1
                    else:
                        counts['age_groups_updated'] += 1
                else:
                    counts['age_groups_skipped'] += 1
                    counts['skipped'] += 1

            # 3. Diagnoses
            for r in diags:
                row_idx = r.get('_row_index', 0)
                code   = clean_code(r.get('diagnosis code'), diag_resolver.by_code.keys())
                name   = clean_val(r.get('diagnosis name'))
                status = is_active_bool(r.get('status'))
                if code and name:
                    diag, created = diag_resolver.upsert(code, name, is_active=status, row_index=row_idx)
                    if created:
                        counts['diagnoses_created'] += 1
                    else:
                        counts['diagnoses_updated'] += 1
                else:
                    counts['diagnoses_skipped'] += 1
                    counts['skipped'] += 1

            # 4. Investigations
            for r in invs:
                row_idx   = r.get('_row_index', 0)
                code      = clean_code(r.get('investigation code'), inv_resolver.by_code.keys())
                name      = clean_val(r.get('investigation name'))
                dept_name = clean_val(r.get('department'))
                samp_name = clean_val(r.get('sample type'))
                is_panel  = clean_val(r.get('is panel')).lower() in ['yes', 'true', '1', 'y']
                status    = is_active_bool(r.get('status'))
                if code and name:
                    inv, created = inv_resolver.upsert(
                        code, name, dept_name=dept_name, samp_name=samp_name,
                        is_panel=is_panel, is_active=status, row_index=row_idx
                    )
                    if created:
                        counts['investigations_created'] += 1
                    else:
                        counts['investigations_updated'] += 1
                else:
                    counts['investigations_skipped'] += 1
                    counts['skipped'] += 1

            # 5. Parameters
            for r in params:
                row_idx = r.get('_row_index', 0)
                code   = clean_code(r.get('parameter code'), param_resolver.by_code.keys())
                name   = clean_val(r.get('parameter name'))
                dtype  = clean_val(r.get('data type')).upper() or 'NUMERIC'
                unit   = clean_val(r.get('default unit'))
                status = is_active_bool(r.get('status'))
                if code and name:
                    param, created = param_resolver.upsert(
                        code, name, data_type=dtype, default_unit=unit,
                        is_active=status, row_index=row_idx
                    )
                    if created:
                        counts['parameters_created'] += 1
                    else:
                        counts['parameters_updated'] += 1
                else:
                    counts['parameters_skipped'] += 1
                    counts['skipped'] += 1

            # 6. Diagnosis Departments
            for r in diag_depts:
                row_idx   = r.get('_row_index', 0)
                diag_code = clean_code(r.get('diagnosis code'), diag_resolver.by_code.keys())
                dept_code = clean_code(r.get('department code'), dept_resolver.by_code.keys())
                age_code  = clean_code(r.get('age group code'), age_resolver.by_code.keys())
                status_val = 'Active' if is_active_bool(r.get('status')) else 'Inactive'
                if diag_code and dept_code:
                    diag = diag_resolver.get_diagnosis(diag_code)
                    dept = dept_resolver.get_department(dept_code)
                    ag   = age_resolver.get_age_group(age_code) if age_code else None
                    if diag and dept:
                        existing = DiagnosisDepartmentMapping.objects.filter(
                            diagnosis=diag, department=dept, age_group=ag
                        ).first()
                        if existing:
                            existing.status = status_val
                            existing.save()
                            counts['diag_dept_mappings_updated'] += 1
                            counts['mappings_updated'] += 1
                        else:
                            DiagnosisDepartmentMapping.objects.create(
                                diagnosis=diag, department=dept, age_group=ag,
                                status=status_val
                            )
                            counts['diag_dept_mappings_created'] += 1
                            counts['mappings_created'] += 1
                    else:
                        counts['diag_dept_mappings_skipped'] += 1
                        counts['skipped'] += 1
                else:
                    counts['diag_dept_mappings_skipped'] += 1
                    counts['skipped'] += 1

            # 7. Investigation Parameters
            for r in inv_params:
                row_idx    = r.get('_row_index', 0)
                inv_code   = clean_code(r.get('investigation code'), inv_resolver.by_code.keys())
                param_code = clean_code(r.get('parameter code'), param_resolver.by_code.keys())
                order      = clean_val(r.get('display order'))
                status     = is_active_bool(r.get('status'))
                if inv_code and param_code:
                    inv   = inv_resolver.get_investigation(inv_code)
                    param = param_resolver.get_parameter(param_code)
                    if inv and param:
                        existing = (
                            InvestigationParameter.objects.filter(investigation=inv, parameter=param).first() or
                            InvestigationParameter.objects.filter(investigation=inv, code=param.code).first()
                        )
                        disp_order = int(order) if order and order.isdigit() else 0
                        res_type = 'Numeric' if param.data_type == 'NUMERIC' else 'Text'
                        if existing:
                            existing.parameter = param
                            existing.name = param.name
                            existing.code = param.code
                            existing.result_type = res_type
                            if param.default_unit and not existing.unit:
                                existing.unit = param.default_unit
                            existing.display_order = disp_order
                            existing.is_active = status
                            existing.save()
                            counts['inv_params_updated'] += 1
                            counts['mappings_updated'] += 1
                        else:
                            InvestigationParameter.objects.create(
                                investigation=inv, parameter=param,
                                name=param.name, code=param.code,
                                result_type=res_type,
                                unit=param.default_unit or '',
                                display_order=disp_order, is_active=status
                            )
                            counts['inv_params_created'] += 1
                            counts['mappings_created'] += 1
                    else:
                        counts['inv_params_skipped'] += 1
                        counts['skipped'] += 1
                else:
                    counts['inv_params_skipped'] += 1
                    counts['skipped'] += 1

            # 8. Reference Ranges
            for r in refs:
                row_idx    = r.get('_row_index', 0)
                inv_code   = clean_code(r.get('investigation code'), inv_resolver.by_code.keys())
                param_code = clean_code(r.get('parameter code'), param_resolver.by_code.keys())
                age_code   = clean_code(r.get('age group code'), age_resolver.by_code.keys())
                gender     = clean_val(r.get('gender')) or 'All'
                diag_code  = clean_code(r.get('diagnosis code'), diag_resolver.by_code.keys())
                min_v  = get_decimal(r.get('min value'))
                max_v  = get_decimal(r.get('max value'))
                n_val  = clean_val(r.get('normal value'))
                unit   = clean_val(r.get('unit'))
                meth   = clean_val(r.get('method'))
                rem    = clean_val(r.get('remarks'))
                status = is_active_bool(r.get('status'))
                if inv_code and param_code:
                    inv = inv_resolver.get_investigation(inv_code)
                    param = param_resolver.get_parameter(param_code)
                    if inv and param:
                        ip = (
                            InvestigationParameter.objects.filter(investigation=inv, parameter=param).first() or
                            InvestigationParameter.objects.filter(investigation=inv, code=param.code).first()
                        )
                        if not ip:
                            res_type = 'Numeric' if param.data_type == 'NUMERIC' else 'Text'
                            ip = InvestigationParameter.objects.create(
                                investigation=inv, parameter=param,
                                name=param.name, code=param.code,
                                result_type=res_type,
                                unit=unit or param.default_unit or '',
                                is_active=True
                            )
                        ag   = age_resolver.get_age_group(age_code) if age_code else None
                        diag = diag_resolver.get_diagnosis(diag_code) if diag_code else None

                        range_type = 'Numeric' if (min_v is not None or max_v is not None) else ('Text' if n_val else 'None')

                        defaults = {
                            'min_value': min_v,
                            'max_value': max_v,
                            'reference_text': n_val or None,
                            'unit': unit,
                            'method': meth,
                            'remarks': rem,
                            'is_active': status,
                            'range_type': range_type,
                        }
                        existing = ParameterReferenceRange.objects.filter(
                            investigation_parameter=ip, age_group=ag,
                            gender=gender, diagnosis=diag
                        )
                        if existing.exists():
                            first = existing.first()
                            for k, v in defaults.items():
                                setattr(first, k, v)
                            first.save()
                            if existing.count() > 1:
                                existing.exclude(id=first.id).delete()
                            counts['ref_ranges_updated'] += 1
                        else:
                            ParameterReferenceRange.objects.create(
                                investigation_parameter=ip, age_group=ag,
                                gender=gender, diagnosis=diag, **defaults
                            )
                            counts['ref_ranges_created'] += 1

                        # Sync fallback strings and unit to InvestigationParameter
                        ref_str = ''
                        if min_v is not None and max_v is not None:
                            ref_str = f"{min_v} - {max_v}"
                        elif min_v is not None:
                            ref_str = f">= {min_v}"
                        elif max_v is not None:
                            ref_str = f"<= {max_v}"
                        elif n_val:
                            ref_str = n_val

                        ip_updated = False
                        if ref_str:
                            if gender == 'Male':
                                if not ip.male_reference_range:
                                    ip.male_reference_range = ref_str
                                    ip_updated = True
                            elif gender == 'Female':
                                if not ip.female_reference_range:
                                    ip.female_reference_range = ref_str
                                    ip_updated = True
                            else:
                                if not ip.reference_range:
                                    ip.reference_range = ref_str
                                    ip_updated = True
                        if unit and not ip.unit:
                            ip.unit = unit
                            ip_updated = True
                        if ip_updated:
                            ip.save()
                    else:
                        counts['ref_ranges_skipped'] += 1
                        counts['skipped'] += 1
                else:
                    counts['ref_ranges_skipped'] += 1
                    counts['skipped'] += 1

            # 9. Diagnosis Investigations
            for r in diag_invs:
                row_idx   = r.get('_row_index', 0)
                diag_code = clean_code(r.get('diagnosis code'), diag_resolver.by_code.keys())
                inv_code  = clean_code(r.get('investigation code'), inv_resolver.by_code.keys())
                age_code  = clean_code(r.get('age group code'), age_resolver.by_code.keys())
                status    = is_active_bool(r.get('status'))
                if diag_code and inv_code:
                    diag = diag_resolver.get_diagnosis(diag_code)
                    inv  = inv_resolver.get_investigation(inv_code)
                    ag   = age_resolver.get_age_group(age_code) if age_code else None
                    if diag and inv:
                        existing = DiagnosisInvestigationMap.objects.filter(
                            diagnosis=diag, investigation=inv, age_group=ag
                        ).first()
                        if existing:
                            existing.is_active = status
                            existing.save()
                            counts['diag_inv_mappings_updated'] += 1
                            counts['mappings_updated'] += 1
                        else:
                            DiagnosisInvestigationMap.objects.create(
                                diagnosis=diag, investigation=inv, age_group=ag,
                                is_active=status
                            )
                            counts['diag_inv_mappings_created'] += 1
                            counts['mappings_created'] += 1
                    else:
                        counts['diag_inv_mappings_skipped'] += 1
                        counts['skipped'] += 1
                else:
                    counts['diag_inv_mappings_skipped'] += 1
                    counts['skipped'] += 1

        except UniversalImportError:
            logger.exception("universal_importer: UniversalImportError during commit")
            raise
        except Exception as e:
            logger.exception("universal_importer: unexpected error during commit")
            err_msg = str(e)
            is_unique = 'UNIQUE constraint' in err_msg or isinstance(e, IntegrityError)
            reason = 'A record with this unique identity already exists.' if is_unique else err_msg
            action = 'Existing record should be updated instead of created.' if is_unique else 'Please review import workbook and try again.'
            raise UniversalImportError(
                message=f"Import failed — no partial changes were committed. {err_msg}",
                entity='Master Record',
                entity_name='',
                reason=reason,
                action=action
            ) from e

    counts['total_created'] = sum(v for k, v in counts.items() if k.endswith('_created'))
    counts['total_updated'] = sum(v for k, v in counts.items() if k.endswith('_updated'))
    counts['total_skipped'] = sum(v for k, v in counts.items() if k.endswith('_skipped'))

    logger.info('universal_importer: import complete counts=%s', counts)
    return counts



# Alias for views importing commit_import
commit_import = import_universal_master


# ---------------------------------------------------------------------------
# EXPORT (produces exactly what the importer expects — round-trip safe)
# ---------------------------------------------------------------------------

def export_universal_master():
    """Export all Lab Master data as a multi-sheet workbook."""
    wb = Workbook(write_only=False)
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # 1. Departments
    ws = wb.create_sheet('Departments')
    ws.append(SHEET_DEFS['Departments']['columns'])
    for d in Department.objects.order_by('name'):
        ws.append([d.name, d.code, 'Active' if d.is_active else 'Inactive'])
    _style_header_row(ws)

    # 2. Diagnoses
    ws = wb.create_sheet('Diagnoses')
    ws.append(SHEET_DEFS['Diagnoses']['columns'])
    for d in Diagnosis.objects.order_by('name'):
        ws.append([d.name, d.code, 'Active' if d.is_active else 'Inactive'])
    _style_header_row(ws)

    # 3. Investigations
    ws = wb.create_sheet('Investigations')
    ws.append(SHEET_DEFS['Investigations']['columns'])
    for i in Investigation.objects.select_related('department', 'sample_type').order_by('name'):
        ws.append([
            i.name, i.code,
            i.department.name if i.department else '',
            i.sample_type.name if i.sample_type else '',
            'Yes' if i.is_panel else 'No',
            'Active' if i.is_active else 'Inactive',
        ])
    _style_header_row(ws)

    # 4. Parameters
    ws = wb.create_sheet('Parameters')
    ws.append(SHEET_DEFS['Parameters']['columns'])
    for p in Parameter.objects.order_by('name'):
        ws.append([p.name, p.code, p.data_type, p.default_unit or '',
                   'Active' if p.is_active else 'Inactive'])
    _style_header_row(ws)

    # 5. Age Groups
    ws = wb.create_sheet('Age Groups')
    ws.append(SHEET_DEFS['Age Groups']['columns'])
    for a in AgeGroup.objects.order_by('code'):
        ws.append([
            a.code, a.label,
            a.min_age_value if a.min_age_value is not None else '',
            a.min_age_unit or 'Years',
            a.max_age_value if a.max_age_value is not None else '',
            a.max_age_unit or 'Years',
            a.gender,
            'Active' if a.is_active else 'Inactive',
        ])
    _style_header_row(ws)

    # 6. Investigation Parameters
    ws = wb.create_sheet('Investigation Parameters')
    ws.append(SHEET_DEFS['Investigation Parameters']['columns'])
    for ip in InvestigationParameter.objects.select_related(
        'investigation', 'parameter'
    ).order_by('investigation__code', 'display_order'):
        if ip.parameter:
            ws.append([
                ip.investigation.code, ip.parameter.code,
                ip.display_order,
                'Active' if ip.is_active else 'Inactive',
            ])
    _style_header_row(ws)

    # 7. Reference Ranges
    ws = wb.create_sheet('Reference Ranges')
    ws.append(SHEET_DEFS['Reference Ranges']['columns'])
    for r in ParameterReferenceRange.objects.select_related(
        'investigation_parameter__investigation',
        'investigation_parameter__parameter',
        'age_group', 'diagnosis',
    ).order_by('investigation_parameter__investigation__code'):
        ip = r.investigation_parameter
        ws.append([
            ip.investigation.code,
            ip.parameter.code if ip.parameter else '',
            r.age_group.code if r.age_group else '',
            r.gender,
            r.diagnosis.code if r.diagnosis else '',
            r.min_value, r.max_value,
            r.reference_text or '', r.unit or '',
            r.method or '', r.remarks or '',
            'Active' if r.is_active else 'Inactive',
        ])
    _style_header_row(ws)

    # 8. Diagnosis Departments
    ws = wb.create_sheet('Diagnosis Departments')
    ws.append(SHEET_DEFS['Diagnosis Departments']['columns'])
    for dd in DiagnosisDepartmentMapping.objects.select_related(
        'diagnosis', 'department', 'age_group'
    ).order_by('diagnosis__code'):
        ws.append([
            dd.diagnosis.code if dd.diagnosis else '',
            dd.department.code if dd.department else '',
            dd.age_group.code if dd.age_group else '',
            dd.status or 'Active',
        ])
    _style_header_row(ws)

    # 9. Diagnosis Investigations
    ws = wb.create_sheet('Diagnosis Investigations')
    ws.append(SHEET_DEFS['Diagnosis Investigations']['columns'])
    for di in DiagnosisInvestigationMap.objects.select_related(
        'diagnosis', 'investigation', 'age_group'
    ).order_by('diagnosis__code'):
        ws.append([
            di.diagnosis.code if di.diagnosis else '',
            di.investigation.code if di.investigation else '',
            di.age_group.code if di.age_group else '',
            'Active' if di.is_active else 'Inactive',
        ])
    _style_header_row(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_blank_template():
    """Produce a blank multi-sheet template with headers only."""
    wb = Workbook(write_only=False)
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(sheet_name)
        ws.append(SHEET_DEFS[sheet_name]['columns'])
        _style_header_row(ws)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

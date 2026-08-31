import json
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_exempt
from django.core.paginator import Paginator
from django.db import transaction
from apps.patients.models import Department
from apps.lab.models import Diagnosis, DiagnosisDepartmentMapping, DiagnosisDepartmentMappingImportHistory

def diagnosis_department_mapping_view(request):
    if request.GET.get('export') == 'excel':
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Mapping')
        ws.append(['Department', 'Diagnosis', 'Status', 'Mapped Date'])
        
        qs = DiagnosisDepartmentMapping.objects.select_related('department', 'diagnosis')
        dept_filter = request.GET.get('department_id')
        if dept_filter:
            qs = qs.filter(department_id=dept_filter)
            
        for r in qs:
            ws.append([
                r.department.name,
                r.diagnosis.name,
                r.status,
                r.created_at.strftime('%d-%m-%Y %I:%M %p') if r.created_at else ''
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="diagnosis_department_mapping_export.xlsx"'
        return response

    context = {
        'departments': Department.objects.filter(is_active=True).order_by('name')
    }
    return render(request, 'lab/master/diagnosis_department_mapping.html', context)

def api_diagnosis_department_map_list(request):
    dept_id = request.GET.get('department_id')
    status = request.GET.get('status')
    q = request.GET.get('q', '').strip()
    page_num = int(request.GET.get('page', 1))
    per_page = int(request.GET.get('per_page', 10))
    
    qs = DiagnosisDepartmentMapping.objects.select_related('department', 'diagnosis')
    
    if dept_id:
        qs = qs.filter(department_id=dept_id)
    if status and status != 'All':
        qs = qs.filter(status=status)
    if q:
        qs = qs.filter(diagnosis__name__icontains=q)
        
    paginator = Paginator(qs, per_page)
    page_obj = paginator.get_page(page_num)
    
    results = []
    for r in page_obj:
        results.append({
            'id': r.id,
            'department_id': r.department_id,
            'department_name': r.department.name,
            'diagnosis_id': r.diagnosis_id,
            'diagnosis_name': r.diagnosis.name,
            'status': r.status,
            'mapped_date': r.created_at.strftime('%d-%m-%Y %I:%M %p') if r.created_at else '',
            'mapped_by': 'Admin' # Default as per reference for now
        })
        
    return JsonResponse({
        'status': 'success',
        'mappings': results,
        'total': paginator.count,
        'pages': paginator.num_pages,
        'current_page': page_num
    })

@csrf_exempt
def api_diagnosis_department_map_save(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            department_id = data.get('department_id')
            diagnoses_ids = data.get('diagnoses_ids', [])
            
            if not department_id or not diagnoses_ids:
                return JsonResponse({'status': 'error', 'message': 'Missing department_id or diagnoses_ids'})
                
            department = Department.objects.get(id=department_id)
            
            created_count = 0
            with transaction.atomic():
                for diag_id in diagnoses_ids:
                    obj, created = DiagnosisDepartmentMapping.objects.get_or_create(
                        department_id=department_id,
                        diagnosis_id=diag_id,
                        defaults={'status': 'Active'}
                    )
                    if created:
                        created_count += 1
            
            return JsonResponse({'status': 'success', 'message': f'Successfully mapped {created_count} diagnoses.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method.'})

@csrf_exempt
def api_diagnosis_department_map_toggle(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            mapping_id = data.get('mapping_id')
            status = data.get('status')
            if mapping_id and status:
                mapping = DiagnosisDepartmentMapping.objects.get(id=mapping_id)
                mapping.status = status
                mapping.save()
                return JsonResponse({'status': 'success', 'message': 'Status updated.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    return JsonResponse({'status': 'error', 'message': 'Invalid method.'})

def api_diagnosis_department_map_unmapped_diagnoses(request):
    dept_id = request.GET.get('department_id')
    if not dept_id:
        return JsonResponse({'status': 'error', 'message': 'Missing department_id'})
        
    mapped_diag_ids = DiagnosisDepartmentMapping.objects.filter(department_id=dept_id).values_list('diagnosis_id', flat=True)
    unmapped = Diagnosis.objects.exclude(id__in=mapped_diag_ids).filter(is_active=True).values('id', 'name', 'code')
    
    return JsonResponse({
        'status': 'success',
        'diagnoses': list(unmapped)
    })

def api_diagnosis_department_map_mapped_diagnoses(request):
    dept_id = request.GET.get('department_id')
    if not dept_id:
        return JsonResponse({'status': 'error', 'message': 'Missing department_id'})
        
    mapped = DiagnosisDepartmentMapping.objects.filter(department_id=dept_id, status='Active').select_related('diagnosis')
    results = [{'id': m.diagnosis.id, 'name': m.diagnosis.name, 'code': m.diagnosis.code} for m in mapped if m.diagnosis.is_active]
    
    return JsonResponse({
        'status': 'success',
        'diagnoses': results
    })

def api_diag_dept_map_download_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diagnosis-Department Mapping"
    
    headers = ["Department", "Diagnosis", "Status"]
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 25
        
    dv = DataValidation(type="list", formula1='"Active,Inactive"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C1000")
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="diagnosis_department_mapping_template.xlsx"'
    wb.save(response)
    return response

@csrf_exempt
def api_diag_dept_map_preview(request):
    if request.method == 'POST':
        file = request.FILES.get('file')
        if not file:
            return JsonResponse({'status': 'error', 'message': 'No file uploaded'})
            
        try:
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
            
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return JsonResponse({'status': 'error', 'message': 'File is empty or missing headers'})
                
            headers = [str(h).strip().lower() for h in rows[0] if h]
            if 'department' not in headers or 'diagnosis' not in headers:
                return JsonResponse({'status': 'error', 'message': 'Missing required columns: Department, Diagnosis'})
                
            dept_idx = headers.index('department')
            diag_idx = headers.index('diagnosis')
            status_idx = headers.index('status') if 'status' in headers else -1
            
            preview_data = []
            valid_count = 0
            invalid_count = 0
            
            dept_cache = {d.name.lower(): d for d in Department.objects.all()}
            diag_cache = {d.name.lower(): d for d in Diagnosis.objects.all()}
            existing_maps = set(DiagnosisDepartmentMapping.objects.values_list('department_id', 'diagnosis_id'))
            
            for i, row in enumerate(rows[1:], start=2):
                if not any(row): continue
                
                dept_name = str(row[dept_idx]).strip() if len(row) > dept_idx and row[dept_idx] else ''
                diag_name = str(row[diag_idx]).strip() if len(row) > diag_idx and row[diag_idx] else ''
                status = str(row[status_idx]).strip() if status_idx != -1 and len(row) > status_idx and row[status_idx] else 'Active'
                
                if status not in ['Active', 'Inactive']:
                    status = 'Active'
                    
                is_valid = True
                errors = []
                
                dept_obj = dept_cache.get(dept_name.lower())
                diag_obj = diag_cache.get(diag_name.lower())
                
                if not dept_name or not diag_name:
                    is_valid = False
                    errors.append("Department and Diagnosis are required")
                else:
                    if not dept_obj:
                        is_valid = False
                        errors.append(f"Department '{dept_name}' not found")
                    if not diag_obj:
                        is_valid = False
                        errors.append(f"Diagnosis '{diag_name}' not found")
                        
                if is_valid and dept_obj and diag_obj:
                    if (dept_obj.id, diag_obj.id) in existing_maps:
                        is_valid = False
                        errors.append("Mapping already exists")
                        
                if is_valid:
                    valid_count += 1
                else:
                    invalid_count += 1
                    
                preview_data.append({
                    'row': i,
                    'department': dept_name,
                    'diagnosis': diag_name,
                    'status': status,
                    'is_valid': is_valid,
                    'errors': errors
                })
                
            return JsonResponse({
                'status': 'success',
                'preview': preview_data,
                'summary': {
                    'total': len(preview_data),
                    'valid': valid_count,
                    'invalid': invalid_count
                },
                'file_name': file.name
            })
            
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})

@csrf_exempt
def api_diag_dept_map_import(request):
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Invalid method'})
        
    try:
        import json
        data = json.loads(request.body)
        filename = data.get('filename', 'Unknown')
        rows = data.get('rows', [])
        update_duplicates = data.get('update_duplicates', False)
        
        dept_cache = {d.name.lower(): d for d in Department.objects.all()}
        diag_cache = {d.name.lower(): d for d in Diagnosis.objects.all()}
        
        imported = 0
        skipped = 0
        failed = 0
        
        with transaction.atomic():
            for row in rows:
                if row.get('status') == 'Error' or not row.get('is_valid', True):
                    failed += 1
                    continue
                    
                dept_name = str(row.get('department', '')).strip().lower()
                diag_name = str(row.get('diagnosis', '')).strip().lower()
                status = str(row.get('status_field', row.get('status', 'Active'))).strip()
                if status not in ['Active', 'Inactive']: status = 'Active'
                
                dept_obj = dept_cache.get(dept_name)
                diag_obj = diag_cache.get(diag_name)
                
                if dept_obj and diag_obj:
                    obj, created = DiagnosisDepartmentMapping.objects.get_or_create(
                        department=dept_obj,
                        diagnosis=diag_obj,
                        defaults={'status': status}
                    )
                    if created:
                        imported += 1
                    else:
                        if update_duplicates:
                            obj.status = status
                            obj.save()
                            imported += 1 # consider as imported/updated
                        else:
                            skipped += 1
                else:
                    failed += 1
                    
            history = DiagnosisDepartmentMappingImportHistory.objects.create(
                file_name=filename,
                uploaded_by=request.user if request.user.is_authenticated else None,
                total_records=len(rows),
                imported=imported,
                skipped=skipped,
                failed=failed,
                status='Completed'
            )
            
            return JsonResponse({
                'status': 'success',
                'summary': {
                    'total': len(rows),
                    'imported': imported,
                    'updated': 0,
                    'duplicates': skipped,
                    'failed': failed,
                    'history_id': history.id
                }
            })
            
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'status': 'error', 'message': str(e)})

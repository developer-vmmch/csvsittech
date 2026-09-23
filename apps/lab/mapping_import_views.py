"""
Smart Import Views for Lab Master Mapping Modules.

These views implement the "Smart Import" feature:
- Auto-create missing master records (Investigation, Parameter, Diagnosis, Department)
  before creating the mapping, so the user does NOT need to manually pre-create every
  master before bulk-importing mappings.
- Duplicate mappings are detected and skipped (never re-created).
- Case-insensitive, whitespace-normalised name matching prevents accidental duplicates.
- Full transaction safety: each row is wrapped atomically so a single failure does not
  leave orphan master records.
"""

import io
import json
import uuid

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import JsonResponse, HttpResponse
from django.views.generic import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _n(val):
    """Normalize: strip, collapse spaces, lower-case."""
    if val is None:
        return ""
    return " ".join(str(val).strip().split()).lower()


def _title(val):
    """Strip + title-case."""
    if val is None:
        return ""
    return " ".join(str(val).strip().split()).title()


def _auto_code(name, prefix="AUTO"):
    """Generate a short unique code when one is not supplied."""
    slug = "".join(c for c in name.upper().replace(" ", "_") if c.isalnum() or c == "_")[:20]
    return f"{prefix}-{slug}-{uuid.uuid4().hex[:4].upper()}"


def _safe_cell(row, headers, col):
    idx = headers.get(col)
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    if v is None:
        return ""
    return " ".join(str(v).strip().split())


# ---------------------------------------------------------------------------
# Investigation Parameter Mapping Import
# ---------------------------------------------------------------------------

class InvestigationParameterMappingImportView(LoginRequiredMixin, TemplateView):
    template_name = "lab/master/import_investigation_parameter_mapping.html"


def api_inv_param_map_download_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inv-Param Mapping"

    headers = ["investigation", "parameter", "active"]
    ws.append(headers)

    # Style header row
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 35

    # Sample rows
    ws.append(["Complete Blood Count (CBC)", "Hemoglobin", "Yes"])
    ws.append(["Complete Blood Count (CBC)", "WBC Count", "Yes"])
    ws.append(["Liver Function Test", "SGPT", "Yes"])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        'attachment; filename="investigation_parameter_mapping_template.xlsx"'
    )
    wb.save(response)
    return response


@csrf_exempt
def api_inv_param_map_preview(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})
    if "file" not in request.FILES:
        return JsonResponse({"success": False, "message": "No file uploaded"})

    file = request.FILES["file"]
    filename = file.name

    try:
        if filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        elif filename.endswith(".csv"):
            import csv
            decoded = file.read().decode("utf-8-sig")
            reader = csv.reader(io.StringIO(decoded))
            rows = list(reader)
        else:
            return JsonResponse({"success": False, "message": "Upload .xlsx or .csv"})

        if len(rows) < 2:
            return JsonResponse({"success": False, "message": "File is empty or has only headers"})

        raw_headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        headers = {h: i for i, h in enumerate(raw_headers)}

        for req in ("investigation", "parameter"):
            if req not in headers:
                return JsonResponse(
                    {"success": False, "message": f"Missing required column: {req}"}
                )

        from .models import Investigation, InvestigationParameter, Parameter

        inv_by_name = {_n(i.name): i for i in Investigation.objects.all()}
        param_by_name = {_n(p.name): p for p in Parameter.objects.all()}
        existing_mappings = set(
            InvestigationParameter.objects.filter(parameter__isnull=False)
            .values_list("investigation__name", "parameter__name")
        )
        existing_keys = {(_n(a), _n(b)) for a, b in existing_mappings}

        preview_data = []
        file_keys = set()
        valid = invalid = duplicate = 0

        for i, row in enumerate(rows[1:], start=2):
            if not any(cell for cell in row):
                continue

            inv_name_raw = _safe_cell(row, headers, "investigation")
            param_name_raw = _safe_cell(row, headers, "parameter")
            active_raw = _safe_cell(row, headers, "active").lower()
            is_active = active_raw not in ("no", "n", "0", "false")

            errors = []
            if not inv_name_raw:
                errors.append("Investigation name is required")
            if not param_name_raw:
                errors.append("Parameter name is required")

            inv_obj = inv_by_name.get(_n(inv_name_raw))
            param_obj = param_by_name.get(_n(param_name_raw))

            will_create_inv = bool(inv_name_raw) and inv_obj is None
            will_create_param = bool(param_name_raw) and param_obj is None

            key = (_n(inv_name_raw), _n(param_name_raw))
            is_dup = key in existing_keys
            is_file_dup = key in file_keys

            if errors:
                status = "Error"
                invalid += 1
            elif is_file_dup:
                status = "Error"
                errors.append("Duplicate row in file")
                invalid += 1
            elif is_dup:
                status = "Duplicate"
                duplicate += 1
            else:
                status = "Valid"
                valid += 1
                file_keys.add(key)

            preview_data.append({
                "row": i,
                "investigation": inv_name_raw,
                "parameter": param_name_raw,
                "active": is_active,
                "status": status,
                "will_create_investigation": will_create_inv,
                "will_create_parameter": will_create_param,
                "errors": errors,
                "is_valid": status == "Valid",
            })

        return JsonResponse({
            "success": True,
            "filename": filename,
            "preview": preview_data,
            "summary": {"total": len(preview_data), "valid": valid, "duplicate": duplicate, "invalid": invalid},
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})


@csrf_exempt
def api_inv_param_map_import(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})

    try:
        data = json.loads(request.body)
        rows = data.get("rows", [])
        filename = data.get("filename", "Unknown")
        skip_duplicates = not data.get("update_duplicates", False)

        from .models import Investigation, InvestigationParameter, Parameter

        imported = skipped = failed = 0
        errors_detail = []

        for row in rows:
            if row.get("status") == "Error":
                failed += 1
                continue
            if row.get("status") == "Duplicate" and skip_duplicates:
                skipped += 1
                continue

            inv_name = row.get("investigation", "").strip()
            param_name = row.get("parameter", "").strip()
            is_active = row.get("active", True)

            try:
                with transaction.atomic():
                    # 1. Resolve / auto-create Investigation
                    inv_obj = Investigation.objects.filter(name__iexact=inv_name).first()
                    if inv_obj is None:
                        inv_obj = Investigation.objects.create(
                            name=_title(inv_name),
                            code=_auto_code(inv_name, "INV"),
                            is_active=True,
                        )

                    # 2. Resolve / auto-create Parameter
                    param_obj = Parameter.objects.filter(name__iexact=param_name).first()
                    if param_obj is None:
                        param_obj = Parameter.objects.create(
                            name=_title(param_name),
                            code=_auto_code(param_name, "PARAM"),
                            is_active=True,
                        )

                    # 3. Create mapping if absent
                    _, created = InvestigationParameter.objects.get_or_create(
                        investigation=inv_obj,
                        parameter=param_obj,
                        defaults={
                            "name": param_obj.name,
                            "code": _auto_code(param_name, f"IP-{inv_obj.code[:6]}"),
                            "result_type": "Numeric",
                            "is_active": is_active,
                        },
                    )
                    if created:
                        imported += 1
                    else:
                        skipped += 1

            except Exception as e:
                failed += 1
                errors_detail.append({"row": row.get("row", "?"), "message": str(e)})

        return JsonResponse({
            "success": True,
            "summary": {
                "total": len(rows),
                "imported": imported,
                "skipped": skipped,
                "failed": failed,
            },
            "errors": errors_detail,
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})


# ---------------------------------------------------------------------------
# Diagnosis-Department Mapping — Smart Import (replaces existing basic version)
# ---------------------------------------------------------------------------

class DiagnosisDeptMappingImportView(LoginRequiredMixin, TemplateView):
    template_name = "lab/master/import_diagnosis_department_mapping.html"


def api_diag_dept_smart_download_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diag-Dept Mapping"

    headers = ["diagnosis", "department", "status"]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 30

    ws.append(["Fever", "General Medicine", "Active"])
    ws.append(["Diabetes", "Diabetology", "Active"])
    ws.append(["Hypertension", "Cardiology", "Active"])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        'attachment; filename="diagnosis_department_mapping_template.xlsx"'
    )
    wb.save(response)
    return response


@csrf_exempt
def api_diag_dept_smart_preview(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})
    if "file" not in request.FILES:
        return JsonResponse({"success": False, "message": "No file uploaded"})

    file = request.FILES["file"]
    filename = file.name

    try:
        if filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        elif filename.endswith(".csv"):
            import csv
            decoded = file.read().decode("utf-8-sig")
            reader = csv.reader(io.StringIO(decoded))
            rows = list(reader)
        else:
            return JsonResponse({"success": False, "message": "Upload .xlsx or .csv"})

        if len(rows) < 2:
            return JsonResponse({"success": False, "message": "File is empty or has only headers"})

        raw_headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        headers = {h: i for i, h in enumerate(raw_headers)}

        for req in ("diagnosis", "department"):
            if req not in headers:
                return JsonResponse(
                    {"success": False, "message": f"Missing required column: {req}"}
                )

        from .models import Diagnosis, DiagnosisDepartmentMapping
        from apps.patients.models import Department

        diag_by_name = {_n(d.name): d for d in Diagnosis.objects.all()}
        dept_by_name = {_n(d.name): d for d in Department.objects.all()}
        existing_keys = set(
            DiagnosisDepartmentMapping.objects.values_list("diagnosis__name", "department__name")
        )
        existing_keys = {(_n(a), _n(b)) for a, b in existing_keys}

        preview_data = []
        file_keys = set()
        valid = invalid = duplicate = 0

        for i, row in enumerate(rows[1:], start=2):
            if not any(cell for cell in row):
                continue

            diag_name_raw = _safe_cell(row, headers, "diagnosis")
            dept_name_raw = _safe_cell(row, headers, "department")
            status_raw = _safe_cell(row, headers, "status") or "Active"
            if status_raw not in ("Active", "Inactive"):
                status_raw = "Active"

            errors = []
            if not diag_name_raw:
                errors.append("Diagnosis name is required")
            if not dept_name_raw:
                errors.append("Department name is required")

            diag_obj = diag_by_name.get(_n(diag_name_raw))
            dept_obj = dept_by_name.get(_n(dept_name_raw))
            will_create_diag = bool(diag_name_raw) and diag_obj is None
            will_create_dept = bool(dept_name_raw) and dept_obj is None

            key = (_n(diag_name_raw), _n(dept_name_raw))
            is_dup = key in existing_keys
            is_file_dup = key in file_keys

            if errors:
                status = "Error"
                invalid += 1
            elif is_file_dup:
                status = "Error"
                errors.append("Duplicate row in file")
                invalid += 1
            elif is_dup:
                status = "Duplicate"
                duplicate += 1
            else:
                status = "Valid"
                valid += 1
                file_keys.add(key)

            preview_data.append({
                "row": i,
                "diagnosis": diag_name_raw,
                "department": dept_name_raw,
                "map_status": status_raw,
                "status": status,
                "will_create_diagnosis": will_create_diag,
                "will_create_department": will_create_dept,
                "errors": errors,
                "is_valid": status == "Valid",
            })

        return JsonResponse({
            "success": True,
            "filename": filename,
            "preview": preview_data,
            "summary": {"total": len(preview_data), "valid": valid, "duplicate": duplicate, "invalid": invalid},
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})


@csrf_exempt
def api_diag_dept_smart_import(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})

    try:
        data = json.loads(request.body)
        rows = data.get("rows", [])
        filename = data.get("filename", "Unknown")
        skip_duplicates = not data.get("update_duplicates", False)

        from .models import Diagnosis, DiagnosisDepartmentMapping, DiagnosisDepartmentMappingImportHistory
        from apps.patients.models import Department

        imported = skipped = failed = 0
        errors_detail = []

        for row in rows:
            if row.get("status") == "Error":
                failed += 1
                continue
            if row.get("status") == "Duplicate" and skip_duplicates:
                skipped += 1
                continue

            diag_name = row.get("diagnosis", "").strip()
            dept_name = row.get("department", "").strip()
            map_status = row.get("map_status", "Active")
            if map_status not in ("Active", "Inactive"):
                map_status = "Active"

            try:
                with transaction.atomic():
                    # 1. Resolve / auto-create Diagnosis
                    diag_obj = Diagnosis.objects.filter(name__iexact=diag_name).first()
                    if diag_obj is None:
                        diag_obj = Diagnosis.objects.create(
                            name=_title(diag_name),
                            code=_auto_code(diag_name, "DX"),
                            is_active=True,
                        )

                    # 2. Resolve / auto-create Department (patients app)
                    dept_obj = Department.objects.filter(name__iexact=dept_name).first()
                    if dept_obj is None:
                        dept_obj = Department.objects.create(
                            name=_title(dept_name),
                            is_active=True,
                        )

                    # 3. Create mapping if absent
                    _, created = DiagnosisDepartmentMapping.objects.get_or_create(
                        diagnosis=diag_obj,
                        department=dept_obj,
                        defaults={"status": map_status},
                    )
                    if created:
                        imported += 1
                    else:
                        skipped += 1

            except Exception as e:
                failed += 1
                errors_detail.append({"row": row.get("row", "?"), "message": str(e)})

        # Record history
        try:
            DiagnosisDepartmentMappingImportHistory.objects.create(
                file_name=filename,
                uploaded_by=request.user if request.user.is_authenticated else None,
                total_records=len(rows),
                imported=imported,
                skipped=skipped,
                failed=failed,
                status="Completed",
            )
        except Exception:
            pass

        return JsonResponse({
            "success": True,
            "summary": {
                "total": len(rows),
                "imported": imported,
                "skipped": skipped,
                "failed": failed,
            },
            "errors": errors_detail,
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})


# ---------------------------------------------------------------------------
# Diagnosis-Investigation Mapping — Smart Import (replaces row-ID-based version)
# ---------------------------------------------------------------------------

@csrf_exempt
def api_diag_inv_smart_preview(request):
    """
    Smart preview: auto-resolves Diagnosis and Investigation by name.
    Missing masters are flagged as 'will_create_*' so the user sees the plan.
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})
    if "file" not in request.FILES:
        return JsonResponse({"success": False, "message": "No file uploaded"})

    file = request.FILES["file"]
    filename = file.name

    try:
        if filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        elif filename.endswith(".csv"):
            import csv
            decoded = file.read().decode("utf-8-sig")
            reader = csv.reader(io.StringIO(decoded))
            rows = list(reader)
        else:
            return JsonResponse({"success": False, "message": "Upload .xlsx or .csv"})

        if len(rows) < 2:
            return JsonResponse({"success": False, "message": "File is empty"})

        raw_headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        headers = {h: i for i, h in enumerate(raw_headers)}

        for req in ("diagnosis", "investigation"):
            if req not in headers:
                return JsonResponse({"success": False, "message": f"Missing column: {req}"})

        from .models import Diagnosis, Investigation, DiagnosisInvestigationMap, AgeGroup

        diag_by_name = {_n(d.name): d for d in Diagnosis.objects.all()}
        inv_by_name = {_n(i.name): i for i in Investigation.objects.all()}
        # Default age group — the first active one or None
        default_ag = AgeGroup.objects.filter(is_active=True).order_by("sort_order").first()

        existing_keys = set(
            DiagnosisInvestigationMap.objects.values_list(
                "diagnosis__name", "investigation__name"
            )
        )
        existing_keys = {(_n(a), _n(b)) for a, b in existing_keys}

        preview_data = []
        file_keys = set()
        valid = invalid = duplicate = 0

        for i, row in enumerate(rows[1:], start=2):
            if not any(cell for cell in row):
                continue

            diag_name_raw = _safe_cell(row, headers, "diagnosis")
            inv_name_raw = _safe_cell(row, headers, "investigation")
            active_raw = _safe_cell(row, headers, "active").lower()
            is_active = active_raw not in ("no", "n", "0", "false")

            errors = []
            if not diag_name_raw:
                errors.append("Diagnosis name is required")
            if not inv_name_raw:
                errors.append("Investigation name is required")

            diag_obj = diag_by_name.get(_n(diag_name_raw))
            inv_obj = inv_by_name.get(_n(inv_name_raw))
            will_create_diag = bool(diag_name_raw) and diag_obj is None
            will_create_inv = bool(inv_name_raw) and inv_obj is None

            if not default_ag:
                errors.append("No Age Group found in master. Please create at least one Age Group first.")

            key = (_n(diag_name_raw), _n(inv_name_raw))
            is_dup = key in existing_keys
            is_file_dup = key in file_keys

            if errors:
                status = "Error"
                invalid += 1
            elif is_file_dup:
                status = "Error"
                errors.append("Duplicate row in file")
                invalid += 1
            elif is_dup:
                status = "Duplicate"
                duplicate += 1
            else:
                status = "Valid"
                valid += 1
                file_keys.add(key)

            preview_data.append({
                "row": i,
                "diagnosis": diag_name_raw,
                "investigation": inv_name_raw,
                "active": is_active,
                "status": status,
                "will_create_diagnosis": will_create_diag,
                "will_create_investigation": will_create_inv,
                "errors": errors,
                "is_valid": status == "Valid",
            })

        return JsonResponse({
            "success": True,
            "filename": filename,
            "preview": preview_data,
            "summary": {"total": len(preview_data), "valid": valid, "duplicate": duplicate, "invalid": invalid},
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})


@csrf_exempt
def api_diag_inv_smart_import(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid method"})

    try:
        data = json.loads(request.body)
        rows = data.get("rows", [])
        skip_duplicates = not data.get("update_duplicates", False)

        from .models import Diagnosis, Investigation, DiagnosisInvestigationMap, AgeGroup, DiagnosisImportHistory

        default_ag = AgeGroup.objects.filter(is_active=True).order_by("sort_order").first()

        imported = skipped = failed = 0
        errors_detail = []

        for row in rows:
            if row.get("status") == "Error":
                failed += 1
                continue
            if row.get("status") == "Duplicate" and skip_duplicates:
                skipped += 1
                continue

            diag_name = row.get("diagnosis", "").strip()
            inv_name = row.get("investigation", "").strip()
            is_active = row.get("active", True)

            try:
                with transaction.atomic():
                    if not default_ag:
                        raise ValueError("No Age Group master found. Create one first.")

                    diag_obj = Diagnosis.objects.filter(name__iexact=diag_name).first()
                    if diag_obj is None:
                        diag_obj = Diagnosis.objects.create(
                            name=_title(diag_name),
                            code=_auto_code(diag_name, "DX"),
                            is_active=True,
                        )

                    inv_obj = Investigation.objects.filter(name__iexact=inv_name).first()
                    if inv_obj is None:
                        inv_obj = Investigation.objects.create(
                            name=_title(inv_name),
                            code=_auto_code(inv_name, "INV"),
                            is_active=True,
                        )

                    _, created = DiagnosisInvestigationMap.objects.get_or_create(
                        diagnosis=diag_obj,
                        age_group=default_ag,
                        investigation=inv_obj,
                        defaults={"is_active": is_active},
                    )
                    if created:
                        imported += 1
                    else:
                        skipped += 1

            except Exception as e:
                failed += 1
                errors_detail.append({"row": row.get("row", "?"), "message": str(e)})

        return JsonResponse({
            "success": True,
            "summary": {
                "total": len(rows),
                "imported": imported,
                "skipped": skipped,
                "failed": failed,
            },
            "errors": errors_detail,
        })

    except Exception as e:
        import traceback; traceback.print_exc()
        return JsonResponse({"success": False, "message": str(e)})

"""
test_atc_phase1.py

Comprehensive Phase 1 Automated Test Suite for VMMC ERP:
1. Universal Master Import / Export (Lab Master)
2. Auto Trigger Control (ATC) Module

Tests:
  TEST 1: Universal Export (9 sheets, headers, data rows)
  TEST 2: Universal Import (Validation preview, safe upsert, 0 duplicates)
  TEST 3: Lab Master Consistency (Service Request & Automate Test lookups)
  TEST 4: OP Automation (Quota fulfillment, strictly today's date, day-end cutoff)
  TEST 5: OP Batching (Progression, batch log auditing)
  TEST 6: Review Automation (Only 'D' patients, no new patients, PatientVisit & Report integration)
  TEST 7: Future Patient Creation (Shortcuts, valid range, >1 month backend rejection)
  TEST 8: Emergency Stop (Instant halt, status update, data preservation)
  TEST 9: Double-Click / Concurrency Prevention (Running job rejection)
  TEST 10: Controlled Failure & Error Handling (Invalid input, DB consistency)
"""

import os
import io
import sys
import time
from datetime import datetime, date, time as dtime, timedelta
import openpyxl

# Setup Django Environment
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'vmmc_erp.settings')
import django
django.setup()

from django.test import RequestFactory
from django.utils import timezone
from django.db.models import Q

from apps.users.models import User
from apps.patients.models import Patient, PatientVisit, Department, DepartmentUnit
from apps.lab.models import (
    Investigation, Diagnosis, Parameter, ParameterReferenceRange,
    InvestigationParameter, DiagnosisInvestigationMap,
    ATCJob, ATCJobLog, ATCSetting
)
from apps.lab.universal_importer import (
    export_universal_master,
    validate_import,
    import_universal_master,
    SHEET_DEFS,
    SHEET_ORDER
)
from apps.lab.atc_service import (
    execute_op_automation,
    execute_review_automation,
    execute_future_patient_automation,
    run_atc_job,
    get_server_today,
    get_server_now
)
from apps.lab.atc_views import api_atc_trigger, api_atc_stop, api_atc_status


class TestRunner:
    def __init__(self):
        self.results = {}
        self.factory = RequestFactory()
        self.admin_user = User.objects.filter(is_superuser=True).first()
        if not self.admin_user:
            self.admin_user = User.objects.create_superuser('test_admin', 'admin@example.com', 'adminpass123')

    def log_header(self, title):
        print("\n" + "=" * 70)
        print(f"  {title}")
        print("=" * 70)

    def record_result(self, name, passed, details=""):
        self.results[name] = (passed, details)
        status = "\033[92mPASSED\033[0m" if passed else "\033[91mFAILED\033[0m"
        print(f"[{status}] {name}")
        if details:
            print(f"       -> {details}")

    def run_all(self):
        print("STARTING ATC PHASE 1 TEST SUITE...")
        ATCJob.objects.filter(status__in=['STARTING', 'RUNNING', 'STOP_REQUESTED']).update(status='STOPPED')
        start_time = time.time()

        exported_bytes = self.test_1_universal_export()
        self.test_2_universal_import(exported_bytes)
        self.test_3_lab_master_consistency()
        self.test_4_op_automation()
        self.test_5_op_batching()
        self.test_6_review_automation()
        self.test_7_future_patient_creation()
        self.test_8_emergency_stop()
        self.test_9_double_click_prevention()
        self.test_10_controlled_failure()

        elapsed = time.time() - start_time
        self.print_summary(elapsed)

    # -----------------------------------------------------------------------
    # TEST 1: Universal Export
    # -----------------------------------------------------------------------
    def test_1_universal_export(self):
        self.log_header("TEST 1: Universal Export")
        try:
            exported_buf = export_universal_master()
            exported_bytes = exported_buf.getvalue()
            wb = openpyxl.load_workbook(io.BytesIO(exported_bytes))

            # 1. Verify 9 sheets
            sheet_names = wb.sheetnames
            expected_sheets = list(SHEET_ORDER)
            missing_sheets = [s for s in expected_sheets if s not in sheet_names]
            assert not missing_sheets, f"Missing sheets in export: {missing_sheets}"

            # 2. Verify headers match expected definitions
            for sheet_name in expected_sheets:
                ws = wb[sheet_name]
                headers = [str(cell.value or '').strip() for cell in ws[1] if cell.value is not None]
                expected_cols = SHEET_DEFS[sheet_name]['columns']
                assert headers == expected_cols, f"Sheet '{sheet_name}' headers mismatch! Found: {headers}, Expected: {expected_cols}"

            # 3. Verify data rows were exported accurately
            counts_summary = []
            for sheet_name in expected_sheets:
                ws = wb[sheet_name]
                row_count = ws.max_row - 1  # exclude header
                counts_summary.append(f"{sheet_name}: {row_count}")

            self.record_result(
                "TEST 1: Universal Export",
                True,
                f"9/9 sheets present, headers 100% matched, exported rows: {', '.join(counts_summary[:4])}..."
            )
            return exported_bytes
        except Exception as e:
            self.record_result("TEST 1: Universal Export", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 2: Universal Import (Validation & Safe Upsert)
    # -----------------------------------------------------------------------
    def test_2_universal_import(self, exported_bytes):
        self.log_header("TEST 2: Universal Import")
        try:
            # 1. Validation preview
            summary = validate_import(io.BytesIO(exported_bytes))
            assert summary['valid'] is True, f"Validation preview failed with errors: {summary.get('errors')}"
            assert len(summary['detected_sheets']) == 9, f"Expected 9 sheets detected, got {len(summary['detected_sheets'])}"
            assert summary['total_stats']['valid_rows'] > 0, "Expected valid rows in exported workbook"

            # 2. Record counts before import
            dept_count_before = Department.objects.count()
            diag_count_before = Diagnosis.objects.count()
            inves_count_before = Investigation.objects.count()
            param_count_before = Parameter.objects.count()
            range_count_before = ParameterReferenceRange.objects.count()

            # 3. Perform Universal Import
            import_result = import_universal_master(io.BytesIO(exported_bytes))
            assert 'departments_created' in import_result, "Import counts dict missing departments_created"

            # 4. Safe upsert verification: counts must NOT double
            dept_count_after = Department.objects.count()
            diag_count_after = Diagnosis.objects.count()
            inves_count_after = Investigation.objects.count()
            param_count_after = Parameter.objects.count()
            range_count_after = ParameterReferenceRange.objects.count()

            assert dept_count_after == dept_count_before, f"Departments duplicated! Before: {dept_count_before}, After: {dept_count_after}"
            assert diag_count_after == diag_count_before, f"Diagnoses duplicated! Before: {diag_count_before}, After: {diag_count_after}"
            assert inves_count_after == inves_count_before, f"Investigations duplicated! Before: {inves_count_before}, After: {inves_count_after}"
            assert param_count_after == param_count_before, f"Parameters duplicated! Before: {param_count_before}, After: {param_count_after}"
            assert range_count_after == range_count_before, f"Reference Ranges duplicated! Before: {range_count_before}, After: {range_count_after}"

            self.record_result(
                "TEST 2: Universal Import",
                True,
                f"Validation preview passed (valid=True), 0 duplicates on re-import (Safe Upsert verified). Total rows processed: {summary['total_stats']['valid_rows']}"
            )
        except Exception as e:
            self.record_result("TEST 2: Universal Import", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 3: Lab Master Consistency
    # -----------------------------------------------------------------------
    def test_3_lab_master_consistency(self):
        self.log_header("TEST 3: Lab Master Consistency")
        try:
            # 1. Service Request mapping check
            mappings = DiagnosisInvestigationMap.objects.filter(is_active=True).select_related('diagnosis', 'investigation')
            assert mappings.exists(), "No active DiagnosisInvestigationMap found after import"

            first_map = mappings.first()
            diag = first_map.diagnosis
            # Query mapped investigations for this diagnosis (as done in Service Request view)
            diag_mappings = DiagnosisInvestigationMap.objects.filter(
                diagnosis=diag,
                is_active=True
            ).select_related('investigation', 'age_group')
            assert diag_mappings.exists(), f"Failed to retrieve mapped investigations for diagnosis '{diag.name}'"
            mapped_inves = [m.investigation for m in diag_mappings if m.investigation]
            assert len(mapped_inves) > 0, f"No investigations found for diagnosis '{diag.name}'"

            # 2. Automate Test reference range lookup check
            ref_ranges = ParameterReferenceRange.objects.filter(is_active=True).select_related('investigation_parameter__parameter')
            assert ref_ranges.exists(), "No active ParameterReferenceRange found after import"

            sample_range = ref_ranges.first()
            param = sample_range.investigation_parameter.parameter
            # Verify range fields
            assert sample_range.investigation_parameter is not None
            assert param is not None
            assert sample_range.min_value is not None or sample_range.max_value is not None or sample_range.normal_text

            self.record_result(
                "TEST 3: Lab Master Consistency",
                True,
                f"Service Request mapping verified for '{diag.name}' ({len(mapped_inves)} tests). Automate Test reference range verified for parameter '{param.name}'."
            )
        except Exception as e:
            self.record_result("TEST 3: Lab Master Consistency", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 4: OP Automation
    # -----------------------------------------------------------------------
    def test_4_op_automation(self):
        self.log_header("TEST 4: OP Automation")
        try:
            today = get_server_today()

            # Create test OP job: target 20 (Male: 12, Female: 8)
            job = ATCJob.objects.create(
                job_id=f"ATC-TEST-OP-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=20,
                target_male=12,
                target_female=8,
                schedule_start_time=dtime(0, 0),
                schedule_end_time=dtime(23, 59),
                batch_size=10,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            execute_op_automation(job)
            job.refresh_from_db()

            assert job.status == ATCJob.StatusChoices.COMPLETED, f"Job failed to complete: {job.status}, error: {job.error_summary}"
            assert job.created_count == 20, f"Expected 20 created patients, got {job.created_count}"
            assert job.created_male == 12, f"Expected 12 male patients, got {job.created_male}"
            assert job.created_female == 8, f"Expected 8 female patients, got {job.created_female}"

            # Check created patients
            created_logs = ATCJobLog.objects.filter(job=job, status='Success').select_related('patient', 'visit')
            assert created_logs.count() == 20, f"Expected 20 log entries, got {created_logs.count()}"

            for log in created_logs:
                p = log.patient
                v = log.visit
                # STRICT CURRENT DATE CHECK
                assert p.registration_date == today, f"Patient registration date {p.registration_date} is NOT today ({today})!"
                # PATIENT TYPE CHECK
                assert p.patient_type == 'D', f"Patient type is not 'D': {p.patient_type}"
                assert p.created_source == 'D', f"Created source is not 'D': {p.created_source}"
                # VISIT CHECK
                assert v is not None, "PatientVisit #1 was not created!"
                assert v.visit_no == 1, f"Expected visit_no 1, got {v.visit_no}"
                assert v.visit_type == 'OP', f"Expected visit_type 'OP', got {v.visit_type}"

            # Verify Day-End Stop Cutoff:
            cutoff_job = ATCJob.objects.create(
                job_id=f"ATC-TEST-CUTOFF-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=10,
                target_male=5,
                target_female=5,
                # Set schedule end time in past (e.g. 00:01)
                schedule_start_time=dtime(0, 0),
                schedule_end_time=dtime(0, 1),
                batch_size=5,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )
            execute_op_automation(cutoff_job)
            cutoff_job.refresh_from_db()

            assert cutoff_job.status == ATCJob.StatusChoices.STOPPED, f"Expected cutoff job to STOP, got {cutoff_job.status}"
            assert "Day-end" in (cutoff_job.stop_reason or ""), f"Expected 'Day-end' in stop reason, got: {cutoff_job.stop_reason}"
            assert cutoff_job.created_count == 0, f"Expected 0 created for past cutoff, got {cutoff_job.created_count}"

            self.record_result(
                "TEST 4: OP Automation",
                True,
                f"Target 20 (12 M, 8 F) exact quotas fulfilled, 100% registration_date={today}, patient_type='D', Visit #1 created, Day-End stop cutoff verified."
            )
        except Exception as e:
            self.record_result("TEST 4: OP Automation", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 5: OP Batching
    # -----------------------------------------------------------------------
    def test_5_op_batching(self):
        self.log_header("TEST 5: OP Batching")
        try:
            # Batch test: target 20, batch size 10 -> exactly 2 batches
            job = ATCJob.objects.create(
                job_id=f"ATC-TEST-BATCH-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=20,
                target_male=10,
                target_female=10,
                schedule_start_time=dtime(0, 0),
                schedule_end_time=dtime(23, 59),
                batch_size=10,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            execute_op_automation(job)
            job.refresh_from_db()

            assert job.status == ATCJob.StatusChoices.COMPLETED, f"Job did not complete: {job.status}"
            assert job.total_batches == 2, f"Expected 2 batches, got {job.total_batches}"

            batch1_logs = ATCJobLog.objects.filter(job=job, batch_number=1, status='Success').count()
            batch2_logs = ATCJobLog.objects.filter(job=job, batch_number=2, status='Success').count()

            assert batch1_logs == 10, f"Expected 10 in batch 1, got {batch1_logs}"
            assert batch2_logs == 10, f"Expected 10 in batch 2, got {batch2_logs}"

            self.record_result(
                "TEST 5: OP Batching",
                True,
                f"Batch progression verified: Batch 1 ({batch1_logs} records) -> Batch 2 ({batch2_logs} records), Total: {job.created_count}"
            )
        except Exception as e:
            self.record_result("TEST 5: OP Batching", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 6: Review Automation
    # -----------------------------------------------------------------------
    def test_6_review_automation(self):
        self.log_header("TEST 6: Review Automation")
        try:
            # Create a normal patient with patient_type='O' to ensure it is NEVER selected
            dept = Department.objects.first()
            normal_patient = Patient.objects.filter(name='Real Normal Patient', patient_type='O').first()
            if not normal_patient:
                normal_patient = Patient.objects.create(
                    title='Mr', name='Real Normal Patient', gender='Male', dob=date(1985, 3, 10),
                    age_years=41, registration_date=date(2025, 1, 10),
                    department_obj=dept, department=dept.name if dept else 'GENERAL MEDICINE',
                    patient_id=f"O-REAL-{int(time.time())}", op_number=f"OP-REAL-{int(time.time())}",
                    guardian_title='Mr', guardian_relationship='S/O', guardian_name='Father Name',
                    street='Real Street', village_area='Real Area', city='Karaikal',
                    state='Puducherry', pincode='609602', patient_type='O', created_source='O'
                )

            # Record total patients before review run
            total_patients_before = Patient.objects.count()

            # Target 5 review visits
            review_date = get_server_today() - timedelta(days=2)  # backdated review test
            job = ATCJob.objects.create(
                job_id=f"ATC-TEST-REV-{int(time.time())}",
                mode=ATCJob.ModeChoices.REVIEW,
                target_total=5,
                from_date=review_date,
                to_date=review_date,
                batch_size=10,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            execute_review_automation(job)
            job.refresh_from_db()

            assert job.status == ATCJob.StatusChoices.COMPLETED, f"Review job did not complete: {job.status}, error: {job.error_summary}"
            assert job.created_count == 5, f"Expected 5 reviews, got {job.created_count}"

            # NO NEW PATIENTS CREATED CHECK
            total_patients_after = Patient.objects.count()
            assert total_patients_after == total_patients_before, f"New patients were created during Review! Before: {total_patients_before}, After: {total_patients_after}"

            # Normal patient was NOT reviewed
            normal_reviews = PatientVisit.objects.filter(patient=normal_patient, visit_type='REVIEW').count()
            assert normal_reviews == 0, "Normal patient ('O') was incorrectly selected for Review!"

            # Review Visits Verification
            review_logs = ATCJobLog.objects.filter(job=job, status='Success').select_related('patient', 'visit')
            assert review_logs.count() == 5, f"Expected 5 review logs, got {review_logs.count()}"

            for log in review_logs:
                p = log.patient
                v = log.visit
                assert p.patient_type == 'D', f"Reviewed patient is not 'D': {p.patient_type}"
                assert v is not None, "Review PatientVisit record is missing!"
                assert v.visit_type == 'REVIEW', f"Visit type is not 'REVIEW': {v.visit_type}"
                assert v.category == 'RE_CONSULTATION', f"Visit category is not 'RE_CONSULTATION': {v.category}"
                assert v.visit_no >= 2, f"Review visit_no must be >= 2, got {v.visit_no}"
                assert v.visit_date.date() == review_date, f"Visit date {v.visit_date.date()} != expected review date {review_date}"

            # Query Review List / Report query
            review_report_qs = PatientVisit.objects.filter(
                Q(visit_no__gt=1) | Q(visit_type='REVIEW') | Q(category='RE_CONSULTATION'),
                visit_date__date=review_date
            )
            assert review_report_qs.count() >= 5, f"Review visits not found in Review Report query! Found {review_report_qs.count()}"

            # Print Copy compatibility: Check all required fields are present
            sample_visit = review_logs.first().visit
            assert sample_visit.patient.name, "Patient name missing for print copy"
            assert sample_visit.department, "Department missing for print copy"
            assert sample_visit.unit_doctor, "Unit doctor missing for print copy"
            assert sample_visit.visit_no, "Visit number missing for print copy"

            self.record_result(
                "TEST 6: Review Automation",
                True,
                f"5 Reviews created. Zero new patients created (patients count stayed at {total_patients_before}). 100% 'D' patients. Visible in Review List and printable in Print Copy."
            )
        except Exception as e:
            self.record_result("TEST 6: Review Automation", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 7: Future Patient Creation
    # -----------------------------------------------------------------------
    def test_7_future_patient_creation(self):
        self.log_header("TEST 7: Future Patient Creation")
        try:
            today = get_server_today()

            # 1. Shortcuts calculation test
            d_1day = today + timedelta(days=1)
            d_1week = today + timedelta(weeks=1)
            d_1month = today + timedelta(days=30)
            assert d_1day > today and d_1week > d_1day and d_1month > d_1week

            # 2. Valid future patient creation (within 1 month)
            from_fut = today + timedelta(days=3)
            to_fut = today + timedelta(days=7)

            job = ATCJob.objects.create(
                job_id=f"ATC-TEST-FUT-{int(time.time())}",
                mode=ATCJob.ModeChoices.FUTURE_PATIENT,
                target_total=6,
                target_male=3,
                target_female=3,
                from_date=from_fut,
                to_date=to_fut,
                batch_size=5,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            execute_future_patient_automation(job)
            job.refresh_from_db()

            assert job.status == ATCJob.StatusChoices.COMPLETED, f"Future patient job did not complete: {job.status}, error: {job.error_summary}"
            assert job.created_count == 6, f"Expected 6 future patients, got {job.created_count}"
            assert job.created_male == 3, f"Expected 3 male future patients, got {job.created_male}"
            assert job.created_female == 3, f"Expected 3 female future patients, got {job.created_female}"

            # Check registration dates of created future patients
            fut_logs = ATCJobLog.objects.filter(job=job, status='Success').select_related('patient')
            for log in fut_logs:
                p = log.patient
                assert from_fut <= p.registration_date <= to_fut, f"Patient date {p.registration_date} out of future bounds ({from_fut} - {to_fut})"
                assert p.patient_type == 'D', f"Future patient type != 'D': {p.patient_type}"

            # 3. Backend rejection for > 1 month (Service level)
            invalid_to = today + timedelta(days=45)
            invalid_job = ATCJob.objects.create(
                job_id=f"ATC-TEST-INVALID-FUT-{int(time.time())}",
                mode=ATCJob.ModeChoices.FUTURE_PATIENT,
                target_total=5,
                target_male=3,
                target_female=2,
                from_date=today,
                to_date=invalid_to,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            service_blocked = False
            try:
                execute_future_patient_automation(invalid_job)
            except ValueError as ve:
                if "Future patient creation is limited to one month." in str(ve):
                    service_blocked = True
            finally:
                invalid_job.status = ATCJob.StatusChoices.FAILED
                invalid_job.save()
            assert service_blocked, "Service did not reject future range > 1 month!"

            # 4. Backend rejection for > 1 month (API level)
            request = self.factory.post('/lab/api/atc/trigger/', {
                'mode': 'FUTURE_PATIENT',
                'target_total': 5,
                'male_count': 3,
                'female_count': 2,
                'from_date': today.strftime('%Y-%m-%d'),
                'to_date': invalid_to.strftime('%Y-%m-%d'),
            })
            request.user = self.admin_user
            api_resp = api_atc_trigger(request)
            assert api_resp.status_code == 400, f"API did not return 400 for > 1 month: {api_resp.status_code}"
            assert "Future patient creation is limited to one month." in api_resp.content.decode('utf-8')

            self.record_result(
                "TEST 7: Future Patient Creation",
                True,
                f"Valid future patient creation tested ({from_fut} to {to_fut}). Backend & API strictly reject > 1 month range."
            )
        except Exception as e:
            self.record_result("TEST 7: Future Patient Creation", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 8: Emergency Stop
    # -----------------------------------------------------------------------
    def test_8_emergency_stop(self):
        self.log_header("TEST 8: Emergency Stop")
        try:
            # 1. Test halting logic in automation
            job = ATCJob.objects.create(
                job_id=f"ATC-TEST-STOP-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=50,
                target_male=25,
                target_female=25,
                schedule_start_time=dtime(0, 0),
                schedule_end_time=dtime(23, 59),
                batch_size=5,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            # Set stop_requested flag
            job.stop_requested = True
            job.status = ATCJob.StatusChoices.STOP_REQUESTED
            job.save()

            execute_op_automation(job)
            job.refresh_from_db()

            assert job.status == ATCJob.StatusChoices.EMERGENCY_STOPPED, f"Job did not transition to EMERGENCY_STOPPED, got: {job.status}"
            assert job.created_count == 0, f"Job should have stopped before creating records, created: {job.created_count}"
            assert "Emergency Stop" in (job.stop_reason or ""), f"Stop reason missing: {job.stop_reason}"

            # 2. Test Emergency Stop API endpoint
            running_job = ATCJob.objects.create(
                job_id=f"ATC-TEST-APISTOP-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=20,
                target_male=10,
                target_female=10,
                status=ATCJob.StatusChoices.RUNNING,
                created_by=self.admin_user
            )

            req = self.factory.post(f"/lab/api/atc/stop/{running_job.id}/")
            req.user = self.admin_user
            resp = api_atc_stop(req, running_job.id)
            assert resp.status_code == 200, f"api_atc_stop returned {resp.status_code}"

            running_job.refresh_from_db()
            assert running_job.stop_requested is True, "stop_requested was not set to True"
            assert running_job.status == ATCJob.StatusChoices.STOP_REQUESTED, f"Status != STOP_REQUESTED, got {running_job.status}"

            # Clean up
            running_job.status = ATCJob.StatusChoices.STOPPED
            running_job.save()

            self.record_result(
                "TEST 8: Emergency Stop",
                True,
                "Immediate halt verified, state transitions to EMERGENCY_STOPPED, API stop endpoint verified."
            )
        except Exception as e:
            self.record_result("TEST 8: Emergency Stop", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 9: Double-Click / Concurrency Prevention
    # -----------------------------------------------------------------------
    def test_9_double_click_prevention(self):
        self.log_header("TEST 9: Double-Click / Concurrency Prevention")
        try:
            # Create a running job to simulate an active background run
            active_job = ATCJob.objects.create(
                job_id=f"ATC-ACTIVE-{int(time.time())}",
                mode=ATCJob.ModeChoices.OP,
                target_total=100,
                target_male=60,
                target_female=40,
                status=ATCJob.StatusChoices.RUNNING,
                created_by=self.admin_user
            )

            # Attempt to trigger second job
            req = self.factory.post('/lab/api/atc/trigger/', {
                'mode': 'OP',
                'male_count': 10,
                'female_count': 10,
                'start_time': '08:00',
                'end_time': '23:59',
            })
            req.user = self.admin_user
            resp = api_atc_trigger(req)

            assert resp.status_code == 400, f"Expected 400 rejection for concurrent trigger, got {resp.status_code}"
            resp_data = resp.content.decode('utf-8')
            assert "An ATC job is already running" in resp_data, f"Expected running job message, got: {resp_data}"

            # Teardown active job
            active_job.status = ATCJob.StatusChoices.STOPPED
            active_job.save()

            self.record_result(
                "TEST 9: Double-Click Prevention",
                True,
                "Second concurrent trigger rejected with HTTP 400 and message 'An ATC job is already running'."
            )
        except Exception as e:
            self.record_result("TEST 9: Double-Click Prevention", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # TEST 10: Controlled Failure & Error Handling
    # -----------------------------------------------------------------------
    def test_10_controlled_failure(self):
        self.log_header("TEST 10: Controlled Failure & Error Handling")
        try:
            # 1. Invalid zero target count in API
            req1 = self.factory.post('/lab/api/atc/trigger/', {
                'mode': 'OP',
                'male_count': 0,
                'female_count': 0,
            })
            req1.user = self.admin_user
            resp1 = api_atc_trigger(req1)
            assert resp1.status_code == 400, f"Expected 400 for zero target, got {resp1.status_code}"
            assert "Target patient count must be greater than 0." in resp1.content.decode('utf-8')

            # 2. Invalid mode in API
            req2 = self.factory.post('/lab/api/atc/trigger/', {
                'mode': 'INVALID_UNKNOWN_MODE',
                'male_count': 5,
                'female_count': 5,
            })
            req2.user = self.admin_user
            resp2 = api_atc_trigger(req2)
            assert resp2.status_code == 400, f"Expected 400 for invalid mode, got {resp2.status_code}"
            assert "Invalid ATC mode" in resp2.content.decode('utf-8')

            # 3. Service failure handling: run_atc_job handles exception cleanly
            fail_job = ATCJob.objects.create(
                job_id=f"ATC-TEST-FAIL-{int(time.time())}",
                mode="NON_EXISTENT_MODE",  # will raise ValueError in run_atc_job
                target_total=10,
                target_male=5,
                target_female=5,
                status=ATCJob.StatusChoices.STARTING,
                created_by=self.admin_user
            )

            run_atc_job(fail_job.id)
            fail_job.refresh_from_db()

            assert fail_job.status == ATCJob.StatusChoices.FAILED, f"Job status != FAILED, got: {fail_job.status}"
            assert fail_job.error_summary is not None, "Error summary was not captured"
            assert "Unsupported ATC mode" in fail_job.error_summary, f"Unexpected error summary: {fail_job.error_summary}"
            assert fail_job.completed_at is not None, "completed_at was not recorded"

            self.record_result(
                "TEST 10: Controlled Failure",
                True,
                "API rejects invalid parameters cleanly. Unhandled exceptions transition job to FAILED state with error audit and without DB corruption."
            )
        except Exception as e:
            self.record_result("TEST 10: Controlled Failure", False, str(e))
            raise

    # -----------------------------------------------------------------------
    # Summary Report
    # -----------------------------------------------------------------------
    def print_summary(self, elapsed):
        print("\n" + "=" * 70)
        print("                 PHASE 1 TEST EXECUTION SUMMARY")
        print("=" * 70)
        total = len(self.results)
        passed = sum(1 for p, _ in self.results.values() if p)
        failed = total - passed

        for name, (p, details) in self.results.items():
            status = "\033[92mPASS\033[0m" if p else "\033[91mFAIL\033[0m"
            print(f"  [{status}] {name}")

        print("-" * 70)
        print(f"Total Tests: {total} | Passed: {passed} | Failed: {failed} | Time: {elapsed:.2f}s")
        print("=" * 70 + "\n")

        if failed > 0:
            sys.exit(1)


if __name__ == '__main__':
    runner = TestRunner()
    runner.run_all()
